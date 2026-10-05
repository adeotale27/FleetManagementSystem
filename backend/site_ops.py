"""Tenant-scoped multi-site transport operations."""
import asyncio
import csv
import hashlib
import io
import re
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from difflib import SequenceMatcher
from typing import Any, Literal, Optional, overload
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from bson.decimal128 import Decimal128
from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from fastapi.concurrency import run_in_threadpool
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, ConfigDict, Field, field_validator, model_validator
from pymongo import ReturnDocument
from pymongo.errors import DuplicateKeyError

import ledger as L
import storage as S
from auth import hash_manager_pw, site_user
from db import current_db_name, db, new_id, now_iso, platform_db

router = APIRouter()
_indexed_databases = set()
_index_lock = asyncio.Lock()

DEFAULT_GOODS = ["General Goods", "Food Grains", "Cement", "Steel", "Fertilizer", "Other"]
DEFAULT_CONTAINERS = ["Box", "Open Parcel", "Container", "Jute Bags", "Carton", "Drum", "Other"]
MANAGER_PERMISSIONS = {
    "dashboard:read", "trips:read", "trips:create", "trips:update",
    "lrs:read", "lrs:create", "lrs:update", "finance:read",
    "finance:update",
    "payments:read", "payments:create", "ledger:export",
}
DEFAULT_MANAGER_PERMISSIONS = [
    "dashboard:read", "trips:read", "trips:create", "trips:update",
    "lrs:read", "lrs:create", "lrs:update", "ledger:export",
]
LEDGER_TEMPLATE_VERSION = "1"
MAX_LEDGER_BYTES = 2 * 1024 * 1024
LEDGER_COLUMNS = (
    "lr_id", "lr_ref", "trip_id", "site_id", "receiver", "goods_type",
    "rent", "hamali", "paid_total", "payment_date", "payment_method",
    "transaction_reference",
)


class StrictModel(BaseModel):
    model_config = ConfigDict(extra="forbid")

    @field_validator("rent", "hamali", check_fields=False, mode="before")
    @classmethod
    def blank_charge_is_zero(cls, value):
        return Decimal("0.00") if value == "" else value


class SiteCreate(StrictModel):
    name: str = Field(min_length=2, max_length=100)
    code: str = Field(min_length=2, max_length=8)
    location: str = Field(default="", max_length=250)
    city: str = Field(default="", max_length=100)
    timezone: str = Field(default="Asia/Kolkata", max_length=80)
    config: dict = Field(default_factory=dict)


class SiteUpdate(StrictModel):
    name: Optional[str] = Field(default=None, min_length=2, max_length=100)
    location: Optional[str] = Field(default=None, max_length=250)
    city: Optional[str] = Field(default=None, max_length=100)
    timezone: Optional[str] = Field(default=None, max_length=80)
    config: Optional[dict] = None
    status: Optional[str] = None


class ReceiptControls(StrictModel):
    hindi_conversion_enabled: bool = True
    receipt_language: Literal["hindi", "english"] = "hindi"
    sender_address_enabled: bool = True
    receiver_address_enabled: bool = True
    receipt_fee: Decimal = Field(default=Decimal("2.00"), ge=0, decimal_places=2)
    overdue_after_days: int = Field(default=30, ge=1, le=3650)


class ManagerAssignment(StrictModel):
    name: str = Field(min_length=2, max_length=100)
    username: str = Field(min_length=3, max_length=60)
    password: Optional[str] = Field(default=None, min_length=1)
    team_member_id: Optional[str] = Field(default=None, min_length=1, max_length=100)
    permissions: list[str] = Field(default_factory=lambda: list(DEFAULT_MANAGER_PERMISSIONS))


class ManagerAccess(StrictModel):
    site_id: str
    permissions: list[str]


class ManagerPasswordReset(StrictModel):
    password: str = Field(min_length=1)


class TripCreate(StrictModel):
    operating_date: Optional[str] = None
    truck_no: str = Field(default="", max_length=32)
    driver_name: str = Field(default="", max_length=100)
    vehicle_id: Optional[str] = None
    driver_id: Optional[str] = None


class TripUpdate(StrictModel):
    truck_no: Optional[str] = Field(default=None, min_length=1, max_length=32)
    driver_name: Optional[str] = Field(default=None, min_length=1, max_length=100)
    vehicle_id: Optional[str] = None
    driver_id: Optional[str] = None


class SiteTripExpenseCreate(StrictModel):
    description: str = Field(min_length=1, max_length=500)
    category: str = Field(min_length=1, max_length=80)
    amount: Decimal
    date: str
    payment_method: str = Field(min_length=1, max_length=30)
    payee: Optional[str] = Field(default=None, max_length=160)
    idempotency_key: str = Field(min_length=8, max_length=120)

    @model_validator(mode="before")
    @classmethod
    def accept_legacy_expense_names(cls, value: Any) -> Any:
        if not isinstance(value, dict):
            return value
        values = value.copy()
        for field, aliases in (
            ("description", ("remarks",)),
            ("payment_method", ("mode", "method")),
            ("payee", ("vendor",)),
        ):
            if field not in values:
                alias = next((name for name in aliases if name in values), None)
                if alias:
                    values[field] = values[alias]
            for alias in aliases:
                values.pop(alias, None)
        return values


class ReopenTrip(StrictModel):
    reason: str = Field(min_length=5, max_length=500)


class VoidReceipt(StrictModel):
    reason: str = Field(min_length=5, max_length=500)


class RestoreReceipt(StrictModel):
    reason: str = Field(min_length=5, max_length=500)
    idempotency_key: str = Field(min_length=8, max_length=120)


class ContainerLine(StrictModel):
    type: str = Field(min_length=1, max_length=80)
    quantity: int = Field(gt=0, le=1000000)
    type_hindi: str = Field(default="", max_length=120)
    description: str = Field(default="", max_length=500)
    description_hindi: str = Field(default="", max_length=500)
    rent: Optional[Decimal] = None
    hamali: Optional[Decimal] = None

    @field_validator("rent", "hamali", mode="before")
    @classmethod
    def blank_amount_defaults_to_zero(cls, value):
        return Decimal("0.00") if value == "" else value


class GoodsRow(StrictModel):
    type: str = Field(min_length=1, max_length=80)
    quantity: int = Field(gt=0, le=1000000)
    type_hindi: str = Field(default="", max_length=120)
    description: str = Field(default="", max_length=500)
    description_hindi: str = Field(default="", max_length=500)
    rent: Optional[Decimal] = None
    hamali: Optional[Decimal] = None

    @field_validator("rent", "hamali", mode="before")
    @classmethod
    def blank_amount_defaults_to_zero(cls, value):
        return Decimal("0.00") if value == "" else value


class LRCreate(StrictModel):
    sender_name: str = Field(default="", max_length=160)
    sender_phone: Optional[str] = Field(default=None, max_length=40)
    receiver_name: str = Field(min_length=1, max_length=160)
    receiver_phone: Optional[str] = Field(default=None, max_length=40)
    receiver_identifier: str = Field(default="", max_length=120)
    receiver_match_action: Optional[str] = None
    receiver_identity_id: Optional[str] = None
    sender_address: str = Field(default="", max_length=500)
    receiver_address: str = Field(default="", max_length=500)
    sender_address_hindi: str = Field(default="", max_length=500)
    receiver_address_hindi: str = Field(default="", max_length=500)
    city: str = Field(default="", max_length=100)
    sender_name_hindi: str = Field(default="", max_length=160)
    receiver_name_hindi: str = Field(default="", max_length=160)
    receipt_date: Optional[str] = None
    goods_type: Optional[str] = Field(default=None, min_length=1, max_length=80)
    goods_type_hindi: str = Field(default="", max_length=120)
    containers: Optional[list[ContainerLine]] = Field(default=None, min_length=1, max_length=30)
    goods_rows: Optional[list[GoodsRow]] = Field(default=None, min_length=1, max_length=30)
    rent: Optional[Decimal] = Decimal("0.00")
    hamali: Optional[Decimal] = Decimal("0.00")
    idempotency_key: Optional[str] = Field(default=None, min_length=8, max_length=120)

    @model_validator(mode="before")
    @classmethod
    def accept_goods_rows(cls, value: Any) -> Any:
        if not isinstance(value, dict):
            return value
        values = value.copy()
        goods_rows = values.pop("goods_rows", None)
        if goods_rows is not None:
            rows = [row for row in goods_rows if isinstance(row, dict) and str(row.get("type", "")).strip()]
            if rows:
                values["goods_type"] = rows[0]["type"]
                values["goods_type_hindi"] = rows[0].get("type_hindi", "")
                values["containers"] = rows
        return values

    @model_validator(mode="after")
    def require_goods(self):
        if not self.goods_type or not self.containers:
            raise ValueError("Enter at least one goods type and quantity")
        return self


class LRUpdate(StrictModel):
    sender_name: Optional[str] = Field(default=None, max_length=160)
    sender_phone: Optional[str] = Field(default=None, max_length=40)
    receiver_name: Optional[str] = Field(default=None, min_length=1, max_length=160)
    receiver_phone: Optional[str] = Field(default=None, max_length=40)
    receiver_identifier: Optional[str] = Field(default=None, max_length=120)
    sender_address: Optional[str] = Field(default=None, max_length=500)
    receiver_address: Optional[str] = Field(default=None, max_length=500)
    sender_address_hindi: Optional[str] = Field(default=None, max_length=500)
    receiver_address_hindi: Optional[str] = Field(default=None, max_length=500)
    city: Optional[str] = Field(default=None, max_length=100)
    sender_name_hindi: Optional[str] = Field(default=None, max_length=160)
    receiver_name_hindi: Optional[str] = Field(default=None, max_length=160)
    receipt_date: Optional[str] = None
    receiver_match_action: Optional[str] = None
    receiver_identity_id: Optional[str] = None
    goods_type: Optional[str] = Field(default=None, min_length=1, max_length=80)
    goods_type_hindi: Optional[str] = Field(default=None, max_length=120)
    containers: Optional[list[ContainerLine]] = Field(default=None, min_length=1, max_length=30)
    goods_rows: Optional[list[GoodsRow]] = Field(default=None, min_length=1, max_length=30)
    rent: Optional[Decimal] = None
    hamali: Optional[Decimal] = None
    idempotency_key: Optional[str] = Field(default=None, min_length=8, max_length=120)
    expected_updated_at: Optional[str] = None
    owner_correction_reason: Optional[str] = Field(default=None, max_length=500)

    @model_validator(mode="before")
    @classmethod
    def accept_goods_rows(cls, value: Any) -> Any:
        if not isinstance(value, dict):
            return value
        values = value.copy()
        goods_rows = values.pop("goods_rows", None)
        if goods_rows is not None:
            rows = [row for row in goods_rows if isinstance(row, dict) and str(row.get("type", "")).strip()]
            if rows:
                values["goods_type"] = rows[0]["type"]
                values["goods_type_hindi"] = rows[0].get("type_hindi", "")
                values["containers"] = rows
            else:
                raise ValueError("A receipt must contain at least one goods row")
        return values


class BookingFollowupUpdate(StrictModel):
    promised_date: Optional[str] = None
    note: str = Field(default="", max_length=500)
    next_contact_date: Optional[str] = None
    outcome: str = Field(default="", max_length=500)


class BookingSettlementUpdate(StrictModel):
    received: bool
    idempotency_key: str = Field(min_length=8, max_length=120)
    reason: str = Field(default="", max_length=500)


class PaymentCreate(StrictModel):
    amount: Decimal
    date: str
    method: str = Field(min_length=1, max_length=30)
    reference: str = Field(default="", max_length=120)
    idempotency_key: str = Field(min_length=8, max_length=120)


class PaymentReverse(StrictModel):
    reason: str = Field(min_length=5, max_length=500)


class CategoryCreate(StrictModel):
    kind: str
    name: str = Field(min_length=1, max_length=80)


class LegacyMigration(StrictModel):
    site_id: str
    confirm: bool = False


class ImportRowEdit(StrictModel):
    row_number: int = Field(gt=0)
    lr_id: Optional[str] = None
    rent: Optional[str] = None
    hamali: Optional[str] = None
    paid_total: Optional[str] = None
    payment_date: Optional[str] = None
    payment_method: Optional[str] = None
    transaction_reference: Optional[str] = None
    skip: bool = False


class ImportPreviewEdit(StrictModel):
    rows: list[ImportRowEdit] = Field(max_length=5000)


class ImportCommit(StrictModel):
    confirm: bool = False
    acknowledged_missing_lr_ids: list[str] = Field(default_factory=list, max_length=5000)


def _tenant_id(user):
    value = user.get("tenant_id")
    if not value:
        raise HTTPException(403, "Business account required")
    return str(value)


def _owner(user):
    if user.get("role") != "owner":
        raise HTTPException(403, "Business owner access required")


def _valid_timezone(value):
    try:
        ZoneInfo(value)
    except (ZoneInfoNotFoundError, ValueError, TypeError):
        raise HTTPException(400, "Select a valid IANA time zone")
    return value


def _site_config_with_receipt_controls(config, existing=None):
    result = dict(existing or {})
    result.update(config or {})
    raw_controls = dict(result.get("receipt_controls") or {})
    controls = ReceiptControls.model_validate(raw_controls)
    result["receipt_controls"] = controls.model_dump(mode="json")
    return result


def _receipt_controls(site):
    config = site.get("config") or {}
    try:
        raw_controls = dict(config.get("receipt_controls") or {})
        return ReceiptControls.model_validate(raw_controls)
    except ValidationError as error:
        raise HTTPException(500, "Receipt controls are invalid; ask an administrator to review site settings") from error


def _normalize_label(value, field, maximum=160):
    if not isinstance(value, str):
        raise HTTPException(422, f"{field} must be text")
    label = " ".join(value.split())
    if not label or len(label) > maximum or any(ord(ch) < 32 for ch in label):
        raise HTTPException(422, f"{field} is empty or invalid")
    return label


def _optional_phone(value, field):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    phone = _normalize_label(value, field, 40)
    if not re.fullmatch(r"[0-9]{10}", phone):
        raise HTTPException(422, f"{field} must contain exactly 10 digits")
    return phone


def _normalize_code(value):
    code = value.strip().upper()
    if not re.fullmatch(r"[A-Z0-9]{2,8}", code):
        raise HTTPException(422, "Site code must contain 2–8 letters or numbers")
    return code


def _decimal(value, field, allow_none=False):
    if value is None or value == "":
        if allow_none:
            return None
        raise HTTPException(422, f"{field} is required")
    try:
        amount = value if isinstance(value, Decimal) else Decimal(str(value))
        if not amount.is_finite() or amount < 0:
            raise InvalidOperation
        exponent = amount.as_tuple().exponent
        if not isinstance(exponent, int) or exponent < -2:
            raise InvalidOperation
        return amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError, TypeError):
        raise HTTPException(422, f"{field} must be a non-negative amount with at most two decimals")


def _decimal_value(value):
    if value is None:
        return Decimal("0.00")
    if isinstance(value, Decimal128):
        return value.to_decimal()
    if isinstance(value, Decimal):
        return value
    return Decimal(str(value))


def _clean(value: Any) -> Any:
    if isinstance(value, Decimal128):
        return format(value.to_decimal(), ".2f")
    if isinstance(value, Decimal):
        return format(value, ".2f")
    if isinstance(value, dict):
        return {k: _clean(v) for k, v in value.items() if k != "_id"}
    if isinstance(value, list):
        return [_clean(v) for v in value]
    return value


@overload
def _doc(value: None) -> None: ...


@overload
def _doc(value: dict[str, Any]) -> dict[str, Any]: ...


def _doc(value: Optional[dict[str, Any]]) -> Optional[dict[str, Any]]:
    if value is None:
        return None
    data = _clean(value)
    data["id"] = _clean(value.get("_id"))
    return data


CREATOR_FIELDS = (
    "created_by", "created_by_id", "created_by_login", "created_by_name",
    "created_by_role", "created_at",
)
CHARGE_EDITOR_FIELDS = (
    "charge_updated_by_id", "charge_updated_by_login", "charge_updated_by_name",
    "charge_updated_by_role", "charge_updated_at",
)


def _creator_metadata(user: dict[str, Any]) -> dict[str, str]:
    login = str(user.get("username") or user.get("_id") or "")
    return {
        "created_by": login,
        "created_by_id": str(user.get("_id") or login),
        "created_by_login": login,
        "created_by_name": str(user.get("name") or ""),
        "created_by_role": str(user.get("role") or ""),
        "created_at": now_iso(),
    }


def _charge_editor_metadata(user: dict[str, Any]) -> dict[str, str]:
    login = str(user.get("username") or user.get("_id") or "")
    return {
        "charge_updated_by_id": str(user.get("_id") or login),
        "charge_updated_by_login": login,
        "charge_updated_by_name": str(user.get("name") or ""),
        "charge_updated_by_role": str(user.get("role") or ""),
        "charge_updated_at": now_iso(),
    }


def _charge_editor_marker(record: dict[str, Any]) -> str:
    role = record.get("charge_updated_by_role")
    if role in {"owner", "admin"}:
        return "A"
    if role == "site_manager":
        return "M"
    return ""


def _entry_by(record: dict[str, Any]) -> str:
    role = record.get("created_by_role")
    label = "Admin" if role in {"owner", "admin"} else "Manager" if role == "site_manager" else ""
    name = str(record.get("created_by_name") or "").strip()
    return f"{label} - {name}" if label and name else ""


def _hindi_digits(value: Any) -> str:
    return str(value).translate(str.maketrans("0123456789", "०१२३४५६७८९"))


def _ledger_hindi_text(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return text
    common_goods = {
        "auto": "ऑटो", "truck": "ट्रक", "cement": "सीमेंट",
        "gadi": "गाड़ी", "gaadi": "गाड़ी", "bhada": "भाड़ा", "bhaada": "भाड़ा",
        "gadi bhada": "गाड़ी भाड़ा", "gaadi bhada": "गाड़ी भाड़ा",
        "gadi bhaada": "गाड़ी भाड़ा", "gaadi bhaada": "गाड़ी भाड़ा",
        "vehicle rent": "गाड़ी भाड़ा", "vehicle rental": "गाड़ी भाड़ा",
        "truck rent": "ट्रक भाड़ा", "truck bhada": "ट्रक भाड़ा",
        "lorry rent": "ट्रक भाड़ा", "lorry bhada": "ट्रक भाड़ा",
        "hamali": "हमाली", "receipt charge": "रसीद शुल्क",
        "receipt charges": "रसीद शुल्क",
        "food grains": "अनाज", "rice": "चावल", "wheat": "गेहूँ",
        "flour": "आटा", "sugar": "चीनी", "salt": "नमक", "steel": "लोहा",
        "iron": "लोहा", "bricks": "ईंट", "brick": "ईंट", "sand": "रेत",
        "stone": "पत्थर", "wood": "लकड़ी", "fertilizer": "खाद",
        "clothes": "कपड़े", "cloth": "कपड़ा", "oil": "तेल",
        "vegetables": "सब्ज़ियाँ", "fruits": "फल", "general goods": "सामान",
        "building material": "निर्माण सामग्री", "car": "कार", "cars": "कारें",
        "hardware": "हार्डवेयर", "part": "पार्ट", "parts": "पार्ट्स",
        "car part": "कार का पुर्जा", "car parts": "कार के पुर्ज़े",
        "spare parts": "स्पेयर पार्ट्स", "bag": "बोरा", "bags": "बोरे",
        "box": "बॉक्स", "boxes": "बॉक्स", "drum": "ड्रम", "carton": "कार्टन",
        "parcel": "पार्सल", "grain": "अनाज", "grains": "अनाज",
        "machine": "मशीन", "machinery": "मशीनरी", "tools": "औज़ार",
        "pipe": "पाइप", "rod": "छड़", "plastic": "प्लास्टिक",
        "paper": "कागज़", "potato": "आलू", "mango": "आम", "onion": "प्याज़",
        "medicine": "दवाई", "water": "पानी",
    }
    if text.lower() in common_goods:
        return common_goods[text.lower()]

    consonants = (
        ("ksh", "क्ष"), ("chh", "छ"), ("sh", "श"), ("ch", "च"), ("kh", "ख"),
        ("gh", "घ"), ("th", "थ"), ("dh", "ध"), ("ph", "फ"), ("bh", "भ"),
        ("jh", "झ"), ("ng", "ङ"), ("ny", "ञ"), ("q", "क़"), ("k", "क"),
        ("g", "ग"), ("c", "क"), ("j", "ज"), ("t", "त"), ("d", "द"),
        ("n", "न"), ("p", "प"), ("b", "ब"), ("m", "म"), ("y", "य"),
        ("r", "र"), ("l", "ल"), ("v", "व"), ("w", "व"), ("s", "स"),
        ("h", "ह"), ("f", "फ़"), ("z", "ज़"),
    )
    vowels = (
        ("aa", "ा"), ("ee", "ी"), ("ii", "ी"), ("oo", "ू"), ("uu", "ू"),
        ("ai", "ै"), ("au", "ौ"), ("a", ""), ("i", "ि"), ("u", "ु"),
        ("e", "े"), ("o", "ो"),
    )
    independent_vowels = {
        "ा": "आ", "ी": "ई", "ू": "ऊ", "ै": "ऐ", "ौ": "औ",
        "ि": "इ", "ु": "उ", "े": "ए", "ो": "ओ",
    }

    def translate_word(word: str) -> str:
        lower = word.lower()
        if lower in common_goods:
            return common_goods[lower]
        if any("\u0900" <= character <= "\u097f" for character in word):
            return word
        result = []
        index = 0
        pending_consonant = False
        while index < len(lower):
            vowel = next((item for item in vowels if lower.startswith(item[0], index)), None)
            if vowel:
                result.append(vowel[1] if pending_consonant else independent_vowels.get(vowel[1], "अ"))
                pending_consonant = False
                index += len(vowel[0])
                continue
            consonant = next(
                (item for item in consonants if lower.startswith(item[0], index)), None,
            )
            if consonant:
                if pending_consonant:
                    result.append("्")
                result.append(consonant[1])
                pending_consonant = True
                index += len(consonant[0])
                continue
            if pending_consonant:
                result.append("्")
            result.append(word[index])
            pending_consonant = False
            index += 1
        return "".join(result)

    return "".join(
        part if not part or re.fullmatch(r"\s+|[-/]", part) else translate_word(part)
        for part in re.split(r"(\s+|[-/])", text)
    )


def _creator_user(user: dict[str, Any], record: dict[str, Any]) -> dict[str, Any]:
    return {
        **user,
        "_id": record.get("created_by_id", user.get("_id") or user["username"]),
        "username": record.get("created_by_login", record.get("created_by", user["username"])),
        "name": record.get("created_by_name", user.get("name", "")),
        "role": record.get("created_by_role", user.get("role", "")),
    }


def _doc_for_user(value: Optional[dict[str, Any]], user: dict[str, Any]) -> Optional[dict[str, Any]]:
    data = _doc(value)
    if data is not None and user.get("role") == "site_manager":
        for key in (*CREATOR_FIELDS, *CHARGE_EDITOR_FIELDS):
            data.pop(key, None)
    return data


def _scoped(business_id: str, site_id: str, **filters: Any) -> dict[str, Any]:
    return {"business_id": business_id, "site_id": site_id, **filters}


async def initialize_site_storage():
    """Create indexes for the current tenant DB; called lazily for new tenants."""
    name = current_db_name()
    if name in _indexed_databases:
        return
    async with _index_lock:
        if name in _indexed_databases:
            return
        await db.sites.create_index(
            [("business_id", 1), ("code_normalized", 1)], unique=True,
            name="site_business_code_unique",
        )
        await db.site_trips.create_index(
            [("business_id", 1), ("site_id", 1), ("operating_date", 1), ("sequence", 1)],
            unique=True, name="site_trip_daily_sequence_unique",
        )
        await db.site_trips.create_index(
            [("business_id", 1), ("site_id", 1), ("trip_ref", 1)],
            unique=True, name="site_trip_ref_unique",
        )
        await db.site_trips.create_index(
            [("business_id", 1), ("site_id", 1), ("status", 1), ("operating_date", -1)],
            name="site_trip_dashboard",
        )
        await db.site_lrs.create_index(
            [("business_id", 1), ("site_id", 1), ("operating_date", 1), ("sequence", 1)],
            unique=True, name="site_lr_daily_sequence_unique",
        )
        await db.site_lrs.create_index(
            [("business_id", 1), ("site_id", 1), ("trip_id", 1), ("operating_date", 1)],
            name="site_lr_trip_date",
        )
        await db.site_lrs.create_index(
            [("business_id", 1), ("site_id", 1), ("idempotency_key", 1)],
            unique=True, partialFilterExpression={"idempotency_key": {"$exists": True}},
            name="site_lr_idempotency_unique",
        )
        await db.site_lrs.create_index(
            [("business_id", 1), ("site_id", 1), ("receiver_normalized", 1), ("operating_date", 1)],
            name="site_lr_receiver_date",
        )
        await db.site_categories.create_index(
            [("business_id", 1), ("site_id", 1), ("kind", 1), ("normalized_name", 1)],
            unique=True, name="site_category_unique",
        )
        await db.site_receivers.create_index(
            [("business_id", 1), ("site_id", 1), ("operating_date", 1), ("normalized_name", 1)],
            name="site_receiver_daily_lookup",
        )
        await db.parties.create_index(
            [("business_id", 1), ("site_id", 1), ("site_receiver_id", 1)],
            unique=True, partialFilterExpression={"site_receiver_id": {"$exists": True}},
            name="site_receiver_party_unique",
        )
        await db.site_payments.create_index(
            [("business_id", 1), ("site_id", 1), ("idempotency_key", 1)],
            unique=True, name="site_payment_idempotency_unique",
        )
        await db.expenses.create_index(
            [("business_id", 1), ("site_id", 1), ("trip_id", 1), ("idempotency_key", 1)],
            unique=True, partialFilterExpression={"idempotency_key": {"$exists": True}},
            name="site_trip_expense_idempotency_unique",
        )
        await db.site_payments.create_index(
            [("business_id", 1), ("site_id", 1), ("trip_id", 1), ("lr_id", 1)],
            name="site_payment_trip_lr",
        )
        await db.site_payments.create_index(
            [("business_id", 1), ("site_id", 1), ("operating_date", 1)],
            name="site_payment_dashboard_date",
        )
        for collection in (db.ledger, db.cashbook):
            await collection.create_index(
                [("business_id", 1), ("site_id", 1), ("ref_type", 1), ("ref_id", 1)],
                unique=True,
                partialFilterExpression={"$and": [
                    {"site_id": {"$exists": True}},
                    {"$or": [
                        {"ref_type": "site_lr_financial"},
                        {"ref_type": "site_lr_payment"},
                    ]},
                ]},
                name="site_financial_posting_unique",
            )
        for collection in (db.ledger, db.cashbook):
            await collection.create_index(
                [("business_id", 1), ("site_id", 1), ("ref_type", 1), ("ref_id", 1)],
                unique=True,
                partialFilterExpression={"$and": [
                    {"site_id": {"$exists": True}},
                    {"ref_type": "site_trip_expense"},
                ]},
                name="site_trip_expense_posting_unique",
            )
        existing_reversal_index = (await db.site_payments.index_information()).get(
            "site_payment_reversal_unique",
        )
        if existing_reversal_index and (
            existing_reversal_index.get("unique") is not True
            or existing_reversal_index.get("partialFilterExpression")
            != {"kind": "reversal", "reversal_of": {"$exists": True}}
        ):
            await db.site_payments.drop_index("site_payment_reversal_unique")
        await db.site_payments.create_index(
            [("business_id", 1), ("site_id", 1), ("reversal_of", 1)],
            unique=True,
            partialFilterExpression={"kind": "reversal", "reversal_of": {"$exists": True}},
            name="site_payment_reversal_unique",
        )
        await db.site_audit_events.create_index(
            [("business_id", 1), ("site_id", 1), ("created_at", -1)],
            name="site_audit_timeline",
        )
        await db.site_ledger_imports.create_index(
            [("business_id", 1), ("site_id", 1), ("trip_id", 1), ("file_hash", 1)],
            unique=True, name="site_ledger_import_idempotency",
        )
        await db.site_ledger_imports.create_index(
            [("business_id", 1), ("site_id", 1), ("status", 1), ("created_at", -1)],
            name="site_ledger_import_status",
        )
        _indexed_databases.add(name)


async def _site_for_user(site_id, user, permission=None, include_inactive_owner=True):
    await initialize_site_storage()
    business_id = _tenant_id(user)
    site = await db.sites.find_one({"_id": site_id, "business_id": business_id})
    if not site:
        raise HTTPException(404, "Site not found")
    if user.get("role") == "owner":
        if not include_inactive_owner and site.get("status") != "Active":
            raise HTTPException(409, "This site is inactive")
        return site
    if user.get("role") != "site_manager":
        raise HTTPException(403, "Business site access required")
    if site.get("status") != "Active":
        raise HTTPException(403, "This site is inactive")
    grants = user.get("site_permissions") or {}
    permitted = set(grants.get(site_id) or [])
    if site_id not in (user.get("site_ids") or []) or (permission and permission not in permitted):
        raise HTTPException(403, "You do not have permission for this site action")
    return site


async def _audit(site, user, action, target_type, target_id, old=None, new=None, reason=None,
                 event_id=None):
    await db.site_audit_events.update_one(
        {"_id": event_id or new_id()},
        {"$setOnInsert": {
        "business_id": str(site["business_id"]), "site_id": site["_id"],
        "actor_id": str(user.get("_id") or user["username"]),
        "actor_username": user["username"], "actor_name": user.get("name", ""),
        "actor_role": user.get("role", ""),
        "action": action, "target_type": target_type, "target_id": target_id,
        "old_values": old or {}, "new_values": new or {}, "reason": reason or "",
        "created_at": now_iso(),
        }},
        upsert=True,
    )


def _audit_rows_for_user(
    rows: list[dict[str, Any]], user: dict[str, Any], site_id: str,
) -> list[dict[str, Any]]:
    finance_read = (
        user.get("role") != "site_manager"
        or "finance:read" in (user.get("site_permissions") or {}).get(site_id, [])
    )
    output = []
    for row in rows:
        item = _doc(row)
        if not finance_read:
            if item.get("target_type") == "payment" or str(item.get("action", "")).startswith("ledger."):
                continue
            for field in ("old_values", "new_values"):
                values = item.get(field) or {}
                item[field] = {
                    key: value for key, value in values.items()
                    if key not in {
                        "rent", "hamali", "paid_total", "outstanding", "amount",
                        "recorded_bhada", "reconciled_bhada", "collected_bhada",
                        "outstanding_bhada", "recorded_hamali",
                    }
                }
        output.append(item)
    return output


def _now_at_site(site):
    zone = ZoneInfo(site.get("timezone") or "Asia/Kolkata")
    current = datetime.now(zone)
    return current.date().isoformat(), current


def _validate_iso_date(value):
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError):
        raise HTTPException(422, "Date must be in YYYY-MM-DD format")


async def _manager_user(username: str, business_id: str) -> dict[str, Any]:
    user = await platform_db.users.find_one({"_id": username, "tenant_id": business_id})
    if not user or user.get("role") != "site_manager":
        raise HTTPException(404, "Site manager not found")
    return user


async def _replace_manager_access(username, business_id, site_id, permissions):
    target = await platform_db.users.find_one({"_id": username, "tenant_id": business_id})
    if not target:
        raise HTTPException(404, "Site manager not found")
    site_ids = set(target.get("site_ids") or [])
    current = dict(target.get("site_permissions") or {})
    if permissions:
        site_ids.add(site_id)
        current[site_id] = sorted(set(permissions))
        active = True
    else:
        site_ids.discard(site_id)
        current.pop(site_id, None)
        active = bool(site_ids)
    await platform_db.users.update_one(
        {"_id": username, "tenant_id": business_id},
        {"$set": {"site_ids": sorted(site_ids), "site_permissions": current, "active": active}},
    )


@router.get("/sites")
async def list_sites(u=Depends(site_user)):
    await initialize_site_storage()
    business_id = _tenant_id(u)
    query: dict[str, Any] = {"business_id": business_id}
    if u.get("role") == "site_manager":
        query["_id"] = {"$in": u.get("site_ids") or []}
    elif u.get("role") != "owner":
        raise HTTPException(403, "Business site access required")
    sites = await db.sites.find(query).sort([("is_default", -1), ("name", 1)]).to_list(500)
    return [_doc_for_user(s, u) for s in sites]


@router.post("/sites")
async def create_site(body: SiteCreate, u=Depends(site_user)):
    _owner(u)
    await initialize_site_storage()
    business_id = _tenant_id(u)
    code = _normalize_code(body.code)
    timezone_name = _valid_timezone(body.timezone)
    existing = await db.sites.find_one({"business_id": business_id, "code_normalized": code})
    if existing:
        raise HTTPException(409, "That short code is already used by another site in this business")
    first_site = await db.sites.count_documents({"business_id": business_id}) == 0
    site = {
        "_id": new_id(), "business_id": business_id,
        "name": _normalize_label(body.name, "Site name", 100), "code": code,
        "code_normalized": code, "location": " ".join(body.location.split()),
        "city": " ".join(body.city.split()), "timezone": timezone_name,
        "status": "Active", "config": _site_config_with_receipt_controls(body.config),
        "manager_username": None,
        "is_default": first_site, **_creator_metadata(u),
        "updated_at": now_iso(),
    }
    try:
        await db.sites.insert_one(site)
    except DuplicateKeyError:
        raise HTTPException(409, "That short code is already used by another site in this business")
    await _audit(site, u, "site.created", "site", site["_id"], new=_doc(site))
    return _doc(site)


@router.put("/sites/{site_id}")
async def update_site(site_id: str, body: SiteUpdate, u=Depends(site_user)):
    _owner(u)
    site = await _site_for_user(site_id, u)
    changes = body.model_dump(exclude_unset=True)
    if "status" in changes and changes["status"] not in ("Active", "Inactive"):
        raise HTTPException(422, "Site status must be Active or Inactive")
    if "timezone" in changes:
        changes["timezone"] = _valid_timezone(changes["timezone"])
    if "config" in changes:
        changes["config"] = _site_config_with_receipt_controls(
            changes["config"], site.get("config"),
        )
    for key in ("name", "location", "city"):
        if key in changes and changes[key] is not None:
            changes[key] = _normalize_label(changes[key], key, 250 if key == "location" else 100) if changes[key] else ""
    changes.pop("code", None)
    changes.pop("business_id", None)
    changes["updated_at"] = now_iso()
    await db.sites.update_one({"_id": site_id, "business_id": str(u["tenant_id"])}, {"$set": changes})
    updated = await db.sites.find_one({"_id": site_id, "business_id": str(u["tenant_id"])})
    if not updated:
        raise HTTPException(500, "Site could not be loaded after update")
    await _audit(site, u, "site.updated", "site", site_id,
                 old={k: site.get(k) for k in changes if k in site},
                 new={k: updated.get(k) for k in changes if k in updated})
    return _doc(updated)


@router.post("/sites/{site_id}/manager")
async def assign_site_manager(site_id: str, body: ManagerAssignment, u=Depends(site_user)):
    _owner(u)
    site = await _site_for_user(site_id, u)
    business_id = _tenant_id(u)
    actor_login = str(u.get("username") or u.get("_id") or "")
    username = body.username.strip().lower()
    if not re.fullmatch(r"[a-z0-9._-]{3,60}", username):
        raise HTTPException(422, "Username may contain lowercase letters, numbers, dot, underscore, and hyphen")
    permissions = sorted(set(body.permissions))
    if not permissions or set(permissions) - MANAGER_PERMISSIONS:
        raise HTTPException(422, "One or more requested manager permissions are not allowed")
    team_member = None
    if body.team_member_id:
        team_member = await db.team.find_one({
            "_id": body.team_member_id, "archived": {"$ne": True},
        })
        if not team_member:
            raise HTTPException(422, "Selected team member is unavailable")
        linked_manager = await platform_db.users.find_one({
            "tenant_id": business_id, "role": "site_manager",
            "team_member_id": body.team_member_id, "active": {"$ne": False},
        })
        if linked_manager and linked_manager.get("_id") != username:
            raise HTTPException(409, "That team member is already linked to another manager login")
    old_username = site.get("manager_username")
    existing = await platform_db.users.find_one({"_id": username})
    if existing and (existing.get("tenant_id") != business_id or existing.get("role") != "site_manager"):
        raise HTTPException(409, "That username already belongs to another account")
    password = body.password or ""
    if not existing and not password:
        raise HTTPException(422, "A password is required for a new manager")
    if existing:
        update_fields = {
            "name": _normalize_label(body.name, "Manager name", 100),
            "active": True,
            "updated_by_id": str(u.get("_id") or actor_login),
            "updated_by_login": actor_login,
            "updated_by_name": u.get("name", ""),
            "updated_by_role": u.get("role", ""),
            "updated_at": now_iso(),
        }
        if body.team_member_id:
            update_fields["team_member_id"] = body.team_member_id
        if password:
            update_fields["password"] = hash_manager_pw(password)
        await platform_db.users.update_one(
            {"_id": username, "tenant_id": business_id},
            {"$set": update_fields},
        )
    else:
        await platform_db.users.insert_one({
            "_id": username, "name": _normalize_label(body.name, "Manager name", 100),
            "password": hash_manager_pw(password), "role": "site_manager",
            "tenant_id": business_id, "active": True, "site_ids": [],
            "site_permissions": {}, **_creator_metadata(u),
            **({"team_member_id": body.team_member_id} if body.team_member_id else {}),
        })
    await _replace_manager_access(username, business_id, site_id, permissions)
    site_update = {"$set": {"manager_username": username, "updated_at": now_iso()}}
    if body.team_member_id:
        site_update["$set"]["manager_team_member_id"] = body.team_member_id
    else:
        site_update["$unset"] = {"manager_team_member_id": ""}
    await db.sites.update_one({"_id": site_id, "business_id": business_id}, site_update)
    if old_username and old_username != username:
        await _replace_manager_access(old_username, business_id, site_id, [])
    await _audit(site, u, "site.manager_assigned", "manager", username,
                 old={"username": old_username},
                 new={"username": username, "permissions": permissions})
    if existing and password:
        await _audit(site, u, "manager.password_reset", "manager", username)
    return {"site_id": site_id, "username": username, "name": body.name, "permissions": permissions}


@router.get("/sites/managers")
async def list_site_managers(u=Depends(site_user)):
    _owner(u)
    business_id = _tenant_id(u)
    managers = await platform_db.users.find(
        {"tenant_id": business_id, "role": "site_manager"}
    ).sort("name", 1).to_list(500)
    return [{
        "username": m["_id"], "name": m.get("name", ""), "active": m.get("active") is not False,
        "team_member_id": m.get("team_member_id"),
        "site_ids": m.get("site_ids") or [], "site_permissions": m.get("site_permissions") or {},
    } for m in managers]


@router.put("/sites/managers/{username}/access")
async def update_manager_access(username: str, body: ManagerAccess, u=Depends(site_user)):
    _owner(u)
    business_id = _tenant_id(u)
    site = await _site_for_user(body.site_id, u)
    permissions = sorted(set(body.permissions))
    if set(permissions) - MANAGER_PERMISSIONS:
        raise HTTPException(422, "One or more requested manager permissions are not allowed")
    manager = await _manager_user(username, business_id)
    old = (manager.get("site_permissions") or {}).get(body.site_id, [])
    await _replace_manager_access(username, business_id, body.site_id, permissions)
    action = (
        "site.manager_access_updated" if site.get("manager_username") == username
        else "site.cross_access_updated"
    )
    await _audit(site, u, action, "manager", username,
                 old={"permissions": old}, new={"permissions": permissions})
    return {"username": username, "site_id": body.site_id, "permissions": permissions}


@router.post("/sites/managers/{username}/reset-password")
async def reset_manager_password(
    username: str, payload: ManagerPasswordReset, u=Depends(site_user),
):
    _owner(u)
    business_id = _tenant_id(u)
    await _manager_user(username, business_id)
    password = payload.password
    await platform_db.users.update_one(
        {"_id": username, "tenant_id": business_id, "role": "site_manager"},
        {"$set": {"password": hash_manager_pw(password)}},
    )
    site = await db.sites.find_one({"business_id": business_id, "manager_username": username})
    if site:
        await _audit(site, u, "manager.password_reset", "manager", username)
    return {"ok": True}


@router.delete("/sites/managers/{username}")
async def deactivate_manager(username: str, u=Depends(site_user)):
    _owner(u)
    business_id = _tenant_id(u)
    manager = await _manager_user(username, business_id)
    for site_id in manager.get("site_ids") or []:
        site = await db.sites.find_one({"_id": site_id, "business_id": business_id})
        if site:
            if site.get("manager_username") == username:
                await db.sites.update_one(
                    {"_id": site_id, "business_id": business_id},
                    {"$unset": {"manager_username": ""}, "$set": {"updated_at": now_iso()}},
                )
            await _audit(site, u, "site.manager_deactivated", "manager", username,
                         old={"username": username})
    await platform_db.users.update_one(
        {"_id": username, "tenant_id": business_id, "role": "site_manager"},
        {"$set": {"active": False, "site_ids": [], "site_permissions": {}}},
    )
    return {"ok": True}


async def _next_sequence(site_id, operating_date, kind):
    key = f"{site_id}:{kind}:{operating_date}"
    try:
        counter = await db.site_counters.find_one_and_update(
            {"_id": key}, {"$inc": {"sequence": 1}},
            upsert=True, return_document=ReturnDocument.AFTER,
        )
    except DuplicateKeyError:
        counter = await db.site_counters.find_one_and_update(
            {"_id": key}, {"$inc": {"sequence": 1}},
            return_document=ReturnDocument.AFTER,
        )
    if not counter:
        raise HTTPException(503, "Could not allocate a unique daily sequence; retry the request")
    return counter["sequence"]


async def _trip_master_values(
    vehicle_id, driver_id, allow_archived_vehicle_id=None, allow_archived_driver_id=None,
):
    values = {}
    if vehicle_id:
        vehicle_query = {"_id": vehicle_id}
        if vehicle_id != allow_archived_vehicle_id:
            vehicle_query["archived"] = {"$ne": True}
        vehicle = await db.vehicles.find_one(vehicle_query)
        if not vehicle:
            raise HTTPException(404, "Selected vehicle was not found or is archived")
        values["vehicle_id"] = vehicle["_id"]
        values["truck_no"] = _normalize_label(vehicle.get("vehicle_no", ""), "Truck number", 32)
    if driver_id:
        driver_query = {"_id": driver_id}
        if driver_id != allow_archived_driver_id:
            driver_query["archived"] = {"$ne": True}
        driver = await db.drivers.find_one(driver_query)
        if not driver:
            raise HTTPException(404, "Selected driver was not found or is archived")
        values["driver_id"] = driver["_id"]
        values["driver_name"] = _normalize_label(driver.get("name", ""), "Driver name", 100)
    return values


async def _trip(site: dict[str, Any], trip_id: str) -> dict[str, Any]:
    trip = await db.site_trips.find_one(_scoped(str(site["business_id"]), site["_id"], _id=trip_id))
    if not trip:
        raise HTTPException(404, "Trip not found")
    return trip


def _ensure_trip_ledger_open(trip: dict[str, Any]) -> None:
    if trip.get("ledger_completed_at") or trip.get("ledger_completion_started_at"):
        raise HTTPException(409, "This trip ledger is being completed or is already locked")


def _can_read_site_finances(user: dict[str, Any], site_id: str) -> bool:
    return user.get("role") != "site_manager" or bool(
        {"finance:read", "finance:update"}.intersection(
            (user.get("site_permissions") or {}).get(site_id, [])
        )
    )


def _can_view_site_lr_charges(user: dict[str, Any], site_id: str) -> bool:
    if user.get("role") != "site_manager":
        return True
    permissions = (user.get("site_permissions") or {}).get(site_id, [])
    return bool({"finance:read", "finance:update", "lrs:create", "lrs:update"}.intersection(permissions))


def _site_expense_doc(row: dict[str, Any], user: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    result = _doc_for_user(row, user or {})
    for key in ("idempotency_key", "request_fingerprint", "posting_started_at", "posting_error"):
        result.pop(key, None)
    return result


async def _trip_expense_rows(
    site: dict[str, Any], trip_id: str, user: Optional[dict[str, Any]] = None,
) -> tuple[list[dict[str, Any]], Decimal, int]:
    query = _scoped(str(site["business_id"]), site["_id"], trip_id=trip_id, cancelled=False)
    rows = await db.expenses.find(query).sort([("date", -1), ("created_at", -1)]).to_list(500)
    aggregates = await db.expenses.aggregate([
        {"$match": query},
        {"$group": {"_id": None, "total": {"$sum": "$amount"}, "count": {"$sum": 1}}},
    ]).to_list(1)
    summary = aggregates[0] if aggregates else {}
    total = _decimal_value(summary.get("total"))
    return [_site_expense_doc(row, user) for row in rows], total, summary.get("count", 0)


async def _site_trip_expense_totals(query: dict[str, Any]) -> dict[str, dict[str, Any]]:
    result = {}
    async for row in db.expenses.aggregate([
        {"$match": query},
        {"$group": {"_id": "$site_id", "total": {"$sum": "$amount"}, "count": {"$sum": 1}}},
    ]):
        result[row["_id"]] = {
            "total": _decimal_value(row.get("total")),
            "count": row.get("count", 0),
        }
    return result


def _receivable_group_stages(key_expression: Any) -> list[dict[str, Any]]:
    return [
        {"$group": {
            "_id": key_expression,
            "lr_count": {"$sum": 1},
            "parcels": {"$sum": {"$ifNull": ["$total_quantity", 0]}},
            "reconciled_rent": {"$sum": "$_collectible_rent"},
            "posted_payments": {"$sum": "$_collectible_paid"},
            "collectible_outstanding": {"$sum": "$_collectible_outstanding"},
            "unpaid_lrs": {"$sum": "$_collectible_unpaid"},
            "unreconciled_bhada": {"$sum": "$_unreconciled_rent"},
            "unreconciled_lrs": {"$sum": "$_unreconciled_count"},
            "unpriced_lrs": {"$sum": "$_unpriced_count"},
        }},
    ]


def _format_receivable_row(row: dict[str, Any], label: Optional[str] = None) -> dict[str, Any]:
    result = {
        "lr_count": row.get("lr_count", 0),
        "parcels": row.get("parcels", 0),
        "reconciled_rent": format(_decimal_value(row.get("reconciled_rent")), ".2f"),
        "posted_payments": format(_decimal_value(row.get("posted_payments")), ".2f"),
        "collectible_outstanding": format(
            _decimal_value(row.get("collectible_outstanding")), ".2f",
        ),
        "unpaid_lrs": row.get("unpaid_lrs", 0),
        "unreconciled_bhada": format(_decimal_value(row.get("unreconciled_bhada")), ".2f"),
        "unreconciled_lrs": row.get("unreconciled_lrs", 0),
        "unpriced_lrs": row.get("unpriced_lrs", 0),
    }
    if label is not None:
        result["label"] = label
    return result


async def _receivables_breakdown(lr_query: dict[str, Any]) -> dict[str, Any]:
    priced = {"$ne": ["$rent", None]}
    reconciled = {"$eq": ["$reconciled", True]}
    collectable = {"$and": [reconciled, priced]}
    receiver_key = {
        "site_id": "$site_id",
        "operating_date": "$operating_date",
        "label": {"$ifNull": ["$receiver_label", "$receiver_name"]},
    }
    pipeline = [
        {"$match": lr_query},
        {"$lookup": {
            "from": "site_payments",
            "let": {
                "lr_id": "$_id", "business_id": "$business_id",
                "site_id": "$site_id", "trip_id": "$trip_id",
            },
            "pipeline": [
                {"$match": {"$expr": {"$and": [
                    {"$eq": ["$lr_id", "$$lr_id"]},
                    {"$eq": ["$business_id", "$$business_id"]},
                    {"$eq": ["$site_id", "$$site_id"]},
                    {"$eq": ["$trip_id", "$$trip_id"]},
                    {"$eq": ["$posting_status", "posted"]},
                    {"$in": ["$kind", ["payment", "reversal"]]},
                ]}}},
                {"$group": {
                    "_id": None,
                    "total": {"$sum": {"$cond": [
                        {"$eq": ["$kind", "reversal"]},
                        {"$multiply": ["$amount", -1]}, "$amount",
                    ]}},
                }},
            ],
            "as": "_posted_payment_rows",
        }},
        {"$set": {
            "_posted_lr_amount": {"$ifNull": [
                {"$arrayElemAt": ["$_posted_payment_rows.total", 0]}, 0,
            ]},
            "_collectible_rent": {"$cond": [
                collectable,
                {"$add": [
                    "$rent", {"$ifNull": ["$receipt_fee", Decimal128(Decimal("2.00"))]},
                ]},
                0,
            ]},
            "_collectible_paid": {"$cond": [collectable, "$_posted_lr_amount", 0]},
            "_collectible_outstanding": {"$cond": [
                collectable,
                {"$max": [{"$subtract": [
                    {"$add": [
                        "$rent", {"$ifNull": ["$receipt_fee", Decimal128(Decimal("2.00"))]},
                    ]},
                    "$_posted_lr_amount",
                ]}, 0]},
                0,
            ]},
            "_collectible_unpaid": {"$cond": [
                {"$and": [collectable, {"$gt": [
                    {"$subtract": [
                        {"$add": [
                            "$rent", {"$ifNull": ["$receipt_fee", Decimal128(Decimal("2.00"))]},
                        ]},
                        "$_posted_lr_amount",
                    ]}, 0,
                ]}]}, 1, 0,
            ]},
            "_unreconciled_rent": {"$cond": [
                {"$and": [{"$not": [reconciled]}, priced]}, "$rent", 0,
            ]},
            "_unreconciled_count": {"$cond": [{"$not": [reconciled]}, 1, 0]},
            "_unpriced_count": {"$cond": [{"$eq": ["$rent", None]}, 1, 0]},
        }},
        {"$facet": {
            "by_receiver": [
                *_receivable_group_stages(receiver_key),
                {"$sort": {
                    "collectible_outstanding": -1, "unreconciled_bhada": -1,
                    "unreconciled_lrs": -1, "unpriced_lrs": -1, "_id": 1,
                }},
                {"$limit": 200},
            ],
            "receiver_count": [
                {"$group": {"_id": receiver_key}},
                {"$count": "count"},
            ],
            "by_goods": [
                *_receivable_group_stages({"$ifNull": ["$goods_type", "Not specified"]}),
                {"$sort": {
                    "collectible_outstanding": -1, "unreconciled_bhada": -1,
                    "unreconciled_lrs": -1, "unpriced_lrs": -1, "_id": 1,
                }},
                {"$limit": 200},
            ],
            "goods_count": [
                {"$group": {"_id": {"$ifNull": ["$goods_type", "Not specified"]}}},
                {"$count": "count"},
            ],
            "by_site": _receivable_group_stages("$site_id"),
        }},
    ]
    result = await db.site_lrs.aggregate(pipeline).to_list(1)
    facets = result[0] if result else {}
    by_site = {
        row["_id"]: _format_receivable_row(row)
        for row in facets.get("by_site", [])
    }

    def dimension_rows(name: str, count_name: str) -> dict[str, Any]:
        rows = []
        for row in facets.get(name, []):
            group_key = row.get("_id")
            label = group_key.get("label") if isinstance(group_key, dict) else group_key
            item = _format_receivable_row(row, str(label or "Not specified"))
            if isinstance(group_key, dict) and group_key.get("site_id") is not None:
                item["site_id"] = group_key["site_id"]
                item["operating_date"] = group_key.get("operating_date")
            rows.append(item)
        count_rows = facets.get(count_name, [])
        group_count = count_rows[0].get("count", 0) if count_rows else 0
        return {"rows": rows, "group_count": group_count, "truncated": group_count > 200}

    summary = {
        key: sum((_decimal_value(value.get(key)) for value in by_site.values()), Decimal("0.00"))
        for key in ("reconciled_rent", "posted_payments", "collectible_outstanding", "unreconciled_bhada")
    }
    summary_counts = {
        key: sum(value.get(key, 0) for value in by_site.values())
        for key in ("unpaid_lrs", "unreconciled_lrs", "unpriced_lrs", "lr_count", "parcels")
    }
    return {
        "basis": "Collectible outstanding is reconciled rent less net successfully posted payments. Unreconciled and unpriced LRs are shown separately and excluded from collectible outstanding.",
        **{key: format(value, ".2f") for key, value in summary.items()},
        **summary_counts,
        "by_receiver": dimension_rows("by_receiver", "receiver_count"),
        "by_goods": dimension_rows("by_goods", "goods_count"),
        "by_site": by_site,
    }


async def _trip_expense_postings_exist(site: dict[str, Any], expense: dict[str, Any]) -> bool:
    reference = _scoped(
        str(site["business_id"]), site["_id"],
        ref_type="site_trip_expense", ref_id=expense["_id"],
    )
    if not await db.cashbook.find_one(reference):
        return False
    if expense.get("category") == "Driver Advance" and expense.get("driver_id"):
        driver_reference = {
            **reference, "entity_type": "driver", "entity_id": expense["driver_id"],
        }
        if not await db.ledger.find_one(driver_reference):
            return False
    return True


async def _post_site_trip_expense(site: dict[str, Any], expense: dict[str, Any]) -> None:
    reference = _scoped(
        str(site["business_id"]), site["_id"],
        ref_type="site_trip_expense", ref_id=expense["_id"],
    )
    cash_record = await db.cashbook.find_one(reference)
    if not cash_record:
        method = expense["mode"]
        account = "cash" if method in ("Cash", "Other") else "bank"
        assignment = expense.get("vehicle_no") or expense.get("driver_name") or ""
        payee = f" paid to {expense['payee']}" if expense.get("payee") else ""
        description = (
            f"{expense['category']} expense{payee} - {assignment} "
            f"(trip {expense.get('trip_no', '')}: {expense['remarks']})"
        ).strip()
        try:
            await L.cash(
                account, "out", expense["amount"], expense["date"], description,
                "site_trip_expense", expense["_id"], mode=method,
                site_id=site["_id"], business_id=str(site["business_id"]),
            )
        except DuplicateKeyError:
            if not await db.cashbook.find_one(reference):
                raise
    if expense.get("category") == "Driver Advance" and expense.get("driver_id"):
        driver_reference = {
            **reference, "entity_type": "driver", "entity_id": expense["driver_id"],
        }
        if not await db.ledger.find_one(driver_reference):
            try:
                await L.post(
                    "driver", expense["driver_id"], expense["date"],
                    "Advance given (expense entry)", debit=expense["amount"],
                    ref_type="site_trip_expense", ref_id=expense["_id"],
                    site_id=site["_id"], business_id=str(site["business_id"]),
                )
            except DuplicateKeyError:
                if not await db.ledger.find_one(driver_reference):
                    raise


async def _ensure_site_trip_expense_posted(
    site: dict[str, Any], expense: dict[str, Any],
) -> dict[str, Any]:
    query = _scoped(
        str(site["business_id"]), site["_id"],
        trip_id=expense["trip_id"], idempotency_key=expense["idempotency_key"],
    )
    if expense.get("posting_status") == "posted":
        return expense
    stale_before = (datetime.now(timezone.utc) - timedelta(minutes=2)).isoformat()
    claim = await db.expenses.update_one(
        {
            **query, "_id": expense["_id"],
            "$or": [
                {"posting_status": {"$in": ["pending", "failed"]}},
                {"posting_status": "posting", "posting_started_at": {"$lte": stale_before}},
            ],
        },
        {"$set": {"posting_status": "posting", "posting_started_at": now_iso()}},
    )
    if not claim.modified_count:
        current = await db.expenses.find_one(query)
        if current and current.get("posting_status") == "posting":
            if await _trip_expense_postings_exist(site, current):
                await db.expenses.update_one(
                    {**query, "_id": current["_id"], "posting_status": "posting"},
                    {"$set": {"posting_status": "posted", "posted_at": now_iso()},
                     "$unset": {"posting_error": ""}},
                )
                current["posting_status"] = "posted"
                return current
        if current and current.get("posting_status") == "posted":
            return current
        raise HTTPException(
            409, "This expense is being posted or needs posting recovery; retry shortly",
        )
    try:
        await _post_site_trip_expense(site, expense)
    except Exception as error:
        await db.expenses.update_one(
            {**query, "_id": expense["_id"], "posting_status": "posting"},
            {"$set": {"posting_status": "failed", "posting_error": type(error).__name__}},
        )
        raise HTTPException(500, "Expense was saved but accounting posting failed; retry this request") from error
    await db.expenses.update_one(
        {**query, "_id": expense["_id"], "posting_status": "posting"},
        {"$set": {"posting_status": "posted", "posted_at": now_iso()},
         "$unset": {"posting_error": ""}},
    )
    expense["posting_status"] = "posted"
    return expense


async def _site_trip_expense_response(
    site: dict[str, Any], trip_id: str, user: Optional[dict[str, Any]] = None,
) -> dict[str, Any]:
    expenses, total, count = await _trip_expense_rows(site, trip_id, user)
    settings = await db.settings.find_one({"_id": "settings"}) or {}
    return {
        "expenses": expenses,
        "total_amount": format(total, ".2f"),
        "total": count,
        "expense_categories": settings.get("expense_categories") or [],
        "payment_modes": settings.get("payment_modes") or [],
    }


async def _lr(site: dict[str, Any], trip_id: str, lr_id: str) -> dict[str, Any]:
    lr = await db.site_lrs.find_one(_scoped(
        str(site["business_id"]), site["_id"], _id=lr_id, trip_id=trip_id,
    ))
    if not lr:
        raise HTTPException(404, "LR not found for this trip")
    return lr


async def _payments_by_lr(site, trip_id, lr_ids):
    if not lr_ids:
        return {}
    out = {}
    match = _scoped(str(site["business_id"]), site["_id"],
                    lr_id={"$in": lr_ids}, posting_status={"$ne": "rejected"})
    if trip_id is not None:
        match["trip_id"] = trip_id
    async for row in db.site_payments.aggregate([
        {"$match": match},
        {"$group": {
            "_id": "$lr_id",
            "net": {"$sum": {"$cond": [{"$eq": ["$kind", "reversal"]},
                                       {"$multiply": ["$amount", -1]}, "$amount"]}},
            "count": {"$sum": 1},
        }},
    ]):
        out[row["_id"]] = _decimal_value(row["net"])
    return out


def _payment_status(rent: Any, paid: Decimal) -> str:
    if rent is None:
        return "unpriced"
    amount_due = _decimal_value(rent)
    if paid >= amount_due:
        return "paid"
    return "partial" if paid > 0 else "unpaid"


def _system_lr_report_query(
    business_id: str, site_ids: list[str], start: str, end: str, search: Optional[str],
) -> dict[str, Any]:
    query = {
        "business_id": business_id,
        "site_id": {"$in": site_ids},
        "operating_date": {"$gte": start, "$lte": end},
    }
    term = (search or "").strip()
    if term:
        escaped = re.escape(term)
        query["$or"] = [
            {"lr_ref": {"$regex": escaped, "$options": "i"}},
            {"receiver_name": {"$regex": escaped, "$options": "i"}},
            {"receiver_label": {"$regex": escaped, "$options": "i"}},
            {"goods_type": {"$regex": escaped, "$options": "i"}},
            {"trip_ref": {"$regex": escaped, "$options": "i"}},
        ]
    return query


def _activity_by_date(start: str, end: str, trip_rows: list[dict], lr_rows: list[dict]) -> list[dict]:
    activity = {}
    for field, rows in (("total_trips", trip_rows), ("total_lrs", lr_rows)):
        for row in rows:
            day = row["_id"]
            activity.setdefault(day, {"date": day, "total_trips": 0, "total_lrs": 0})
            activity[day][field] = row["count"]

    current = date.fromisoformat(start)
    last = date.fromisoformat(end)
    result = []
    while current <= last:
        day = current.isoformat()
        result.append(activity.get(day, {"date": day, "total_trips": 0, "total_lrs": 0}))
        current += timedelta(days=1)
    return result


async def _lr_rows(
    site: dict[str, Any], trip_id: str, rows: list[dict[str, Any]],
    user: Optional[dict[str, Any]] = None,
) -> list[dict[str, Any]]:
    paid = await _payments_by_lr(site, trip_id, [r["_id"] for r in rows])
    settled_by_ledger = set()
    if user and user.get("role") == "owner" and rows:
        async for settlement in db.site_payments.aggregate([
            {"$match": _scoped(
                str(site["business_id"]), site["_id"], trip_id=trip_id,
                lr_id={"$in": [row["_id"] for row in rows]},
                kind="payment", source="booking_settlement", posting_status="posted",
            )},
            {"$lookup": {
                "from": "site_payments", "let": {"payment_id": "$_id"},
                "pipeline": [{"$match": {"$expr": {"$and": [
                    {"$eq": ["$reversal_of", "$$payment_id"]},
                    {"$eq": ["$kind", "reversal"]},
                    {"$eq": ["$posting_status", "posted"]},
                ]}}}, {"$limit": 1}],
                "as": "reversals",
            }},
            {"$match": {"reversals": {"$eq": []}}},
            {"$group": {"_id": "$lr_id"}},
        ]):
            settled_by_ledger.add(settlement["_id"])
    result = []
    for row in rows:
        item = _doc_for_user(row, user or {})
        item["charge_editor_marker"] = _charge_editor_marker(row)
        item["entry_by"] = _entry_by(row)
        if user and user.get("role") == "owner":
            item["charge_editor_name"] = row.get("charge_updated_by_name", "")
        item.setdefault("receipt_date", item.get("operating_date", ""))
        item.setdefault("sender_address", "")
        item.setdefault("receiver_address", "")
        item.setdefault("sender_address_hindi", item.get("sender_address", ""))
        item.setdefault("receiver_address_hindi", item.get("receiver_address", ""))
        item.setdefault("city", "")
        item.setdefault("sender_name_hindi", item.get("sender_name", ""))
        item.setdefault("receiver_name_hindi", item.get("receiver_name", ""))
        source_goods = item.get("goods_rows") or item.get("containers") or []
        legacy_charge_allocation = bool(source_goods) and not any(
            "rent" in line or "hamali" in line for line in source_goods
        )
        item["goods_rows"] = [
            {
                "type": line.get("type", ""),
                "type_hindi": line.get("type_hindi", line.get("type", "")),
                "quantity": line.get("quantity", 0),
                "description": line.get("description", ""),
                "description_hindi": line.get("description_hindi", line.get("description", "")),
                "rent": format(
                    _decimal_value(row.get("rent")) if legacy_charge_allocation and index == 0
                    else _decimal_value(line.get("rent")), ".2f",
                ),
                "hamali": format(
                    _decimal_value(row.get("hamali")) if legacy_charge_allocation and index == 0
                    else _decimal_value(line.get("hamali")), ".2f",
                ),
            }
            for index, line in enumerate(source_goods)
        ]
        item["containers"] = item["goods_rows"]
        item["legacy_charge_allocation"] = legacy_charge_allocation
        receipt_fee = _decimal_value(row.get("receipt_fee", Decimal128(Decimal("2.00"))))
        item["receipt_fee"] = format(receipt_fee, ".2f")
        item["total_rent"] = format(
            _decimal_value(row.get("rent")) + _decimal_value(row.get("hamali"))
            + receipt_fee, ".2f",
        )
        paid_total = paid.get(row["_id"], Decimal("0.00"))
        item["paid_total"] = format(paid_total, ".2f")
        amount_due = _decimal_value(row.get("rent")) + receipt_fee
        item["outstanding"] = format(
            max(amount_due - paid_total, Decimal("0.00"))
            if row.get("rent") is not None else Decimal("0.00"), ".2f")
        item["payment_status"] = _payment_status(
            amount_due if row.get("rent") is not None else None, paid_total,
        )
        if user and user.get("role") == "owner":
            item["ledger_settlement_created"] = row["_id"] in settled_by_ledger
            item["amount_paid"] = row.get("booking_ledger_paid", (
                row.get("rent") is not None
                and paid_total >= amount_due
                and not row.get("voided")
            ))
        permissions = (user.get("site_permissions") or {}).get(site["_id"], []) if user else []
        if (user and user.get("role") == "site_manager"
                and not {"finance:read", "finance:update", "lrs:create", "lrs:update"}.intersection(permissions)):
            for key in ("rent", "hamali", "total_rent", "paid_total", "outstanding", "payment_status"):
                item.pop(key, None)
            for line in item["goods_rows"]:
                line.pop("rent", None)
                line.pop("hamali", None)
        result.append(item)
    return result


async def _financial_event(event_id, site, party_id, event_date, description, delta,
                           cash_direction=None, mode=None, user=None):
    amount = abs(delta).quantize(Decimal("0.01"))
    if amount <= 0:
        return
    await db.site_financial_events.update_one(
        {"_id": event_id},
        {"$setOnInsert": {
            "business_id": str(site["business_id"]), "site_id": site["_id"],
            "party_id": party_id, "date": event_date, "description": description,
            "amount": Decimal128(amount), "direction": "debit" if delta > 0 else "credit",
            "cash_direction": cash_direction, "mode": mode, "status": "pending",
            **(_creator_metadata(user) if user else {"created_at": now_iso()}),
        }},
        upsert=True,
    )
    event = await db.site_financial_events.find_one({
        "_id": event_id, "business_id": str(site["business_id"]), "site_id": site["_id"],
    })
    if not event:
        raise HTTPException(500, "Could not load the staged LR financial adjustment")
    if (_decimal_value(event.get("amount")) != amount
            or event.get("direction") != ("debit" if delta > 0 else "credit")
            or event.get("party_id") != party_id):
        raise HTTPException(409, "Financial event identifier was already used for a different adjustment")
    if event.get("status") == "complete":
        return
    ledger_exists = await db.ledger.find_one({
        "business_id": str(site["business_id"]), "site_id": site["_id"],
        "ref_type": "site_lr_financial", "ref_id": event_id,
    })
    if not ledger_exists:
        try:
            await L.post(
                "party", party_id, event_date, description,
                debit=float(amount) if delta > 0 else 0,
                credit=float(amount) if delta < 0 else 0,
                ref_type="site_lr_financial", ref_id=event_id,
                site_id=site["_id"], business_id=str(site["business_id"]),
            )
        except DuplicateKeyError:
            if not await db.ledger.find_one({
                "business_id": str(site["business_id"]), "site_id": site["_id"],
                "ref_type": "site_lr_financial", "ref_id": event_id,
            }):
                raise
    if cash_direction:
        cash_exists = await db.cashbook.find_one({
            "business_id": str(site["business_id"]), "site_id": site["_id"],
            "ref_type": "site_lr_financial", "ref_id": event_id,
        })
        if not cash_exists:
            try:
                await L.cash(
                    "cash" if mode in ("Cash", "Other") else "bank",
                    cash_direction, float(amount), event_date, description,
                    "site_lr_financial", event_id, mode=mode,
                    site_id=site["_id"], business_id=str(site["business_id"]),
                )
            except DuplicateKeyError:
                if not await db.cashbook.find_one({
                    "business_id": str(site["business_id"]), "site_id": site["_id"],
                    "ref_type": "site_lr_financial", "ref_id": event_id,
                }):
                    raise
    await db.site_financial_events.update_one(
        {"_id": event_id}, {"$set": {"status": "complete", "completed_at": now_iso()}},
    )


async def _receiver_for_lr(site, user, trip, name, identifier, action, selected_id=None):
    original = _normalize_label(name, "Receiver", 160)
    normalized = original.casefold()
    day = trip["operating_date"]
    matches = await db.site_receivers.find(_scoped(
        str(site["business_id"]), site["_id"], operating_date=day,
    )).sort("created_at", -1).limit(5001).to_list(5001)
    if len(matches) > 5000:
        raise HTTPException(413, "Receiver duplicate check reached its daily limit; ask the owner to review")
    similar = [r for r in matches if SequenceMatcher(None, normalized, r["normalized_name"]).ratio() >= 0.86]
    if similar and action not in ("same", "different"):
        raise HTTPException(409, {
            "code": "receiver_match_confirmation_required",
            "message": "A similar receiver was used at this site on this operating date. Confirm whether this is the same person.",
            "matches": [{"id": r["_id"], "name": r["name"], "label": r["label"],
                         "identifier": r.get("identifier", "")} for r in similar[:8]],
        })
    if action == "same":
        selected = next((r for r in similar if r["_id"] == selected_id), None)
        if not selected:
            raise HTTPException(422, "Choose a receiver from the matching list")
        identity = selected
    else:
        identity_value = " ".join((identifier or "").split())
        if action not in (None, "different"):
            raise HTTPException(422, "Receiver duplicate action is invalid")
        duplicates = [r for r in similar if r["normalized_name"] == normalized]
        label = original
        if duplicates:
            if identity_value:
                label = f"{original} ({identity_value})"
            else:
                used = {int(m.group(1)) for r in duplicates
                        if (m := re.search(r" (\d+)$", r.get("label", "")))}
                suffix = 1
                while suffix in used:
                    suffix += 1
                label = f"{original} {suffix:02d}"
        identity = {
            "_id": new_id(), "business_id": str(site["business_id"]), "site_id": site["_id"],
            "operating_date": day, "name": original, "normalized_name": normalized,
            "identifier": identity_value, "label": label, **_creator_metadata(user),
        }
        await db.site_receivers.insert_one(identity)
    return identity


async def _party_for_receiver(site, identity, phone=None, user=None):
    q = _scoped(str(site["business_id"]), site["_id"],
                site_receiver_id=identity["_id"])
    party = await db.parties.find_one(q)
    if party:
        if phone:
            await db.parties.update_one(q, {"$set": {"mobile": phone}})
        return party["_id"]
    party_id = new_id()
    try:
        await db.parties.insert_one({
            "_id": party_id, "name": identity["label"], "mobile": phone or "",
            "status": "Active", "archived": False, "site_id": site["_id"],
            "business_id": str(site["business_id"]), "site_receiver_id": identity["_id"],
            **(_creator_metadata(user) if user else {"created_at": now_iso()}),
        })
    except DuplicateKeyError:
        existing = await db.parties.find_one(q)
        if not existing:
            raise
        if phone:
            await db.parties.update_one(q, {"$set": {"mobile": phone}})
        return existing["_id"]
    return party_id


@router.get("/sites/{site_id}/dashboard")
async def site_dashboard(site_id: str, u=Depends(site_user)):
    site = await _site_for_user(site_id, u, "dashboard:read", include_inactive_owner=False)
    business_id = str(site["business_id"])
    day, _ = _now_at_site(site)
    trip_filter = _scoped(business_id, site_id, operating_date=day)
    today_trips = await db.site_trips.count_documents(trip_filter)
    open_trips = await db.site_trips.count_documents(_scoped(business_id, site_id, status="open"))
    closed_trips = await db.site_trips.count_documents(_scoped(business_id, site_id, status="closed"))
    pending_ledgers = await db.site_trips.count_documents(
        _scoped(business_id, site_id, status="closed", reconciled=False))
    recent = await db.site_audit_events.find(_scoped(business_id, site_id)).sort(
        "created_at", -1).limit(12).to_list(12)
    manager = None
    if site.get("manager_username"):
        manager = await platform_db.users.find_one({
            "_id": site["manager_username"], "tenant_id": business_id,
            "role": "site_manager", "active": {"$ne": False},
        }, {"_id": 1})
    lrs = await db.site_lrs.aggregate([
        {"$match": _scoped(business_id, site_id)},
        {"$group": {
            "_id": None, "total": {"$sum": 1}, "parcels": {"$sum": "$total_quantity"},
            "recorded_rent": {"$sum": {"$ifNull": ["$rent", Decimal128(Decimal("0"))]}},
            "recorded_hamali": {"$sum": {"$ifNull": ["$hamali", Decimal128(Decimal("0"))]}},
            "reconciled_rent": {"$sum": {"$cond": ["$reconciled", "$rent", Decimal128(Decimal("0"))]}},
        }},
    ]).to_list(1)
    payments = await db.site_payments.aggregate([
        {"$match": _scoped(
            business_id, site_id, posting_status={"$ne": "rejected"},
        )},
        {"$group": {
            "_id": None,
            "collected": {"$sum": {"$cond": [{"$eq": ["$kind", "reversal"]},
                                             {"$multiply": ["$amount", -1]}, "$amount"]}},
        }},
    ]).to_list(1)
    expense_total = Decimal("0.00")
    expense_count = 0
    recent_expenses = []
    if _can_read_site_finances(u, site_id):
        expense_summary = await db.expenses.aggregate([
            {"$match": _scoped(
                business_id, site_id, cancelled=False, trip_id={"$exists": True, "$nin": [None, ""]},
            )},
            {"$group": {"_id": None, "total": {"$sum": "$amount"}, "count": {"$sum": 1}}},
        ]).to_list(1)
        if expense_summary:
            expense_total = _decimal_value(expense_summary[0].get("total"))
            expense_count = expense_summary[0].get("count", 0)
        recent_expenses = await db.expenses.find(_scoped(
            business_id, site_id, cancelled=False, trip_id={"$exists": True, "$nin": [None, ""]},
        )).sort([("date", -1), ("created_at", -1)]).limit(5).to_list(5)
    lr_sum = lrs[0] if lrs else {}
    paid = _decimal_value((payments[0] if payments else {}).get("collected"))
    rent = _decimal_value(lr_sum.get("recorded_rent"))
    reconciled_rent = _decimal_value(lr_sum.get("reconciled_rent"))
    result = {
        "site": _doc(site), "date": day, "active_managers": 1 if manager else 0,
        "trips_today": today_trips, "open_trips": open_trips, "closed_trips": closed_trips,
        "pending_reconciliation": pending_ledgers,
        "total_lrs": lr_sum.get("total", 0), "total_parcels": lr_sum.get("parcels", 0),
        "recorded_bhada": format(rent, ".2f"),
        "reconciled_bhada": format(_decimal_value(lr_sum.get("reconciled_rent")), ".2f"),
        "collected_bhada": format(paid, ".2f"),
        "outstanding_bhada": format(max(reconciled_rent - paid, Decimal("0.00")), ".2f"),
        "recorded_outstanding_bhada": format(max(rent - paid, Decimal("0.00")), ".2f"),
        "recorded_hamali": format(_decimal_value(lr_sum.get("recorded_hamali")), ".2f"),
        "recent_activity": _audit_rows_for_user(recent, u, site_id),
    }
    if _can_read_site_finances(u, site_id):
        result["trip_expenses"] = format(expense_total, ".2f")
        result["trip_expense_count"] = expense_count
        result["recent_trip_expenses"] = [_site_expense_doc(row, u) for row in recent_expenses]
    if u.get("role") == "site_manager" and "finance:read" not in (u.get("site_permissions") or {}).get(site_id, []):
        for key in ("recorded_bhada", "reconciled_bhada", "collected_bhada",
                    "outstanding_bhada", "recorded_outstanding_bhada", "recorded_hamali"):
            result.pop(key, None)
    return result


@router.get("/sites/system-dashboard")
async def system_dashboard(
    from_date: Optional[str] = None, to_date: Optional[str] = None,
    site_id: Optional[str] = None, trip_id: Optional[str] = None,
    trip_status: Optional[str] = None,
    receiver: Optional[str] = None, u=Depends(site_user),
):
    _owner(u)
    await initialize_site_storage()
    business_id = _tenant_id(u)
    start = _validate_iso_date(from_date).isoformat() if from_date else date.today().isoformat()
    end = _validate_iso_date(to_date).isoformat() if to_date else start
    if start > end or (date.fromisoformat(end) - date.fromisoformat(start)).days > 366:
        raise HTTPException(422, "Date range must be ordered and no longer than 366 days")
    if trip_status and trip_status not in ("open", "closed"):
        raise HTTPException(422, "Trip status must be open or closed")
    site_query = {"business_id": business_id}
    if site_id:
        site_query["_id"] = site_id
    sites = await db.sites.find(site_query).sort("name", 1).to_list(500)
    if site_id and not sites:
        raise HTTPException(404, "Site not found")
    selected_ids = [s["_id"] for s in sites]
    if not sites:
        zero = {
            "active_sites": 0, "active_managers": 0, "trips_today": 0,
            "total_trips": 0,
            "open_trips": 0, "closed_trips": 0, "pending_reconciliation": 0,
            "total_lrs": 0, "total_parcels": 0, "recorded_bhada": "0.00",
            "reconciled_bhada": "0.00", "collected_bhada": "0.00",
            "outstanding_bhada": "0.00", "recorded_outstanding_bhada": "0.00",
            "unreconciled_bhada": "0.00", "unreconciled_lrs": 0, "unpriced_lrs": 0,
            "recorded_hamali": "0.00", "trip_expenses": "0.00",
            "trip_expense_count": 0,
        }
        return {
            "from_date": start, "to_date": end, "sites": [], "totals": zero,
            "receivables": {
                "basis": "Collectible outstanding is reconciled rent less net successfully posted payments. Unreconciled and unpriced LRs are shown separately and excluded from collectible outstanding.",
                "reconciled_rent": "0.00", "posted_payments": "0.00",
                "collectible_outstanding": "0.00", "unreconciled_bhada": "0.00",
                "unpaid_lrs": 0, "unreconciled_lrs": 0, "unpriced_lrs": 0,
                "lr_count": 0, "parcels": 0,
                "by_receiver": {"rows": [], "group_count": 0, "truncated": False},
                "by_goods": {"rows": [], "group_count": 0, "truncated": False},
            },
            "recent_activity": [],
            "activity_by_date": [],
            "alerts": {"trips_awaiting_closure": 0, "ledgers_awaiting_upload": 0,
                       "ledger_imports_needing_review": 0, "unpaid_lrs": 0,
                       "outstanding_bhada": "0.00"},
        }
    if trip_id:
        selected_trip = await db.site_trips.find_one({
            "business_id": business_id, "site_id": {"$in": selected_ids}, "_id": trip_id,
        }, {"_id": 1})
        if not selected_trip:
            raise HTTPException(404, "Trip not found in the selected business/site")
    matching_trip_ids = None
    if receiver:
        receiver_query = {
            "business_id": business_id, "site_id": {"$in": selected_ids},
            "operating_date": {"$gte": start, "$lte": end},
            "receiver_name": {"$regex": re.escape(receiver.strip()), "$options": "i"},
        }
        matching_trip_ids = await db.site_lrs.distinct("trip_id", receiver_query)
    trip_query = {"business_id": business_id, "site_id": {"$in": selected_ids},
                  "operating_date": {"$gte": start, "$lte": end}}
    trip_filters = []
    if trip_status in ("open", "closed"):
        trip_filters.append({"status": trip_status})
    if matching_trip_ids is not None:
        trip_filters.append({"_id": {"$in": matching_trip_ids}})
    if trip_id:
        trip_filters.append({"_id": trip_id})
    if trip_filters:
        trip_query["$and"] = trip_filters
    matching_site_trip_ids = await db.site_trips.distinct("_id", trip_query)
    expense_by_site = await _site_trip_expense_totals({
        "business_id": business_id, "site_id": {"$in": selected_ids},
        "trip_id": {"$in": matching_site_trip_ids}, "date": {"$gte": start, "$lte": end},
        "cancelled": False,
    }) if matching_site_trip_ids else {}
    today_pairs = [
        {"site_id": current_site["_id"], "operating_date": _now_at_site(current_site)[0]}
        for current_site in sites
    ]
    trip_aggregates = await db.site_trips.aggregate([
        {"$match": trip_query},
        {"$facet": {
            "by_site": [
                {"$group": {"_id": "$site_id", "count": {"$sum": 1},
                    "open": {"$sum": {"$cond": [{"$eq": ["$status", "open"]}, 1, 0]}},
                    "closed": {"$sum": {"$cond": [{"$eq": ["$status", "closed"]}, 1, 0]}},
                    "pending_reconciliation": {"$sum": {"$cond": [
                        {"$and": [{"$eq": ["$status", "closed"]}, {"$ne": ["$reconciled", True]}]}, 1, 0]}},
                }},
            ],
            "by_date": [
                {"$group": {"_id": "$operating_date", "count": {"$sum": 1}}},
                {"$sort": {"_id": 1}},
            ],
        }},
    ]).to_list(1)
    trip_aggregates = trip_aggregates[0] if trip_aggregates else {"by_site": [], "by_date": []}
    trip_by_site = {row["_id"]: row for row in trip_aggregates["by_site"]}
    today_query = {"business_id": business_id, "$or": today_pairs}
    today_filters = []
    if trip_status in ("open", "closed"):
        today_filters.append({"status": trip_status})
    if receiver:
        matching_today = await db.site_lrs.distinct("trip_id", {
            "business_id": business_id, "$or": today_pairs,
            "receiver_name": {"$regex": re.escape(receiver.strip()), "$options": "i"},
        })
        today_filters.append({"_id": {"$in": matching_today}})
    if trip_id:
        today_filters.append({"_id": trip_id})
    if today_filters:
        today_query["$and"] = today_filters
    today_stats = await db.site_trips.aggregate([
        {"$match": today_query}, {"$group": {"_id": "$site_id", "count": {"$sum": 1}}},
    ]).to_list(500)
    today_by_site = {row["_id"]: row["count"] for row in today_stats}
    lr_query = {"business_id": business_id, "site_id": {"$in": selected_ids},
                "operating_date": {"$gte": start, "$lte": end}}
    if receiver:
        rx = re.escape(receiver.strip())
        lr_query["receiver_name"] = {"$regex": rx, "$options": "i"}
    allowed_trips = None
    if trip_status in ("open", "closed"):
        allowed_trips = await db.site_trips.distinct(
            "_id", {"business_id": business_id, "site_id": {"$in": selected_ids},
                    "status": trip_status, "operating_date": {"$gte": start, "$lte": end}},
        )
        lr_query["trip_id"] = {"$in": allowed_trips}
    if trip_id:
        if allowed_trips is not None and trip_id not in allowed_trips:
            lr_query["trip_id"] = {"$in": []}
        else:
            lr_query["trip_id"] = trip_id
    receivables = await _receivables_breakdown(lr_query)
    receivable_by_site = receivables.pop("by_site")
    site_names = {site["_id"]: site.get("name", "") for site in sites}
    for row in receivables["by_receiver"]["rows"]:
        row["site_name"] = site_names.get(row.get("site_id"), "")
    lr_aggregates = await db.site_lrs.aggregate([
        {"$match": lr_query},
        {"$facet": {
            "by_site": [
                {"$group": {"_id": "$site_id", "total": {"$sum": 1},
                    "parcels": {"$sum": "$total_quantity"},
                    "rent": {"$sum": {"$ifNull": ["$rent", Decimal128(Decimal("0"))]}},
                    "reconciled_rent": {"$sum": {"$cond": [
                        "$reconciled", "$rent", Decimal128(Decimal("0"))]}},
                    "hamali": {"$sum": {"$ifNull": ["$hamali", Decimal128(Decimal("0"))]}}}},
            ],
            "by_date": [
                {"$group": {"_id": "$operating_date", "count": {"$sum": 1}}},
                {"$sort": {"_id": 1}},
            ],
        }},
    ]).to_list(1)
    lr_aggregates = lr_aggregates[0] if lr_aggregates else {"by_site": [], "by_date": []}
    lr_by_site = {row["_id"]: row for row in lr_aggregates["by_site"]}
    activity_by_date = _activity_by_date(
        start, end, trip_aggregates["by_date"], lr_aggregates["by_date"],
    )
    active_managers = await platform_db.users.find(
        {"tenant_id": business_id, "role": "site_manager", "active": {"$ne": False}},
        {"_id": 1},
    ).to_list(1000)
    active_manager_names = {manager["_id"] for manager in active_managers}
    active_manager_count = len({
        site.get("manager_username") for site in sites
        if site.get("manager_username") in active_manager_names
    })
    import_query = {
        "business_id": business_id, "site_id": {"$in": selected_ids},
        "status": {"$in": ["needs_review", "failed", "partial"]},
        "created_at": {"$gte": start, "$lte": end + "T23:59:59"},
    }
    import_trip_ids = None
    if allowed_trips is not None:
        import_trip_ids = set(allowed_trips)
    if matching_trip_ids is not None:
        import_trip_ids = (set(matching_trip_ids) if import_trip_ids is None
                           else import_trip_ids.intersection(matching_trip_ids))
    if trip_id:
        import_trip_ids = ({trip_id} if import_trip_ids is None
                           else import_trip_ids.intersection({trip_id}))
    if import_trip_ids is not None:
        import_query["trip_id"] = {"$in": list(import_trip_ids)}
    import_issues = await db.site_ledger_imports.count_documents(import_query)
    rows = []
    for site in sites:
        trips = trip_by_site.get(site["_id"], {})
        freight = lr_by_site.get(site["_id"], {})
        due = _decimal_value(freight.get("rent"))
        reconciled_due = _decimal_value(freight.get("reconciled_rent"))
        collectible = receivable_by_site.get(site["_id"], {})
        collected = _decimal_value(collectible.get("posted_payments"))
        expenses = expense_by_site.get(site["_id"], {})
        rows.append({
            "site": _doc(site), "trips_today": today_by_site.get(site["_id"], 0),
            "total_trips": trips.get("count", 0),
            "open_trips": trips.get("open", 0), "closed_trips": trips.get("closed", 0),
            "pending_reconciliation": trips.get("pending_reconciliation", 0),
            "total_lrs": freight.get("total", 0), "total_parcels": freight.get("parcels", 0),
            "recorded_bhada": format(due, ".2f"),
            "reconciled_bhada": format(_decimal_value(freight.get("reconciled_rent")), ".2f"),
            "collected_bhada": format(collected, ".2f"),
            "outstanding_bhada": collectible.get(
                "collectible_outstanding", format(max(reconciled_due - collected, Decimal("0.00")), ".2f"),
            ),
            "recorded_outstanding_bhada": format(max(due - collected, Decimal("0.00")), ".2f"),
            "recorded_hamali": format(_decimal_value(freight.get("hamali")), ".2f"),
            "unreconciled_bhada": collectible.get("unreconciled_bhada", "0.00"),
            "unreconciled_lrs": collectible.get("unreconciled_lrs", 0),
            "unpriced_lrs": collectible.get("unpriced_lrs", 0),
            "collectible_unpaid_lrs": collectible.get("unpaid_lrs", 0),
            "trip_expenses": format(expenses.get("total", Decimal("0.00")), ".2f"),
            "trip_expense_count": expenses.get("count", 0),
            "active_managers": 1 if site.get("manager_username") in active_manager_names else 0,
            "unpaid_lrs": collectible.get("unpaid_lrs", 0),
        })
    total_rent = sum((_decimal_value(r["recorded_bhada"]) for r in rows), Decimal("0.00"))
    total_reconciled = sum((_decimal_value(r["reconciled_bhada"]) for r in rows), Decimal("0.00"))
    total_collected = sum((_decimal_value(r["collected_bhada"]) for r in rows), Decimal("0.00"))
    total_trip_expenses = sum((_decimal_value(r["trip_expenses"]) for r in rows), Decimal("0.00"))
    audit = await db.site_audit_events.find({
        "business_id": business_id, "site_id": {"$in": selected_ids},
        "created_at": {"$gte": start, "$lte": end + "T23:59:59"},
    }).sort("created_at", -1).limit(20).to_list(20)
    recent_activity = [_doc(row) for row in audit]
    if u.get("role") == "site_manager":
        for event in recent_activity:
            for key in ("actor_id", "actor_username", "actor_name", "actor_role"):
                event.pop(key, None)
    return {
        "from_date": start, "to_date": end, "sites": rows,
        "receivables": receivables,
        "totals": {
            "active_sites": sum(1 for s in sites if s.get("status") == "Active"),
            "active_managers": active_manager_count,
            "trips_today": sum(r["trips_today"] for r in rows),
            "total_trips": sum(r["total_trips"] for r in rows),
            "open_trips": sum(r["open_trips"] for r in rows),
            "closed_trips": sum(r["closed_trips"] for r in rows),
            "pending_reconciliation": sum(r["pending_reconciliation"] for r in rows),
            "total_lrs": sum(r["total_lrs"] for r in rows),
            "total_parcels": sum(r["total_parcels"] for r in rows),
            "recorded_bhada": format(total_rent, ".2f"),
            "reconciled_bhada": format(total_reconciled, ".2f"),
            "collected_bhada": format(total_collected, ".2f"),
            "outstanding_bhada": format(max(total_reconciled - total_collected, Decimal("0.00")), ".2f"),
            "recorded_outstanding_bhada": format(max(total_rent - total_collected, Decimal("0.00")), ".2f"),
            "recorded_hamali": format(sum((_decimal_value(r["recorded_hamali"]) for r in rows), Decimal("0.00")), ".2f"),
            "unreconciled_bhada": receivables["unreconciled_bhada"],
            "unreconciled_lrs": receivables["unreconciled_lrs"],
            "unpriced_lrs": receivables["unpriced_lrs"],
            "trip_expenses": format(total_trip_expenses, ".2f"),
            "trip_expense_count": sum(r["trip_expense_count"] for r in rows),
        },
        "recent_activity": recent_activity,
        "activity_by_date": activity_by_date,
        "alerts": {
            "trips_awaiting_closure": sum(r["open_trips"] for r in rows),
            "ledgers_awaiting_upload": sum(r["pending_reconciliation"] for r in rows),
            "ledger_imports_needing_review": import_issues,
            "unpaid_lrs": sum(r["unpaid_lrs"] for r in rows),
            "outstanding_bhada": format(max(total_reconciled - total_collected, Decimal("0.00")), ".2f"),
        },
    }


@router.get("/sites/system-reports/lrs")
async def system_lr_report(
    from_date: Optional[str] = None, to_date: Optional[str] = None,
    site_id: Optional[str] = None, search: Optional[str] = Query(default=None, max_length=120),
    limit: int = Query(default=50, ge=1, le=100), offset: int = Query(default=0, ge=0),
    u=Depends(site_user),
):
    _owner(u)
    await initialize_site_storage()
    business_id = _tenant_id(u)
    start = _validate_iso_date(from_date).isoformat() if from_date else date.today().isoformat()
    end = _validate_iso_date(to_date).isoformat() if to_date else start
    if start > end or (date.fromisoformat(end) - date.fromisoformat(start)).days > 366:
        raise HTTPException(422, "Date range must be ordered and no longer than 366 days")
    if site_id:
        sites = [await _site_for_user(site_id, u)]
    else:
        sites = await db.sites.find({"business_id": business_id}).sort("name", 1).to_list(500)
    site_ids = [site["_id"] for site in sites]
    if not site_ids:
        return {"rows": [], "total": 0, "limit": limit, "offset": offset,
                "from_date": start, "to_date": end}

    query = _system_lr_report_query(business_id, site_ids, start, end, search)
    total = await db.site_lrs.count_documents(query)
    lrs = await db.site_lrs.find(query).sort(
        [("operating_date", -1), ("sequence", -1)]
    ).skip(offset).limit(limit).to_list(limit)
    site_by_id = {site["_id"]: site for site in sites}
    rows = []
    if lrs:
        lr_ids = [lr["_id"] for lr in lrs]
        trip_ids = list({lr["trip_id"] for lr in lrs})
        payments = await db.site_payments.aggregate([
            {"$match": {
                "business_id": business_id,
                "site_id": {"$in": list({lr["site_id"] for lr in lrs})},
                "trip_id": {"$in": trip_ids}, "lr_id": {"$in": lr_ids},
                "posting_status": {"$ne": "rejected"},
            }},
            {"$group": {
                "_id": "$lr_id",
                "net": {"$sum": {"$cond": [
                    {"$eq": ["$kind", "reversal"]},
                    {"$multiply": ["$amount", -1]}, "$amount",
                ]}},
            }},
        ]).to_list(limit)
        paid_by_lr = {row["_id"]: _decimal_value(row["net"]) for row in payments}
        for lr in lrs:
            item = _doc(lr)
            paid = paid_by_lr.get(lr["_id"], Decimal("0.00"))
            item["paid_total"] = format(paid, ".2f")
            item["outstanding"] = format(
                max(_decimal_value(lr.get("rent")) - paid, Decimal("0.00"))
                if lr.get("rent") is not None else Decimal("0.00"), ".2f",
            )
            item["payment_status"] = _payment_status(lr.get("rent"), paid)
            item["site_name"] = site_by_id[lr["site_id"]].get("name", "")
            item["trip_ref"] = lr.get("trip_ref", "")
            rows.append(item)
    return {
        "rows": rows, "total": total, "limit": limit, "offset": offset,
        "from_date": start, "to_date": end,
    }


@router.post("/sites/{site_id}/trips")
async def create_site_trip(site_id: str, body: TripCreate, u=Depends(site_user)):
    site = await _site_for_user(site_id, u, "trips:create", include_inactive_owner=False)
    business_id = str(site["business_id"])
    operating_date = body.operating_date or _now_at_site(site)[0]
    _validate_iso_date(operating_date)
    master_values = await _trip_master_values(body.vehicle_id, body.driver_id)
    sequence = await _next_sequence(site_id, operating_date, "trip")
    trip_ref = f"{site['code']}{date.fromisoformat(operating_date):%d%m%Y}-{sequence:02d}"
    trip = {
        "_id": new_id(), "business_id": business_id, "site_id": site_id,
        "operating_date": operating_date, "timezone": site.get("timezone", "Asia/Kolkata"),
        "sequence": sequence, "trip_ref": trip_ref,
        "truck_no": master_values.get("truck_no") or (
            _normalize_label(body.truck_no, "Truck number", 32) if body.truck_no.strip() else ""
        ),
        "driver_name": master_values.get("driver_name") or (
            _normalize_label(body.driver_name, "Driver name", 100) if body.driver_name.strip() else ""
        ),
        "vehicle_id": master_values.get("vehicle_id"),
        "driver_id": master_values.get("driver_id"),
        "status": "open", "reconciled": False, **_creator_metadata(u),
        "updated_at": now_iso(),
    }
    await db.site_trips.insert_one(trip)
    await _audit(site, u, "trip.created", "trip", trip["_id"], new={
        "trip_ref": trip_ref, "operating_date": operating_date,
        "truck_no": trip["truck_no"], "driver_name": trip["driver_name"],
        "vehicle_id": trip["vehicle_id"], "driver_id": trip["driver_id"],
    })
    return _doc_for_user(trip, u)


@router.get("/sites/{site_id}/trip-resources")
async def list_site_trip_resources(site_id: str, u=Depends(site_user)):
    site = await _site_for_user(site_id, u)
    if u.get("role") == "site_manager":
        permissions = set((u.get("site_permissions") or {}).get(site_id) or [])
        if not permissions.intersection({"trips:read", "trips:create"}):
            raise HTTPException(403, "You do not have permission to view vehicle and driver options for this site")
    vehicles = await db.vehicles.find(
        {"archived": {"$ne": True}},
        {"_id": 1, "vehicle_no": 1, "vehicle_type": 1, "status": 1},
    ).sort("vehicle_no", 1).to_list(2000)
    drivers = await db.drivers.find(
        {"archived": {"$ne": True}},
        {"_id": 1, "name": 1, "mobile": 1},
    ).sort("name", 1).to_list(2000)
    return {
        "vehicles": [{
            "id": row["_id"], "vehicle_no": row.get("vehicle_no", ""),
            "vehicle_type": row.get("vehicle_type", ""), "status": row.get("status", ""),
        } for row in vehicles],
        "drivers": [{
            "id": row["_id"], "name": row.get("name", ""),
        } for row in drivers],
    }


@router.get("/sites/{site_id}/trips")
async def list_site_trips(
    site_id: str, from_date: Optional[str] = None, to_date: Optional[str] = None,
    status: Optional[str] = None, receiver: Optional[str] = None,
    limit: int = Query(default=50, ge=1, le=100), offset: int = Query(default=0, ge=0),
    u=Depends(site_user),
):
    site = await _site_for_user(site_id, u, "trips:read")
    query = _scoped(str(site["business_id"]), site_id)
    if from_date or to_date:
        query["operating_date"] = {}
        if from_date:
            query["operating_date"]["$gte"] = _validate_iso_date(from_date).isoformat()
        if to_date:
            query["operating_date"]["$lte"] = _validate_iso_date(to_date).isoformat()
    if status in ("open", "closed"):
        query["status"] = status
    if receiver:
        matching = await db.site_lrs.distinct(
            "trip_id",
            _scoped(str(site["business_id"]), site_id,
                    receiver_name={"$regex": re.escape(receiver.strip()), "$options": "i"}),
        )
        query["_id"] = {"$in": matching}
    total = await db.site_trips.count_documents(query)
    trips = await db.site_trips.find(query).sort(
        [("operating_date", -1), ("sequence", -1)]
    ).skip(offset).limit(limit).to_list(limit)
    trip_ids = [trip["_id"] for trip in trips]
    receipt_counts = {}
    if trip_ids:
        async for row in db.site_lrs.aggregate([
            {"$match": _scoped(
                str(site["business_id"]), site_id,
                trip_id={"$in": trip_ids}, voided={"$ne": True},
            )},
            {"$group": {"_id": "$trip_id", "count": {"$sum": 1}}},
        ]):
            receipt_counts[row["_id"]] = row["count"]
    trip_rows = [_doc_for_user(row, u) for row in trips]
    for row in trip_rows:
        row["lr_count"] = receipt_counts.get(row["id"], 0)
    return {"rows": trip_rows, "total": total,
            "limit": limit, "offset": offset}


@router.get("/sites/{site_id}/trips/{trip_id}")
async def get_site_trip(site_id: str, trip_id: str, u=Depends(site_user)):
    site = await _site_for_user(site_id, u, "trips:read")
    trip = await _trip(site, trip_id)
    result = _doc_for_user(trip, u)
    if _can_read_site_finances(u, site_id):
        expense_summary = await _site_trip_expense_response(site, trip_id, u)
        result.update(expense_summary)
    return result


@router.get("/sites/{site_id}/trips/{trip_id}/expenses")
async def list_site_trip_expenses(site_id: str, trip_id: str, u=Depends(site_user)):
    site = await _site_for_user(site_id, u)
    if not _can_read_site_finances(u, site_id):
        raise HTTPException(403, "You do not have permission to read site finances")
    await _trip(site, trip_id)
    return await _site_trip_expense_response(site, trip_id, u)


@router.post("/sites/{site_id}/trips/{trip_id}/expenses")
async def create_site_trip_expense(
    site_id: str, trip_id: str, body: SiteTripExpenseCreate, u=Depends(site_user),
):
    site = await _site_for_user(site_id, u, "finance:update")
    trip = await _trip(site, trip_id)
    parsed_amount = _decimal(body.amount, "Expense amount")
    if parsed_amount is None:
        raise HTTPException(422, "Expense amount is required")
    amount = parsed_amount
    if amount <= 0:
        raise HTTPException(422, "Expense amount must be more than zero")
    operating_date = _validate_iso_date(body.date).isoformat()
    description = _normalize_label(body.description, "Expense description", 500)
    category = _normalize_label(body.category, "Expense category", 80)
    payee = _normalize_label(body.payee, "Payee", 160) if body.payee else ""
    payment_method = _normalize_label(body.payment_method, "Payment method", 30)
    settings = await db.settings.find_one({"_id": "settings"})
    if not settings:
        raise HTTPException(503, "Expense and payment categories are not configured")
    if category not in (settings.get("expense_categories") or []):
        raise HTTPException(422, "Select an expense category configured for this business")
    if payment_method not in (settings.get("payment_modes") or []):
        raise HTTPException(422, "Select a payment method configured for this business")
    idempotency_key = _normalize_label(body.idempotency_key, "Idempotency key", 120)
    fingerprint = hashlib.sha256("\x1f".join((
        description, category, format(amount, ".2f"), operating_date,
        payment_method, payee,
    )).encode("utf-8")).hexdigest()
    expense_id = new_id()
    expense = {
        "_id": expense_id, "business_id": str(site["business_id"]), "site_id": site["_id"],
        "trip_id": trip_id, "trip_no": trip.get("trip_ref", ""),
        "date": operating_date, "category": category, "amount": float(amount),
        "vehicle_id": trip.get("vehicle_id"), "vehicle_no": trip.get("truck_no", ""),
        "driver_id": trip.get("driver_id"), "driver_name": trip.get("driver_name", ""),
        "employee_id": None, "vendor": payee, "payee": payee,
        "mode": payment_method, "payment_method": payment_method,
        "proof_url": "", "remarks": description, "description": description,
        "cancelled": False, "idempotency_key": idempotency_key,
        "request_fingerprint": fingerprint, "posting_status": "pending",
        **_creator_metadata(u),
    }
    identity = _scoped(
        str(site["business_id"]), site["_id"], trip_id=trip_id,
        idempotency_key=idempotency_key,
    )
    try:
        stored = await db.expenses.find_one_and_update(
            identity, {"$setOnInsert": expense}, upsert=True,
            return_document=ReturnDocument.AFTER,
        )
    except DuplicateKeyError:
        stored = await db.expenses.find_one(identity)
    if not stored:
        raise HTTPException(500, "Could not load the saved trip expense")
    if stored.get("request_fingerprint") != fingerprint:
        raise HTTPException(409, "This idempotency key was already used for different expense details")
    stored = await _ensure_site_trip_expense_posted(site, stored)
    await _audit(
        site, _creator_user(u, stored), "trip.expense.created", "expense", stored["_id"],
        new={
            "trip_id": trip_id, "trip_ref": trip.get("trip_ref", ""),
            "category": category, "amount": format(amount, ".2f"),
            "date": operating_date, "payment_method": payment_method,
            "payee": payee, "driver_name": trip.get("driver_name", ""),
            "recorded_by": stored.get("created_by_name") or stored.get("created_by", ""),
        },
        event_id=f"site-trip-expense-created:{stored['_id']}",
    )
    return await _site_trip_expense_response(site, trip_id, u)


@router.patch("/sites/{site_id}/trips/{trip_id}")
async def update_site_trip(site_id: str, trip_id: str, body: TripUpdate, u=Depends(site_user)):
    site = await _site_for_user(site_id, u, "trips:update")
    trip = await _trip(site, trip_id)
    if trip.get("reconciled"):
        raise HTTPException(409, "Trip details are locked because its ledger has been reconciled")
    _ensure_trip_ledger_open(trip)
    changes = body.model_dump(exclude_unset=True)
    if not changes:
        return _doc_for_user(trip, u)
    master_values = await _trip_master_values(
        changes.get("vehicle_id"), changes.get("driver_id"),
        allow_archived_vehicle_id=trip.get("vehicle_id"),
        allow_archived_driver_id=trip.get("driver_id"),
    )
    if "vehicle_id" in changes:
        changes["vehicle_id"] = master_values.get("vehicle_id")
        if changes["vehicle_id"]:
            changes["truck_no"] = master_values["truck_no"]
    elif "truck_no" in changes:
        changes["vehicle_id"] = None
    if "driver_id" in changes:
        changes["driver_id"] = master_values.get("driver_id")
        if changes["driver_id"]:
            changes["driver_name"] = master_values["driver_name"]
    elif "driver_name" in changes:
        changes["driver_id"] = None
    for key in ("truck_no", "driver_name"):
        if key in changes:
            changes[key] = _normalize_label(changes[key], key, 32 if key == "truck_no" else 100)
    changes["updated_at"] = now_iso()
    update_query = _scoped(
        str(site["business_id"]), site_id, _id=trip_id, status=trip["status"],
        reconciled={"$ne": True},
        ledger_completed_at={"$exists": False},
        ledger_completion_started_at={"$exists": False},
    )
    updated_result = await db.site_trips.update_one(
        update_query,
        {"$set": changes},
    )
    if not updated_result.matched_count:
        raise HTTPException(409, "Trip details changed or were locked. Reload the trip before saving again.")
    await _audit(site, u, "trip.updated", "trip", trip_id,
                 old={k: trip.get(k) for k in changes if k in trip},
                 new={k: changes[k] for k in changes if k != "updated_at"})
    return _doc_for_user(await _trip(site, trip_id), u)


@router.delete("/sites/{site_id}/trips/{trip_id}")
async def delete_site_trip(site_id: str, trip_id: str, u=Depends(site_user)):
    _owner(u)
    site = await _site_for_user(site_id, u, "trips:update")
    trip = await _trip(site, trip_id)
    business_id = str(site["business_id"])
    linked_records = (
        ("receipts", db.site_lrs),
        ("payments", db.site_payments),
        ("expenses", db.expenses),
        ("ledger entries", db.ledger),
        ("cashbook entries", db.cashbook),
        ("ledger imports", db.site_ledger_imports),
    )
    for label, collection in linked_records:
        if await collection.count_documents(_scoped(
            business_id, site_id, trip_id=trip_id,
        )):
            raise HTTPException(
                409,
                f"This trip cannot be deleted because it has linked {label}.",
            )
    deleted = await db.site_trips.delete_one(
        _scoped(business_id, site_id, _id=trip_id),
    )
    if not deleted.deleted_count:
        raise HTTPException(409, "This trip changed before it could be deleted. Reload and try again.")
    await _audit(
        site, u, "trip.deleted", "trip", trip_id,
        old={
            "trip_ref": trip.get("trip_ref"),
            "operating_date": trip.get("operating_date"),
        },
    )
    return {"deleted": True, "id": trip_id}


@router.post("/sites/{site_id}/trips/{trip_id}/close")
async def close_site_trip(site_id: str, trip_id: str, u=Depends(site_user)):
    _owner(u)
    site = await _site_for_user(site_id, u, "trips:close")
    trip = await _trip(site, trip_id)
    if trip["status"] == "closed":
        return _doc_for_user(trip, u)
    result = await db.site_trips.update_one(
        _scoped(str(site["business_id"]), site_id, _id=trip_id, status="open"),
        {"$set": {"status": "closed", "closed_by": u["username"],
                  "closed_at": now_iso(), "updated_at": now_iso()}},
    )
    if result.modified_count:
        await _audit(site, u, "trip.closed", "trip", trip_id,
                     old={"status": "open"}, new={"status": "closed"})
    return _doc_for_user(await _trip(site, trip_id), u)


@router.post("/sites/{site_id}/trips/{trip_id}/reopen")
async def reopen_site_trip(site_id: str, trip_id: str, body: ReopenTrip, u=Depends(site_user)):
    _owner(u)
    site = await _site_for_user(site_id, u)
    trip = await _trip(site, trip_id)
    if trip["status"] != "closed":
        raise HTTPException(409, "Only a closed trip can be reopened")
    reason = _normalize_label(body.reason, "Reopen reason", 500)
    await db.site_trips.update_one(
        _scoped(str(site["business_id"]), site_id, _id=trip_id, status="closed"),
        {"$set": {"status": "open", "reopened_by": u["username"],
                  "reopened_at": now_iso(), "updated_at": now_iso(),
                  "reopen_reason": reason}},
    )
    await _audit(site, u, "trip.reopened", "trip", trip_id,
                 old={"status": "closed"}, new={"status": "open"}, reason=reason)
    return _doc_for_user(await _trip(site, trip_id), u)


@router.get("/sites/{site_id}/categories")
async def get_site_categories(site_id: str, u=Depends(site_user)):
    site = await _site_for_user(site_id, u, "lrs:read")
    business_id = str(site["business_id"])
    rows = await db.site_categories.find(_scoped(business_id, site_id)).sort(
        [("kind", 1), ("normalized_name", 1)]
    ).to_list(500)
    custom = {"goods": [], "containers": []}
    for row in rows:
        custom.setdefault(row["kind"], []).append(row["name"])
    return {"goods": list(dict.fromkeys(DEFAULT_GOODS + custom.get("goods", []))),
            "containers": list(dict.fromkeys(DEFAULT_CONTAINERS + custom.get("containers", [])))}


@router.post("/sites/{site_id}/categories")
async def add_site_category(site_id: str, body: CategoryCreate, u=Depends(site_user)):
    _owner(u)
    site = await _site_for_user(site_id, u)
    if body.kind not in ("goods", "containers"):
        raise HTTPException(422, "Category kind must be goods or containers")
    name = _normalize_label(body.name, "Category name", 80)
    normalized = name.casefold()
    query = _scoped(str(site["business_id"]), site_id, kind=body.kind, normalized_name=normalized)
    try:
        await db.site_categories.update_one(
            query,
            {"$setOnInsert": {"_id": new_id(), "name": name, "kind": body.kind,
                              "normalized_name": normalized, **_creator_metadata(u)}},
            upsert=True,
        )
    except DuplicateKeyError:
        pass
    category = await db.site_categories.find_one(query)
    if not category:
        raise HTTPException(500, "Category could not be loaded after saving")
    await _audit(site, u, "category.added", "category", category["_id"],
                 new={"kind": body.kind, "name": category["name"]})
    return _doc(category)


async def _save_category(site, user, kind, value):
    normalized = value.casefold()
    query = _scoped(str(site["business_id"]), site["_id"],
                    kind=kind, normalized_name=normalized)
    try:
        await db.site_categories.update_one(
            query,
            {"$setOnInsert": {"_id": new_id(), "name": value, "kind": kind,
                              "normalized_name": normalized, **_creator_metadata(user)}},
            upsert=True,
        )
    except DuplicateKeyError:
        pass


async def _finish_site_lr_creation(site, user, trip, lr):
    lines = lr.get("goods_rows") or lr.get("containers") or []
    await _save_category(site, user, "goods", lr["goods_type"])
    for goods_name in {line["type"] for line in lines}:
        await _save_category(site, user, "goods", goods_name)
    rent = _decimal_value(lr.get("rent"))
    if rent > 0:
        await _financial_event(
            f"site-lr-charge-{lr['_id']}", site, lr["party_id"], trip["operating_date"],
            f"Site LR {lr['lr_ref']} bhada", rent, user=_creator_user(user, lr),
        )
    receipt_fee = _decimal_value(lr.get("receipt_fee", Decimal128(Decimal("2.00"))))
    if receipt_fee > 0:
        await _financial_event(
            f"site-lr-fee-{lr['_id']}", site, lr["party_id"], trip["operating_date"],
            f"Site LR {lr['lr_ref']} receipt fee", receipt_fee, user=_creator_user(user, lr),
        )
    await _audit(
        site, _creator_user(user, lr), "lr.created", "lr", lr["_id"],
        new={"lr_ref": lr["lr_ref"], "trip_id": trip["_id"],
             "goods_type": lr["goods_type"], "total_quantity": lr["total_quantity"],
             "rent": _clean(lr.get("rent")), "hamali": _clean(lr.get("hamali"))},
        event_id=f"site-lr-created:{lr['_id']}",
    )
    await db.site_lrs.update_one(
        _scoped(str(site["business_id"]), site["_id"], _id=lr["_id"]),
        {"$set": {"creation_status": "complete"}},
    )


@router.post("/sites/{site_id}/trips/{trip_id}/lrs")
async def create_site_lr(site_id: str, trip_id: str, body: LRCreate, u=Depends(site_user)):
    site = await _site_for_user(site_id, u, "lrs:create", include_inactive_owner=False)
    trip = await _trip(site, trip_id)
    _ensure_trip_ledger_open(trip)
    if body.idempotency_key:
        previous = await db.site_lrs.find_one(_scoped(
            str(site["business_id"]), site_id, idempotency_key=body.idempotency_key,
        ))
        if previous:
            if previous.get("trip_id") != trip_id:
                raise HTTPException(409, "This receipt request key belongs to another trip")
            if previous.get("creation_status") == "pending":
                await _finish_site_lr_creation(site, u, trip, previous)
            return (await _lr_rows(site, trip_id, [previous], u))[0]
    if trip["status"] != "open":
        raise HTTPException(409, "Lorry receipts can only be added to an open trip")
    receiver = await _receiver_for_lr(
        site, u, trip, body.receiver_name, body.receiver_identifier,
        body.receiver_match_action, body.receiver_identity_id,
    )
    sender_name = _normalize_label(body.sender_name, "Sender", 160) if body.sender_name.strip() else ""
    receipt_controls = _receipt_controls(site)
    sender_phone = _optional_phone(body.sender_phone, "Sender phone")
    receiver_phone = _optional_phone(body.receiver_phone, "Receiver phone")
    goods_type = _normalize_label(body.goods_type, "Goods type", 80)
    receipt_date = body.receipt_date or trip["operating_date"]
    _validate_iso_date(receipt_date)
    containers = []
    for line in body.containers:
        container = {
            "type": _normalize_label(line.type, "Goods type", 80),
            "type_hindi": line.type_hindi.strip() or line.type,
            "quantity": line.quantity,
            "description": " ".join(line.description.split()),
            "description_hindi": " ".join(
                (line.description_hindi or line.description).split()
            ),
        }
        if line.rent is not None:
            container["rent"] = Decimal128(_decimal(line.rent, "Bhada"))
        if line.hamali is not None:
            container["hamali"] = Decimal128(_decimal(line.hamali, "Hamali"))
        containers.append(container)
    per_goods_charges = any(
        line.rent is not None or line.hamali is not None for line in body.containers
    )
    rent = (
        sum((_decimal_value(line.get("rent")) for line in containers), Decimal("0.00"))
        if per_goods_charges else _decimal(body.rent, "Bhada", allow_none=True)
    )
    hamali = (
        sum((_decimal_value(line.get("hamali")) for line in containers), Decimal("0.00"))
        if per_goods_charges else _decimal(body.hamali, "Hamali", allow_none=True)
    )
    sequence = await _next_sequence(site_id, trip["operating_date"], "lr")
    lr_ref = f"{site['code']}LR{date.fromisoformat(trip['operating_date']):%d%m%Y}-{sequence:02d}"
    lr_id = new_id()
    party_id = await _party_for_receiver(site, receiver, receiver_phone, u)
    creator_metadata = _creator_metadata(u)
    lr = {
        "_id": lr_id, "business_id": str(site["business_id"]), "site_id": site_id,
        "trip_id": trip_id, "trip_ref": trip["trip_ref"], "operating_date": trip["operating_date"],
        "timezone": trip["timezone"], "sequence": sequence, "lr_ref": lr_ref,
        "receipt_date": receipt_date,
        "sender_name": sender_name, "sender_name_hindi": body.sender_name_hindi.strip() or sender_name,
        "sender_address": " ".join(body.sender_address.split()), "sender_phone": sender_phone or "",
        "sender_address_hindi": body.sender_address_hindi.strip() or body.sender_address.strip(),
        "receiver_name": receiver["name"], "receiver_phone": receiver_phone or "",
        "receiver_name_hindi": body.receiver_name_hindi.strip() or receiver["name"],
        "receiver_address": " ".join(body.receiver_address.split()),
        "receiver_address_hindi": body.receiver_address_hindi.strip() or body.receiver_address.strip(),
        "city": " ".join(body.city.split()),
        "receiver_normalized": receiver["normalized_name"], "receiver_label": receiver["label"],
        "receiver_identity_id": receiver["_id"], "receiver_identifier": receiver.get("identifier", ""),
        "party_id": party_id, "goods_type": goods_type,
        "goods_type_hindi": body.goods_type_hindi.strip() or containers[0]["type_hindi"],
        "containers": containers, "goods_rows": containers,
        "total_quantity": sum(line["quantity"] for line in containers),
        "rent": Decimal128(rent) if rent is not None else None,
        "hamali": Decimal128(hamali) if hamali is not None else None,
        "receipt_fee": Decimal128(receipt_controls.receipt_fee),
        "creation_status": "pending",
        "receipt_created_at": creator_metadata["created_at"],
        "reconciled": False, **creator_metadata,
        **(_charge_editor_metadata(u)
           if (rent or Decimal("0.00")) > 0 or (hamali or Decimal("0.00")) > 0 else {}),
        "updated_at": now_iso(),
        **({"idempotency_key": body.idempotency_key} if body.idempotency_key else {}),
    }
    try:
        await db.site_lrs.insert_one(lr)
    except DuplicateKeyError:
        if body.idempotency_key:
            previous = await db.site_lrs.find_one(_scoped(
                str(site["business_id"]), site_id, idempotency_key=body.idempotency_key,
            ))
            if previous:
                if previous.get("creation_status") == "pending":
                    await _finish_site_lr_creation(site, u, trip, previous)
                return (await _lr_rows(site, trip_id, [previous], u))[0]
        raise
    await _finish_site_lr_creation(site, u, trip, lr)
    return (await _lr_rows(site, trip_id, [lr], u))[0]


@router.get("/sites/{site_id}/trips/{trip_id}/lrs")
async def list_site_lrs(
    site_id: str, trip_id: str, limit: int = Query(default=50, ge=1, le=100),
    offset: int = Query(default=0, ge=0), u=Depends(site_user),
):
    site = await _site_for_user(site_id, u, "lrs:read")
    await _trip(site, trip_id)
    query = _scoped(str(site["business_id"]), site_id, trip_id=trip_id)
    total = await db.site_lrs.count_documents(query)
    rows = await db.site_lrs.find(query).sort([("sequence", 1)]).skip(offset).limit(limit).to_list(limit)
    return {"rows": await _lr_rows(site, trip_id, rows, u), "total": total,
            "limit": limit, "offset": offset}


@router.get("/sites/{site_id}/trips/{trip_id}/lrs/{lr_id}")
async def get_site_lr(site_id: str, trip_id: str, lr_id: str, u=Depends(site_user)):
    site = await _site_for_user(site_id, u, "lrs:read")
    await _trip(site, trip_id)
    lr = await _lr(site, trip_id, lr_id)
    output = (await _lr_rows(site, trip_id, [lr], u))[0]
    if (u.get("role") == "owner" or "payments:read" in
            (u.get("site_permissions") or {}).get(site_id, [])):
        payments = await db.site_payments.find(
            _scoped(str(site["business_id"]), site_id, trip_id=trip_id, lr_id=lr_id)
        ).sort("created_at", 1).to_list(500)
        output["payments"] = [_doc(p) for p in payments]
    return output


@router.patch("/sites/{site_id}/trips/{trip_id}/lrs/{lr_id}")
async def update_site_lr(
    site_id: str, trip_id: str, lr_id: str, body: LRUpdate, u=Depends(site_user),
):
    site = await _site_for_user(site_id, u, "lrs:update")
    trip = await _trip(site, trip_id)
    _ensure_trip_ledger_open(trip)
    is_owner = u.get("role") == "owner"
    if trip["status"] != "open" and not is_owner:
        raise HTTPException(409, "Reopen the trip before correcting its LRs")
    lr = await _lr(site, trip_id, lr_id)
    if lr.get("voided"):
        raise HTTPException(409, "Voided receipts are read-only")
    changes = body.model_dump(exclude_unset=True)
    idempotency_key = changes.pop("idempotency_key", None)
    expected_updated_at = changes.pop("expected_updated_at", None)
    changes.pop("owner_correction_reason", None)
    if expected_updated_at and lr.get("updated_at") != expected_updated_at:
        raise HTTPException(409, "This receipt changed elsewhere. Reload it before saving your edits.")
    action = changes.pop("receiver_match_action", None)
    identity_id = changes.pop("receiver_identity_id", None)
    old_values = {}
    new_values = {}
    for key, field in (("sender_phone", "Sender phone"),
                       ("receiver_phone", "Receiver phone")):
        if key in changes:
            phone = _optional_phone(changes[key], field)
            if phone is None:
                changes.pop(key)
            else:
                old_values[key] = lr.get(key, "")
                changes[key] = phone
                new_values[key] = phone
    if "receiver_name" in changes or "receiver_identifier" in changes:
        requested_name = " ".join(str(changes.get("receiver_name", lr.get("receiver_name", ""))).split())
        current_name = " ".join(str(lr.get("receiver_name", "")).split())
        requested_identifier = " ".join(str(
            changes.get("receiver_identifier", lr.get("receiver_identifier", ""))
        ).split())
        current_identifier = " ".join(str(lr.get("receiver_identifier", "")).split())
        if requested_name == current_name and requested_identifier == current_identifier:
            changes.pop("receiver_name", None)
            changes.pop("receiver_identifier", None)
    if not changes:
        return (await _lr_rows(site, trip_id, [lr], u))[0]
    financial_fields = {"rent", "hamali"}
    incoming_goods = changes.get("containers") or []
    row_charge_change = bool(incoming_goods) and any(
        line.get("rent") is not None or line.get("hamali") is not None
        for line in incoming_goods
    ) and (
        sum((_decimal_value(line.get("rent")) for line in incoming_goods), Decimal("0.00"))
        != _decimal_value(lr.get("rent"))
        or sum((_decimal_value(line.get("hamali")) for line in incoming_goods), Decimal("0.00"))
        != _decimal_value(lr.get("hamali"))
    )
    requested_charge_change = row_charge_change or any(
        _decimal_value(changes[key]) != _decimal_value(lr.get(key))
        for key in financial_fields.intersection(changes)
    )
    if lr.get("booking_ledger_paid") and requested_charge_change:
        raise HTTPException(
            409,
            "Untick Amount paid before changing Bhada or Hamali.",
        )
    charge_values_changed = False
    if lr.get("reconciled") and (financial_fields.intersection(changes) or row_charge_change) and (
        not is_owner
    ):
        raise HTTPException(409, "Only the business owner can correct reconciled charges")
    if "sender_name" in changes:
        old_values["sender_name"] = lr["sender_name"]
        changes["sender_name"] = (
            _normalize_label(changes["sender_name"], "Sender", 160)
            if str(changes["sender_name"] or "").strip() else ""
        )
        new_values["sender_name"] = changes["sender_name"]
    if "receiver_name" in changes or "receiver_identifier" in changes:
        paid = (await _payments_by_lr(site, trip_id, [lr_id])).get(lr_id, Decimal("0.00"))
        has_financial_history = (
            lr.get("reconciled") or paid or _decimal_value(lr.get("rent"))
            or _decimal_value(lr.get("hamali"))
        )
        if has_financial_history and not is_owner:
            raise HTTPException(
                409,
                "Only the business owner can correct a receiver after charges or payments",
            )
        if has_financial_history and not idempotency_key:
            raise HTTPException(422, "An idempotency key is required for an owner receiver correction")
        identity = await _receiver_for_lr(
            site, u, trip, changes.get("receiver_name", lr["receiver_name"]),
            changes.get("receiver_identifier", lr.get("receiver_identifier", "")),
            action or ("different" if has_financial_history else None), identity_id,
        )
        for key, value in (("receiver_name", identity["name"]),
                           ("receiver_normalized", identity["normalized_name"]),
                           ("receiver_label", identity["label"]),
                           ("receiver_identity_id", identity["_id"]),
                           ("receiver_identifier", identity.get("identifier", ""))):
            old_values[key] = lr.get(key)
            changes[key] = value
            new_values[key] = value
        party_phone = changes.get("receiver_phone", lr.get("receiver_phone"))
        party_id = await _party_for_receiver(site, identity, party_phone, u)
        old_values["party_id"] = lr.get("party_id")
        changes["party_id"] = party_id
        new_values["party_id"] = party_id
        if has_financial_history and party_id != lr.get("party_id"):
            old_party_id = lr["party_id"]
            outstanding = max(_decimal_value(lr.get("rent")) - paid, Decimal("0.00"))
            if outstanding:
                await _financial_event(
                    f"lr-receiver-correction:{lr_id}:{idempotency_key}:old",
                    site, old_party_id, lr["operating_date"],
                    f"Owner correction transfer out for site LR {lr['lr_ref']}",
                    -outstanding, user=u,
                )
                await _financial_event(
                    f"lr-receiver-correction:{lr_id}:{idempotency_key}:new",
                    site, party_id, lr["operating_date"],
                    f"Owner correction transfer in for site LR {lr['lr_ref']}",
                    outstanding, user=u,
                )
    elif "receiver_phone" in changes:
        await db.parties.update_one(
            _scoped(str(site["business_id"]), site_id, _id=lr["party_id"]),
            {"$set": {"mobile": changes["receiver_phone"]}},
        )
    if "goods_type" in changes:
        old_values["goods_type"] = lr.get("goods_type", "")
        changes["goods_type"] = _normalize_label(changes["goods_type"], "Goods type", 80)
        new_values["goods_type"] = changes["goods_type"]
        await _save_category(site, u, "goods", changes["goods_type"])
    for key, max_length in (
        ("sender_name_hindi", 160), ("receiver_name_hindi", 160),
        ("goods_type_hindi", 120),
    ):
        if key in changes:
            changes[key] = " ".join((changes[key] or "").split())[:max_length]
            old_values[key] = lr.get(key, "")
            new_values[key] = changes[key]
    if "containers" in changes:
        for line in changes["containers"]:
            line["description_hindi"] = " ".join(
                (line.get("description_hindi") or line.get("description") or "").split()
            )[:500]
    for key in (
        "sender_address", "receiver_address", "sender_address_hindi",
        "receiver_address_hindi", "city",
    ):
        if key in changes:
            changes[key] = " ".join(str(changes[key] or "").split())
            old_values[key] = lr.get(key, "")
            new_values[key] = changes[key]
    if "receipt_date" in changes:
        changes["receipt_date"] = _validate_iso_date(changes["receipt_date"]).isoformat()
        old_values["receipt_date"] = lr.get("receipt_date", lr.get("operating_date"))
        new_values["receipt_date"] = changes["receipt_date"]
    if "containers" in changes:
        old_values["containers"] = lr.get("containers", lr.get("goods_rows", []))
        source_containers = changes["containers"]
        supplied_goods_charges = any(
            line.get("rent") is not None or line.get("hamali") is not None
            for line in source_containers
        )
        prior_goods = lr.get("goods_rows") or lr.get("containers") or []
        prior_has_allocation = any(
            "rent" in line or "hamali" in line for line in prior_goods
        )
        containers = [{
            "type": _normalize_label(line["type"], "Goods type", 80),
            "type_hindi": " ".join((line.get("type_hindi") or line["type"]).split())[:120],
            "quantity": line["quantity"],
            "description": " ".join((line.get("description") or "").split())[:500],
            "description_hindi": " ".join(
                (line.get("description_hindi") or line.get("description") or "").split()
            )[:500],
            **({
                "rent": Decimal128(_decimal(
                    _decimal_value(line.get("rent")) if supplied_goods_charges else
                    (_decimal_value(prior_goods[index].get("rent"))
                     if prior_has_allocation and index < len(prior_goods)
                     else _decimal_value(lr.get("rent")) if index == 0 else Decimal("0.00")),
                    "Bhada",
                )),
                "hamali": Decimal128(_decimal(
                    _decimal_value(line.get("hamali")) if supplied_goods_charges else
                    (_decimal_value(prior_goods[index].get("hamali"))
                     if prior_has_allocation and index < len(prior_goods)
                     else _decimal_value(lr.get("hamali")) if index == 0 else Decimal("0.00")),
                    "Hamali",
                )),
            }),
        } for index, line in enumerate(changes["containers"])]
        changes["containers"] = containers
        changes["goods_rows"] = containers
        changes["goods_type"] = containers[0]["type"]
        changes["goods_type_hindi"] = containers[0]["type_hindi"]
        changes["total_quantity"] = sum(line["quantity"] for line in containers)
        new_values["containers"] = containers
        new_values["goods_rows"] = containers
        new_values["goods_type"] = changes["goods_type"]
        new_values["goods_type_hindi"] = changes["goods_type_hindi"]
        new_values["total_quantity"] = changes["total_quantity"]
        if any("rent" in line or "hamali" in line for line in containers):
            changes["rent"] = sum(
                (_decimal_value(line.get("rent")) for line in containers), Decimal("0.00"),
            )
            changes["hamali"] = sum(
                (_decimal_value(line.get("hamali")) for line in containers), Decimal("0.00"),
            )
        for goods_name in {line["type"] for line in containers}:
            await _save_category(site, u, "goods", goods_name)
    if "rent" in changes:
        old_rent = _decimal_value(lr.get("rent"))
        new_rent = _decimal(changes["rent"], "Bhada", allow_none=True)
        paid = (await _payments_by_lr(site, trip_id, [lr_id])).get(lr_id, Decimal("0.00"))
        if new_rent is not None and new_rent < paid and not is_owner:
            raise HTTPException(409, "Bhada cannot be reduced below payments already recorded")
        charge_values_changed = charge_values_changed or new_rent != old_rent
        new_values["rent"] = format(new_rent, ".2f") if new_rent is not None else None
        old_values["rent"] = format(old_rent, ".2f") if lr.get("rent") is not None else None
        changes["rent"] = Decimal128(new_rent) if new_rent is not None else None
        delta = (new_rent or Decimal("0.00")) - old_rent
        if delta:
            if not idempotency_key:
                raise HTTPException(422, "An idempotency key is required when changing bhada")
            event_id = f"lr-rent-edit:{lr_id}:{idempotency_key}"
            await _financial_event(
                event_id, site, changes.get("party_id", lr["party_id"]), lr["operating_date"],
                f"Rent correction for site LR {lr['lr_ref']}", delta, user=u,
            )
    if "hamali" in changes:
        old_hamali = _decimal_value(lr.get("hamali"))
        new_hamali = _decimal(changes["hamali"], "Hamali", allow_none=True)
        charge_values_changed = charge_values_changed or new_hamali != old_hamali
        old_values["hamali"] = format(old_hamali, ".2f") if lr.get("hamali") is not None else None
        new_values["hamali"] = format(new_hamali, ".2f") if new_hamali is not None else None
        changes["hamali"] = Decimal128(new_hamali) if new_hamali is not None else None
    if "containers" not in changes and financial_fields.intersection(changes):
        goods = lr.get("goods_rows") or lr.get("containers") or []
        if goods:
            rows_with_charges = []
            for index, line in enumerate(goods):
                rent_value = _decimal_value(line.get("rent"))
                hamali_value = _decimal_value(line.get("hamali"))
                if index == 0:
                    if "rent" in changes:
                        rent_value += _decimal_value(changes["rent"]) - _decimal_value(lr.get("rent"))
                    if "hamali" in changes:
                        hamali_value += _decimal_value(changes["hamali"]) - _decimal_value(lr.get("hamali"))
                rows_with_charges.append({
                    **line, "rent": Decimal128(rent_value), "hamali": Decimal128(hamali_value),
                })
            changes["containers"] = rows_with_charges
            changes["goods_rows"] = rows_with_charges
    if charge_values_changed:
        changes.update(_charge_editor_metadata(u))
    changes["updated_at"] = now_iso()
    update_query = _scoped(str(site["business_id"]), site_id, _id=lr_id, trip_id=trip_id)
    if expected_updated_at:
        update_query["updated_at"] = expected_updated_at
    updated_result = await db.site_lrs.update_one(
        update_query,
        {"$set": changes},
    )
    if expected_updated_at and not updated_result.matched_count:
        raise HTTPException(409, "This receipt changed elsewhere. Reload it before saving your edits.")
    await _audit(
        site, u, "lr.updated", "lr", lr_id, old=old_values, new=new_values,
    )
    return (await _lr_rows(site, trip_id, [await _lr(site, trip_id, lr_id)], u))[0]


@router.post("/sites/{site_id}/trips/{trip_id}/lrs/{lr_id}/void")
async def void_site_lr(
    site_id: str, trip_id: str, lr_id: str, body: VoidReceipt, u=Depends(site_user),
):
    site = await _site_for_user(site_id, u, "lrs:update")
    trip = await _trip(site, trip_id)
    _ensure_trip_ledger_open(trip)
    is_owner = u.get("role") == "owner"
    if trip.get("status") != "open" and not is_owner:
        raise HTTPException(409, "Reopen the trip before voiding a receipt")
    lr = await _lr(site, trip_id, lr_id)
    if lr.get("voided"):
        return {"id": lr_id, "voided": True}
    reason = _normalize_label(body.reason, "Void reason", 500)
    has_financial_history = (
        lr.get("reconciled") or _decimal_value(lr.get("rent"))
        or _decimal_value(lr.get("hamali"))
    )
    if has_financial_history and not is_owner:
        raise HTTPException(403, "Only the business owner can void a receipt with charges or reconciliation history")
    if (has_financial_history or trip.get("status") != "open") and is_owner and len(reason) < 5:
        raise HTTPException(422, "Enter a detailed reason for this owner correction")
    void_cycle = int(lr.get("void_cycle", 0)) + 1
    paid = (await _payments_by_lr(site, trip_id, [lr_id])).get(lr_id, Decimal("0.00"))
    if is_owner and paid:
        payments = await db.site_payments.find(_scoped(
            str(site["business_id"]), site_id, trip_id=trip_id, lr_id=lr_id,
            kind="payment", posting_status="posted",
        )).sort("created_at", 1).to_list(1000)
        for payment in payments:
            reversal = await db.site_payments.find_one(_scoped(
                str(site["business_id"]), site_id, reversal_of=payment["_id"], kind="reversal",
            ))
            if reversal:
                continue
            await _record_lr_payment(
                site, trip, lr, _decimal_value(payment["amount"]), payment["date"],
                payment["method"], payment.get("reference", ""),
                f"void-reversal-{payment['_id']}", u, kind="reversal",
                reversal_of=payment["_id"], reason=reason,
                source="receipt_void", void_cycle=void_cycle,
            )
    rent = _decimal_value(lr.get("rent"))
    if is_owner and rent:
        await _financial_event(
            f"site-lr-void-rent:{lr_id}:{void_cycle}", site, lr["party_id"], lr["operating_date"],
            f"Owner void correction for site LR {lr['lr_ref']}", -rent, user=u,
        )
    await db.site_lrs.update_one(
        _scoped(str(site["business_id"]), site_id, _id=lr_id, trip_id=trip_id, voided={"$ne": True}),
        {"$set": {"voided": True, "voided_by": u["username"], "voided_at": now_iso(),
                  "void_reason": reason, "void_cycle": void_cycle, "updated_at": now_iso()}},
    )
    await _audit(site, u, "lr.voided", "lr", lr_id,
                 old={"voided": False}, new={"voided": True}, reason=reason)
    return {"id": lr_id, "voided": True}


@router.post("/sites/{site_id}/trips/{trip_id}/lrs/{lr_id}/unvoid")
async def unvoid_site_lr(
    site_id: str, trip_id: str, lr_id: str, body: RestoreReceipt, u=Depends(site_user),
):
    _owner(u)
    site = await _site_for_user(site_id, u)
    trip = await _trip(site, trip_id)
    _ensure_trip_ledger_open(trip)
    lr = await _lr(site, trip_id, lr_id)
    if not lr.get("voided"):
        return {"id": lr_id, "voided": False}
    reason = _normalize_label(body.reason, "Restore reason", 500)
    cycle = int(lr.get("void_cycle", 0))
    restore_key = body.idempotency_key
    if cycle:
        reversals = await db.site_payments.find(_scoped(
            str(site["business_id"]), site_id, trip_id=trip_id, lr_id=lr_id,
            kind="reversal", posting_status="posted", source="receipt_void",
            void_cycle=cycle,
        )).sort("created_at", 1).to_list(1000)
    else:
        reversals = await db.site_payments.find(_scoped(
            str(site["business_id"]), site_id, trip_id=trip_id, lr_id=lr_id,
            kind="reversal", posting_status="posted", reason=lr.get("void_reason", ""),
        )).sort("created_at", 1).to_list(1000)
    for reversal in reversals:
        await _record_lr_payment(
            site, trip, lr, _decimal_value(reversal["amount"]), reversal["date"],
            reversal["method"], reversal.get("reference", ""),
            f"unvoid-{restore_key}-{reversal['_id']}", u,
            reason=reason, source="receipt_unvoid", void_cycle=cycle,
        )
    rent = _decimal_value(lr.get("rent"))
    if rent:
        await _financial_event(
            f"site-lr-unvoid-rent:{lr_id}:{cycle}:{restore_key}", site,
            lr["party_id"], lr["operating_date"],
            f"Owner restore correction for site LR {lr['lr_ref']}", rent, user=u,
        )
    now = now_iso()
    update = {
        "voided": False, "unvoided_by": u["username"], "unvoided_at": now,
        "unvoid_reason": reason, "updated_at": now,
    }
    await db.site_lrs.update_one(
        _scoped(str(site["business_id"]), site_id, _id=lr_id, trip_id=trip_id, voided=True),
        {"$set": update},
    )
    await _audit(
        site, u, "lr.unvoided", "lr", lr_id,
        old={"voided": True, "void_reason": lr.get("void_reason", "")},
        new={"voided": False, "reason": reason}, reason=reason,
        event_id=f"site-lr-unvoid-audit:{lr_id}:{restore_key}",
    )
    return {"id": lr_id, "voided": False, "restored_payments": len(reversals)}


async def _ensure_payment_posting(payment, site):
    payment_id = payment["_id"]
    await db.site_payments.update_one(
        {"_id": payment_id}, {"$setOnInsert": {"posting_status": "pending"}}, upsert=True,
    )
    current = await db.site_payments.find_one({"_id": payment_id})
    if not current:
        raise HTTPException(500, "Could not load the staged payment posting")
    if current.get("posting_status") == "posted":
        return
    amount = _decimal_value(payment["amount"])
    reversal = payment.get("kind") == "reversal"
    ledger_exists = await db.ledger.find_one(
        {"ref_type": "site_lr_payment", "ref_id": payment_id}
    )
    if not ledger_exists:
        try:
            await L.post(
                "party", payment["party_id"], payment["date"],
                ("Reversal: " if reversal else "") + f"LR payment {payment['lr_ref']}",
                debit=float(amount) if reversal else 0,
                credit=0 if reversal else float(amount),
                ref_type="site_lr_payment", ref_id=payment_id,
                site_id=site["_id"], business_id=str(site["business_id"]),
            )
        except DuplicateKeyError:
            if not await db.ledger.find_one({
                "business_id": str(site["business_id"]), "site_id": site["_id"],
                "ref_type": "site_lr_payment", "ref_id": payment_id,
            }):
                raise
    cash_exists = await db.cashbook.find_one(
        {"ref_type": "site_lr_payment", "ref_id": payment_id}
    )
    if not cash_exists:
        try:
            await L.cash(
                "cash" if payment["method"] in ("Cash", "Other") else "bank",
                "out" if reversal else "in", float(amount), payment["date"],
                ("Reversal: " if reversal else "") + f"LR payment {payment['lr_ref']}",
                "site_lr_payment", payment_id, mode=payment["method"],
                site_id=site["_id"], business_id=str(site["business_id"]),
            )
        except DuplicateKeyError:
            if not await db.cashbook.find_one({
                "business_id": str(site["business_id"]), "site_id": site["_id"],
                "ref_type": "site_lr_payment", "ref_id": payment_id,
            }):
                raise
    await db.site_payments.update_one(
        {"_id": payment_id}, {"$set": {"posting_status": "posted", "posted_at": now_iso()}},
    )


async def _apply_payment_balance(payment, site):
    payment_id = payment["_id"]
    amount = _decimal_value(payment["amount"])
    reversal = payment.get("kind") == "reversal"
    query = _scoped(
        str(site["business_id"]), site["_id"], _id=payment["lr_id"],
        trip_id=payment["trip_id"], payment_effect_ids={"$ne": payment_id},
    )
    if reversal:
        query["$expr"] = {"$gte": [
            {"$ifNull": ["$paid_total", Decimal128(Decimal("0.00"))]},
            Decimal128(amount),
        ]}
        update = {"$inc": {"paid_total": Decimal128(-amount)},
                  "$addToSet": {"payment_effect_ids": payment_id}}
    else:
        query["$expr"] = {"$lte": [
            {"$add": [{"$ifNull": ["$paid_total", Decimal128(Decimal("0.00"))]},
                      Decimal128(amount)]},
            {"$add": [
                "$rent",
                {"$ifNull": ["$receipt_fee", Decimal128(Decimal("2.00"))]},
            ]},
        ]}
        update = {"$inc": {"paid_total": Decimal128(amount)},
                  "$addToSet": {"payment_effect_ids": payment_id}}
    result = await db.site_lrs.update_one(query, update)
    if result.modified_count == 1:
        return
    lr = await db.site_lrs.find_one(_scoped(
        str(site["business_id"]), site["_id"], _id=payment["lr_id"],
        trip_id=payment["trip_id"],
    ))
    if lr and payment_id in (lr.get("payment_effect_ids") or []):
        return
    reason = "Payment would make the LR balance invalid"
    await db.site_payments.update_one(
        {"_id": payment_id, "posting_status": "pending"},
        {"$set": {"posting_status": "rejected", "rejection_reason": reason,
                  "rejected_at": now_iso()}},
    )
    raise HTTPException(409, reason)


async def _record_lr_payment(site, trip, lr, amount, payment_date, method,
                             reference, idempotency_key, user, kind="payment",
                             reversal_of=None, reason=None, source=None, void_cycle=None):
    business_id = str(site["business_id"])
    existing = await db.site_payments.find_one(
        _scoped(business_id, site["_id"], idempotency_key=idempotency_key)
    )
    if existing:
        if (_decimal_value(existing["amount"]) != amount or existing.get("lr_id") != lr["_id"]
                or existing.get("kind") != kind or existing.get("date") != payment_date
                or existing.get("method") != method
                or existing.get("reference", "") != reference
                or existing.get("reversal_of") != reversal_of):
            raise HTTPException(409, "Idempotency key was already used for a different payment")
        if existing.get("posting_status") == "rejected":
            raise HTTPException(409, "This payment request was previously rejected; submit a new payment")
        await _apply_payment_balance(existing, site)
        await _ensure_payment_posting(existing, site)
        await _audit(
            site, user,
            "lr.payment_recorded" if kind == "payment" else "lr.payment_reversed",
            "payment", existing["_id"],
            new={"lr_id": lr["_id"], "amount": format(amount, ".2f"),
                 "method": method, "reference": reference,
                 "reversal_of": reversal_of, "reason": reason or ""},
            event_id=f"site-payment-audit:{existing['_id']}",
        )
        return existing
    if kind == "payment":
        if lr.get("rent") is None:
            raise HTTPException(409, "Enter bhada before recording a payment")
        if amount <= 0:
            raise HTTPException(422, "Payment must be positive")
    payment_id = new_id()
    payment = {
        "_id": payment_id, "business_id": business_id, "site_id": site["_id"],
        "trip_id": trip["_id"], "lr_id": lr["_id"], "lr_ref": lr["lr_ref"],
        "party_id": lr["party_id"], "operating_date": lr["operating_date"],
        "amount": Decimal128(amount), "date": payment_date, "method": method,
        "reference": reference, "kind": kind,
        "reason": reason or "", "idempotency_key": idempotency_key,
        "posting_status": "pending", **_creator_metadata(user),
    }
    if reversal_of:
        payment["reversal_of"] = reversal_of
    if source:
        payment["source"] = source
    if void_cycle is not None:
        payment["void_cycle"] = void_cycle
    try:
        await db.site_payments.insert_one(payment)
    except DuplicateKeyError:
        if reversal_of:
            duplicate_reversal = await db.site_payments.find_one(
                _scoped(business_id, site["_id"], reversal_of=reversal_of, kind="reversal")
            )
            if duplicate_reversal:
                raise HTTPException(409, "This payment has already been reversed")
        existing = await db.site_payments.find_one(
            _scoped(business_id, site["_id"], idempotency_key=idempotency_key)
        )
        if existing:
            return await _record_lr_payment(
                site, trip, lr, amount, payment_date, method, reference,
                idempotency_key, user, kind, reversal_of, reason,
            )
        raise
    await _apply_payment_balance(payment, site)
    await _ensure_payment_posting(payment, site)
    await _audit(site, user, "lr.payment_recorded" if kind == "payment" else "lr.payment_reversed",
                 "payment", payment_id, new={"lr_id": lr["_id"], "amount": format(amount, ".2f"),
                                             "method": method, "reference": reference,
                                             "reversal_of": reversal_of, "reason": reason or ""},
                 event_id=f"site-payment-audit:{payment_id}")
    return payment


@router.post("/sites/{site_id}/trips/{trip_id}/lrs/{lr_id}/payments")
async def record_site_payment(
    site_id: str, trip_id: str, lr_id: str, body: PaymentCreate,
    u=Depends(site_user),
):
    site = await _site_for_user(site_id, u, "payments:create")
    trip = await _trip(site, trip_id)
    _ensure_trip_ledger_open(trip)
    lr = await _lr(site, trip_id, lr_id)
    payment_date = _validate_iso_date(body.date).isoformat()
    method = body.method.strip()
    if method not in ("Cash", "UPI", "Bank", "Cheque", "Other"):
        raise HTTPException(422, "Choose a supported payment method")
    amount = _decimal(body.amount, "Payment")
    result = await _record_lr_payment(site, trip, lr, amount, payment_date, method,
                                      body.reference.strip(), body.idempotency_key, u)
    return _doc_for_user(result, u)


@router.post("/sites/{site_id}/trips/{trip_id}/lrs/{lr_id}/settlement")
async def update_site_lr_settlement(
    site_id: str, trip_id: str, lr_id: str,
    body: BookingSettlementUpdate, u=Depends(site_user),
):
    _owner(u)
    site = await _site_for_user(site_id, u)
    trip = await _trip(site, trip_id)
    _ensure_trip_ledger_open(trip)
    lr = await _lr(site, trip_id, lr_id)
    if lr.get("voided"):
        raise HTTPException(409, "Restore this receipt before recording or reversing payment")
    reason = " ".join(body.reason.split()) or "Owner updated payment received status in the booking ledger"
    if body.received:
        if lr.get("rent") is None:
            raise HTTPException(409, "Enter Bhada before marking this receipt as paid")
        paid = (await _payments_by_lr(site, trip_id, [lr_id])).get(lr_id, Decimal("0.00"))
        amount_due = _decimal_value(lr.get("rent")) + _decimal_value(
            lr.get("receipt_fee", Decimal128(Decimal("2.00"))),
        )
        outstanding = max(amount_due - paid, Decimal("0.00"))
        if outstanding:
            payment_date = _now_at_site(site)[0]
            payment = await _record_lr_payment(
                site, trip, lr, outstanding, payment_date, "Cash",
                "Booking ledger marked received",
                f"booking-settlement-{body.idempotency_key}", u,
                reason=reason, source="booking_settlement",
            )
            await db.site_payments.update_one(
                {"_id": payment["_id"]},
                {"$set": {"source": "booking_settlement", "settlement_reason": reason}},
            )
        else:
            payment = None
        await db.site_lrs.update_one(
            _scoped(str(site["business_id"]), site_id, _id=lr_id, trip_id=trip_id),
            {"$set": {"booking_ledger_paid": True, "booking_ledger_paid_by": u["username"],
                      "booking_ledger_paid_at": now_iso()}},
        )
        result = {
            "id": lr_id, "received": True,
            "paid_total": format(amount_due, ".2f"),
            "outstanding": "0.00",
            "settlement_payment_id": payment["_id"] if payment else None,
        }
        await _audit(
            site, u, "lr.settlement_marked_received", "lr", lr_id,
            new={"received": True, "amount": format(outstanding, ".2f"),
                 "payment_id": payment["_id"] if payment else ""},
            reason=reason,
            event_id=f"site-lr-settlement-audit:{lr_id}:{body.idempotency_key}",
        )
        return result

    settlements = await db.site_payments.find(_scoped(
        str(site["business_id"]), site_id, trip_id=trip_id, lr_id=lr_id,
        kind="payment", source="booking_settlement", posting_status="posted",
    )).sort("created_at", -1).to_list(1000)
    reversed_count = 0
    for payment in settlements:
        reversal = await db.site_payments.find_one(_scoped(
            str(site["business_id"]), site_id, reversal_of=payment["_id"], kind="reversal",
        ))
        if reversal:
            continue
        await _record_lr_payment(
            site, trip, lr, _decimal_value(payment["amount"]), payment["date"],
            payment["method"], payment.get("reference", ""),
            f"booking-settlement-reversal-{payment['_id']}", u,
            kind="reversal", reversal_of=payment["_id"], reason=reason,
            source="booking_settlement_undo",
        )
        reversed_count += 1
    if reversed_count:
        await _audit(
            site, u, "lr.settlement_reversed", "lr", lr_id,
            new={"received": False, "reversed_settlements": reversed_count},
            reason=reason,
            event_id=f"site-lr-settlement-undo-audit:{lr_id}:{body.idempotency_key}",
        )
    await db.site_lrs.update_one(
        _scoped(str(site["business_id"]), site_id, _id=lr_id, trip_id=trip_id),
        {"$set": {"booking_ledger_paid": False, "booking_ledger_paid_by": u["username"],
                  "booking_ledger_paid_at": now_iso()}},
    )
    paid = (await _payments_by_lr(site, trip_id, [lr_id])).get(lr_id, Decimal("0.00"))
    rent = _decimal_value(lr.get("rent")) + _decimal_value(
        lr.get("receipt_fee", Decimal128(Decimal("2.00"))),
    )
    return {
        "id": lr_id, "received": False,
        "paid_total": format(paid, ".2f"),
        "outstanding": format(max(rent - paid, Decimal("0.00")), ".2f"),
        "reversed_settlements": reversed_count,
    }


@router.post("/sites/{site_id}/trips/{trip_id}/ledger/complete")
async def complete_site_trip_ledger(
    site_id: str, trip_id: str, u=Depends(site_user),
):
    _owner(u)
    site = await _site_for_user(site_id, u)
    trip = await _trip(site, trip_id)
    if trip.get("ledger_completed_at"):
        return {
            "id": trip_id, "ledger_completed": True,
            "ledger_completed_at": trip["ledger_completed_at"],
            "ledger_completed_by": trip.get("ledger_completed_by", ""),
        }

    business_id = str(site["business_id"])
    lrs = await db.site_lrs.find(_scoped(
        business_id, site_id, trip_id=trip_id,
    )).sort("sequence", 1).limit(10001).to_list(10001)
    if len(lrs) > 10000:
        raise HTTPException(413, "This trip has more than 10,000 LRs; complete a smaller trip")
    unpriced = [
        lr["lr_ref"] for lr in lrs
        if not lr.get("voided") and lr.get("rent") is None
    ]
    if unpriced:
        preview = ", ".join(unpriced[:10])
        suffix = "…" if len(unpriced) > 10 else ""
        raise HTTPException(
            409, f"Enter Bhada for every active receipt before completing the ledger: {preview}{suffix}",
        )

    if not trip.get("ledger_completion_started_at"):
        started_at = now_iso()
        completion_date = _now_at_site(site)[0]
        await db.site_trips.update_one(
            {
                **_scoped(business_id, site_id, _id=trip_id),
                "ledger_completion_started_at": {"$exists": False},
                "ledger_completed_at": {"$exists": False},
            },
            {"$set": {
                "ledger_completion_started_at": started_at,
                "ledger_completion_date": completion_date,
                "ledger_completion_started_by": u["username"],
            }},
        )
        trip = await _trip(site, trip_id)
        if trip.get("ledger_completed_at"):
            return {
                "id": trip_id, "ledger_completed": True,
                "ledger_completed_at": trip["ledger_completed_at"],
                "ledger_completed_by": trip.get("ledger_completed_by", ""),
            }

    completion_date = trip.get("ledger_completion_date") or _now_at_site(site)[0]
    active_lrs = [lr for lr in lrs if not lr.get("voided")]
    paid_by_lr = await _payments_by_lr(
        site, trip_id, [lr["_id"] for lr in active_lrs],
    )
    for lr in active_lrs:
        rent = _decimal_value(lr.get("rent"))
        receipt_fee = _decimal_value(lr.get("receipt_fee", Decimal128(Decimal("2.00"))))
        outstanding = max(
            rent + receipt_fee - paid_by_lr.get(lr["_id"], Decimal("0.00")),
            Decimal("0.00"),
        )
        if outstanding:
            payment = await _record_lr_payment(
                site, trip, lr, outstanding, completion_date, "Cash",
                "Trip ledger completed",
                f"trip-ledger-completion:{trip_id}:{lr['_id']}", u,
                reason="Admin completed trip ledger", source="booking_settlement",
            )
            await db.site_payments.update_one(
                {"_id": payment["_id"]},
                {"$set": {
                    "source": "booking_settlement",
                    "settlement_reason": "Admin completed trip ledger",
                }},
            )
        await db.site_lrs.update_one(
            _scoped(business_id, site_id, _id=lr["_id"], trip_id=trip_id),
            {"$set": {
                "booking_ledger_paid": True,
                "booking_ledger_paid_by": u["username"],
                "booking_ledger_paid_at": now_iso(),
            }},
        )

    completed_at = now_iso()
    await db.site_trips.update_one(
        _scoped(business_id, site_id, _id=trip_id),
        {"$set": {
            "ledger_completed_at": completed_at,
            "ledger_completed_by": u["username"],
            "ledger_completion_started_by": u["username"],
        }},
        )
    await _audit(
        site, u, "trip.ledger_completed", "trip", trip_id,
        new={"ledger_completed": True, "receipt_count": len(active_lrs)},
        event_id=f"trip-ledger-completed:{trip_id}",
    )
    return {
        "id": trip_id, "ledger_completed": True,
        "ledger_completed_at": completed_at,
        "ledger_completed_by": u["username"],
    }


@router.post("/sites/{site_id}/trips/{trip_id}/lrs/{lr_id}/payments/{payment_id}/reverse")
async def reverse_site_payment(
    site_id: str, trip_id: str, lr_id: str, payment_id: str,
    body: PaymentReverse, u=Depends(site_user),
):
    _owner(u)
    site = await _site_for_user(site_id, u)
    trip = await _trip(site, trip_id)
    _ensure_trip_ledger_open(trip)
    lr = await _lr(site, trip_id, lr_id)
    original = await db.site_payments.find_one(_scoped(
        str(site["business_id"]), site_id, _id=payment_id, trip_id=trip_id,
        lr_id=lr_id, kind="payment",
    ))
    if not original:
        raise HTTPException(404, "Payment not found")
    if original.get("posting_status") == "rejected":
        raise HTTPException(409, "A rejected payment cannot be reversed")
    if original.get("posting_status") != "posted":
        original = await _record_lr_payment(
            site, trip, lr, _decimal_value(original["amount"]), original["date"],
            original["method"], original.get("reference", ""),
            original["idempotency_key"], u,
        )
    prior_reversal = await db.site_payments.find_one(_scoped(
        str(site["business_id"]), site_id, reversal_of=payment_id, kind="reversal",
    ))
    if prior_reversal:
        return _doc(await _record_lr_payment(
            site, trip, lr, _decimal_value(original["amount"]), original["date"],
            original["method"], original.get("reference", ""),
            f"reverse-{payment_id}", u, kind="reversal", reversal_of=payment_id,
            reason=prior_reversal.get("reason") or _normalize_label(body.reason, "Reversal reason", 500),
        ))
    reversal = await _record_lr_payment(
        site, trip, lr, _decimal_value(original["amount"]), original["date"],
        original["method"], original.get("reference", ""),
        f"reverse-{payment_id}", u, kind="reversal", reversal_of=payment_id,
        reason=_normalize_label(body.reason, "Reversal reason", 500),
    )
    return _doc(reversal)


@router.get("/sites/{site_id}/trips/{trip_id}/payments")
async def list_site_payments(site_id: str, trip_id: str, lr_id: Optional[str] = None,
                             u=Depends(site_user)):
    site = await _site_for_user(site_id, u, "payments:read")
    await _trip(site, trip_id)
    query = _scoped(str(site["business_id"]), site_id, trip_id=trip_id)
    if lr_id:
        await _lr(site, trip_id, lr_id)
        query["lr_id"] = lr_id
    payments = await db.site_payments.find(query).sort("created_at", -1).limit(1000).to_list(1000)
    return [_doc_for_user(payment, u) for payment in payments]


@router.post("/sites/{site_id}/trips/{trip_id}/lrs/{lr_id}/payments/{payment_id}/retry")
async def retry_site_payment_posting(
    site_id: str, trip_id: str, lr_id: str, payment_id: str,
    u=Depends(site_user),
):
    _owner(u)
    site = await _site_for_user(site_id, u)
    trip = await _trip(site, trip_id)
    _ensure_trip_ledger_open(trip)
    lr = await _lr(site, trip_id, lr_id)
    payment = await db.site_payments.find_one(_scoped(
        str(site["business_id"]), site_id, _id=payment_id,
        trip_id=trip_id, lr_id=lr_id,
    ))
    if not payment:
        raise HTTPException(404, "Payment event not found")
    if payment.get("posting_status") == "rejected":
        raise HTTPException(409, "Rejected payment events cannot be retried; record a corrected payment")
    result = await _record_lr_payment(
        site, trip, lr, _decimal_value(payment["amount"]), payment["date"],
        payment["method"], payment.get("reference", ""), payment["idempotency_key"],
        u, kind=payment.get("kind", "payment"),
        reversal_of=payment.get("reversal_of"), reason=payment.get("reason"),
    )
    return _doc_for_user(result, u)


@router.get("/sites/{site_id}/audit")
async def list_site_audit(site_id: str, limit: int = Query(default=50, ge=1, le=100),
                          offset: int = Query(default=0, ge=0), u=Depends(site_user)):
    _owner(u)
    site = await _site_for_user(site_id, u, "dashboard:read")
    query = _scoped(str(site["business_id"]), site_id)
    total = await db.site_audit_events.count_documents(query)
    rows = await db.site_audit_events.find(query).sort("created_at", -1).skip(offset).limit(limit).to_list(limit)
    visible = _audit_rows_for_user(rows, u, site_id)
    return {"rows": visible, "total": total,
            "limit": limit, "offset": offset}


@router.get("/sites/{site_id}/booking-finance")
async def site_booking_finance(site_id: str, u=Depends(site_user)):
    _owner(u)
    site = await _site_for_user(site_id, u, "dashboard:read")
    business_id = str(site["business_id"])
    lrs = await db.site_lrs.find(_scoped(
        business_id, site_id, voided={"$ne": True},
    )).sort([("operating_date", -1), ("sequence", -1)]).limit(10001).to_list(10001)
    if len(lrs) > 10000:
        raise HTTPException(413, "This booking finance view has more than 10,000 receipts")

    paid_by_lr = {}
    async for payment in db.site_payments.aggregate([
        {"$match": _scoped(
            business_id, site_id, posting_status="posted",
            kind={"$in": ["payment", "reversal"]},
        )},
        {"$group": {
            "_id": "$lr_id",
            "paid": {"$sum": {"$cond": [
                {"$eq": ["$kind", "reversal"]},
                {"$multiply": ["$amount", -1]}, "$amount",
            ]}},
        }},
    ]):
        paid_by_lr[payment["_id"]] = _decimal_value(payment.get("paid"))

    trips = {}
    for trip_id in {lr["trip_id"] for lr in lrs}:
        trip = await db.site_trips.find_one(_scoped(business_id, site_id, _id=trip_id))
        if trip:
            trips[trip_id] = trip

    totals = {
        "receipt_count": 0, "pending_orders": 0, "unpriced_receipts": 0,
        "unpaid_receipts": 0, "recorded_bhada": Decimal("0.00"),
        "recorded_receipt_fees": Decimal("0.00"),
        "recorded_hamali": Decimal("0.00"), "grand_total": Decimal("0.00"),
        "collected_bhada": Decimal("0.00"), "outstanding_bhada": Decimal("0.00"),
        "collectible_outstanding_bhada": Decimal("0.00"),
    }
    rows = []
    receiver_totals = {}
    today = date.fromisoformat(_now_at_site(site)[0])
    for lr in lrs:
        trip = trips.get(lr["trip_id"])
        if not trip:
            continue
        rent = _decimal_value(lr.get("rent"))
        hamali = _decimal_value(lr.get("hamali"))
        receipt_fee = _decimal_value(lr.get("receipt_fee", Decimal128(Decimal("2.00"))))
        paid = paid_by_lr.get(lr["_id"], Decimal("0.00"))
        amount_due = rent + receipt_fee if lr.get("rent") is not None else Decimal("0.00")
        outstanding = max(amount_due - paid, Decimal("0.00"))
        reconciled = lr.get("reconciled") is True
        receipt_day = lr.get("receipt_date", lr.get("operating_date", ""))
        age_days = max((today - _validate_iso_date(receipt_day)).days, 0)
        goods = lr.get("goods_rows") or lr.get("containers") or []
        totals["receipt_count"] += 1
        totals["recorded_bhada"] += rent
        totals["recorded_receipt_fees"] += receipt_fee
        totals["recorded_hamali"] += hamali
        totals["grand_total"] += rent + hamali + receipt_fee
        totals["collected_bhada"] += paid
        totals["outstanding_bhada"] += outstanding
        totals["pending_orders"] += int(not reconciled)
        totals["unpriced_receipts"] += int(lr.get("rent") is None)
        totals["unpaid_receipts"] += int(outstanding > 0)
        if reconciled:
            totals["collectible_outstanding_bhada"] += outstanding
        receiver_key = str(lr.get("party_id") or lr.get("receiver_identity_id")
                           or lr.get("receiver_name") or lr["_id"])
        receiver = receiver_totals.setdefault(receiver_key, {
            "receiver_name": lr.get("receiver_name", ""),
            "receiver_label": lr.get("receiver_label", lr.get("receiver_name", "")),
            "receipt_count": 0, "charges": Decimal("0.00"),
            "collected": Decimal("0.00"), "outstanding": Decimal("0.00"),
        })
        receiver["receipt_count"] += 1
        receiver["charges"] += rent + receipt_fee
        receiver["collected"] += paid
        receiver["outstanding"] += outstanding
        rows.append({
            "id": lr["_id"], "site_id": site_id, "trip_id": lr["trip_id"],
            "trip_ref": trip.get("trip_ref", ""), "lr_ref": lr.get("lr_ref", ""),
            "receipt_date": receipt_day,
            "sender_name": lr.get("sender_name", ""),
            "receiver_name": lr.get("receiver_name", ""),
            "receiver_label": lr.get("receiver_label", lr.get("receiver_name", "")),
            "goods": "; ".join(
                f"{line.get('type', '')} x {line.get('quantity', 0)}" for line in goods
            ),
            "rent": format(rent, ".2f") if lr.get("rent") is not None else None,
            "hamali": format(hamali, ".2f"),
            "receipt_fee": format(receipt_fee, ".2f"),
            "total": format(rent + hamali + receipt_fee, ".2f"),
            "paid": format(paid, ".2f"),
            "outstanding": format(outstanding, ".2f"),
            "payment_status": _payment_status(
                amount_due if lr.get("rent") is not None else None, paid,
            ),
            "reconciled": reconciled, "age_days": age_days,
            "promised_date": lr.get("payment_promised_date", ""),
            "followup_note": lr.get("payment_followup_note", ""),
            "next_contact_date": lr.get("payment_next_contact_date", ""),
            "followup_outcome": lr.get("payment_followup_outcome", ""),
            "created_by_login": lr.get("created_by_login", lr.get("created_by", "")),
            "created_by_role": lr.get("created_by_role", ""),
        })
    money_keys = (
        "recorded_bhada", "recorded_receipt_fees", "recorded_hamali",
        "grand_total", "collected_bhada",
        "outstanding_bhada", "collectible_outstanding_bhada",
    )
    formatted = {
        **totals,
        **{key: format(totals[key], ".2f") for key in money_keys},
    }
    receiver_balances = [
        {
            **{key: value for key, value in receiver.items()
               if key not in {"charges", "collected", "outstanding"}},
            "charges": format(receiver["charges"], ".2f"),
            "collected": format(receiver["collected"], ".2f"),
            "outstanding": format(receiver["outstanding"], ".2f"),
        }
        for receiver in sorted(
            receiver_totals.values(),
            key=lambda value: (-value["outstanding"], value["receiver_name"].casefold()),
        )
    ]
    return {
        "site": {"id": site_id, "name": site.get("name", ""), "code": site.get("code", "")},
        "overdue_after_days": _receipt_controls(site).overdue_after_days,
        "totals": formatted,
        "receiver_balances": receiver_balances,
        "basis": (
            "Grand total is active receipt Bhada, Hamali and receipt fees. Outstanding and "
            "collections track Bhada plus receipt fees; unreconciled orders are listed separately. "
            "Payments include only "
            "successfully posted payments and reversals."
        ),
        "rows": rows,
    }


@router.patch("/sites/{site_id}/trips/{trip_id}/lrs/{lr_id}/followup")
async def update_site_lr_followup(
    site_id: str, trip_id: str, lr_id: str,
    body: BookingFollowupUpdate, u=Depends(site_user),
):
    _owner(u)
    site = await _site_for_user(site_id, u, "dashboard:read")
    trip = await _trip(site, trip_id)
    _ensure_trip_ledger_open(trip)
    lr = await _lr(site, trip_id, lr_id)
    promised_date = (
        _validate_iso_date(body.promised_date).isoformat()
        if body.promised_date else ""
    )
    next_contact_date = (
        _validate_iso_date(body.next_contact_date).isoformat()
        if body.next_contact_date else ""
    )
    note = " ".join(body.note.split())[:500]
    outcome = " ".join(body.outcome.split())[:500]
    old = {
        "payment_promised_date": lr.get("payment_promised_date", ""),
        "payment_followup_note": lr.get("payment_followup_note", ""),
        "payment_next_contact_date": lr.get("payment_next_contact_date", ""),
        "payment_followup_outcome": lr.get("payment_followup_outcome", ""),
    }
    new = {
        "payment_promised_date": promised_date,
        "payment_followup_note": note,
        "payment_next_contact_date": next_contact_date,
        "payment_followup_outcome": outcome,
    }
    updated_at = now_iso()
    await db.site_lrs.update_one(
        _scoped(str(site["business_id"]), site_id, _id=lr["_id"], trip_id=trip_id),
        {"$set": {**new, "updated_at": updated_at}},
    )
    await _audit(
        site, u, "lr.payment_followup_updated", "lr", lr["_id"],
        old=old, new=new,
    )
    return {
        "id": lr["_id"], "promised_date": promised_date, "note": note,
        "next_contact_date": next_contact_date, "outcome": outcome,
        "updated_at": updated_at,
    }


LEGACY_COLLECTIONS = (
    "trips", "lrs", "receipts", "handovers", "expenses", "fuel",
    "payments", "advances", "tpl", "ledger", "cashbook",
)


@router.get("/sites/migration/legacy")
async def legacy_migration_preview(site_id: str, u=Depends(site_user)):
    _owner(u)
    site = await _site_for_user(site_id, u)
    counts = {}
    for collection in LEGACY_COLLECTIONS:
        counts[collection] = await db[collection].count_documents({"site_id": None})
    return {"mode": "dry_run", "site": _doc(site), "unassigned": counts,
            "total": sum(counts.values()), "applied": False}


@router.post("/sites/migration/legacy")
async def apply_legacy_migration(body: LegacyMigration, u=Depends(site_user)):
    _owner(u)
    site = await _site_for_user(body.site_id, u)
    if not site.get("is_default"):
        raise HTTPException(409, "Choose the business default site before assigning legacy records")
    if body.confirm is not True:
        raise HTTPException(400, "Explicit confirmation is required to assign legacy records")
    business_id = str(site["business_id"])
    migration_id = f"legacy-default-site-{site['_id']}"
    previous = await db.site_migrations.find_one({"_id": migration_id})
    if previous and previous.get("status") == "complete":
        return _doc(previous)
    counts = {}
    await db.site_migrations.update_one(
        {"_id": migration_id},
        {"$setOnInsert": {"business_id": business_id, "site_id": site["_id"],
                          "version": 1, "status": "applying", "created_by": u["username"],
                          "created_at": now_iso()}},
        upsert=True,
    )
    for collection in LEGACY_COLLECTIONS:
        result = await db[collection].update_many(
            {"site_id": None},
            {"$set": {"site_id": site["_id"], "business_id": business_id}},
        )
        counts[collection] = result.modified_count
    await db.site_migrations.update_one(
        {"_id": migration_id},
        {"$set": {"status": "complete", "counts": counts,
                  "completed_by": u["username"], "completed_at": now_iso()}},
    )
    await _audit(site, u, "legacy.records_assigned_to_default_site", "migration", migration_id,
                 new={"counts": counts, "mode": "explicit_apply"})
    return _doc(await db.site_migrations.find_one({"_id": migration_id}))


def _safe_csv_text(value):
    text = "" if value is None else str(value)
    if (text.startswith(("=", "+", "-", "@", "\t", "\r"))
            or text.lstrip(" \t\r\n\v\f").startswith(("=", "+", "-", "@"))):
        return "'" + text
    return text


def _csv_response(filename, business_id, site_id, trip_id, columns, rows):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\r\n")
    writer.writerow(["#FLEET_MANAGER_TRIP_LEDGER", LEDGER_TEMPLATE_VERSION])
    writer.writerow([f"business_id={business_id}", f"site_id={site_id}", f"trip_id={trip_id}"])
    writer.writerow(columns)
    for row in rows:
        writer.writerow([_safe_csv_text(value) if isinstance(value, str) else value
                        for value in row])
    return Response(
        content="\ufeff" + stream.getvalue(), media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"',
                 "Cache-Control": "no-store"},
    )


def _report_csv_response(filename, columns, rows):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\r\n")
    writer.writerow(columns)
    for row in rows:
        writer.writerow([_safe_csv_text(value) for value in row])
    return Response(
        content="\ufeff" + stream.getvalue(), media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"',
                 "Cache-Control": "no-store"},
    )


def _excel_safe(value):
    if isinstance(value, str) and value.lstrip(" \t\r\n\v\f").startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _xlsx_response(filename, site_name, date_range, sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, headers, rows, widths in sheets:
        sheet = workbook.create_sheet(title)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        heading = sheet.cell(1, 1, _excel_safe(site_name))
        heading.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        heading.fill = PatternFill("solid", fgColor="0B5C4E")
        heading.alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 28
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        subtitle = sheet.cell(2, 1, f"Ledger period: {date_range}")
        subtitle.font = Font(name="Calibri", size=11, italic=True, color="334155")
        header_row = 4
        for column, label in enumerate(headers, start=1):
            cell = sheet.cell(header_row, column, label)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="176B5A")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row_number, values in enumerate(rows, start=header_row + 1):
            for column, value in enumerate(values, start=1):
                if isinstance(value, Decimal):
                    value = float(value)
                cell = sheet.cell(row_number, column, _excel_safe(value))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if headers[column - 1] in {"Bhada", "Hamali", "Total", "Total rent"} and value != "":
                    cell.number_format = '"₹"#,##0.00'
            if row_number % 2 == 0:
                for cell in sheet[row_number]:
                    cell.fill = PatternFill("solid", fgColor="F1F7F5")
        sheet.freeze_panes = "A5"
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{header_row + len(rows)}"
        sheet.sheet_view.showGridLines = False
        sheet.column_dimensions["A"].width = 14
        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:4"
        sheet.print_area = f"A1:{get_column_letter(len(headers))}{header_row + len(rows)}"
        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.5
        sheet.page_margins.bottom = 0.5
    output = io.BytesIO()
    workbook.save(output)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"',
                 "Cache-Control": "no-store"},
    )


async def _trip_ledger_rows(site, trip):
    business_id = str(site["business_id"])
    lrs = await db.site_lrs.find(_scoped(business_id, site["_id"], trip_id=trip["_id"])).sort(
        "sequence", 1).limit(10001).to_list(10001)
    if len(lrs) > 10000:
        raise HTTPException(413, "This trip has more than 10,000 LRs; export a smaller trip")
    payments = await db.site_payments.find(_scoped(
        business_id, site["_id"], trip_id=trip["_id"],
    )).sort("created_at", 1).limit(50001).to_list(50001)
    if len(payments) > 50000:
        raise HTTPException(413, "This trip has more than 50,000 payment events; export a smaller trip")
    by_lr = {}
    for payment in payments:
        by_lr.setdefault(payment["lr_id"], []).append(payment)
    return lrs, by_lr


@router.get("/sites/{site_id}/trips/{trip_id}/ledger/{section}")
async def export_site_ledger(site_id: str, trip_id: str, section: str, u=Depends(site_user)):
    _owner(u)
    site = await _site_for_user(site_id, u)
    trip = await _trip(site, trip_id)
    if section not in (
        "goods", "goods-detail", "goods-details", "receivers", "summary", "import-template",
    ):
        raise HTTPException(404, "Unknown ledger export")
    business_id = str(site["business_id"])
    lrs, payments_by_lr = await _trip_ledger_rows(site, trip)
    lr_payment_totals = await _payments_by_lr(site, trip_id, [lr["_id"] for lr in lrs])
    if section == "goods":
        grouped = {}
        for lr in lrs:
            rent = _decimal_value(lr.get("rent"))
            hamali = _decimal_value(lr.get("hamali"))
            paid = lr_payment_totals.get(lr["_id"], Decimal("0.00"))
            key = lr["goods_type"]
            group = grouped.setdefault(key, {
                "quantity": 0, "lr_refs": set(), "receiver_ids": set(),
                "containers": {}, "rent": Decimal("0.00"),
                "hamali": Decimal("0.00"), "paid": Decimal("0.00"),
            })
            group["lr_refs"].add(lr["lr_ref"])
            group["receiver_ids"].add(lr["receiver_identity_id"])
            group["rent"] += rent
            group["hamali"] += hamali
            group["paid"] += paid
            for container in lr["containers"]:
                group["quantity"] += container["quantity"]
                group["containers"][container["type"]] = (
                    group["containers"].get(container["type"], 0) + container["quantity"]
                )
        columns = (
            "business_id", "site_id", "trip_id", "goods_type", "container_quantities",
            "parcel_quantity_units", "lr_references", "receiver_count",
            "bhada_recorded_inr", "hamali_recorded_inr", "bhada_collected_inr",
            "bhada_outstanding_inr",
        )
        rows = []
        for goods, values in sorted(grouped.items()):
            rows.append((
                business_id, site_id, trip_id, goods,
                "; ".join(f"{kind} x {quantity}"
                          for kind, quantity in sorted(values["containers"].items())),
                values["quantity"], "; ".join(sorted(values["lr_refs"])),
                len(values["receiver_ids"]), format(values["rent"], ".2f"),
                format(values["hamali"], ".2f"), format(values["paid"], ".2f"),
                format(max(values["rent"] - values["paid"], Decimal("0.00")), ".2f"),
            ))
    elif section in ("goods-detail", "goods-details"):
        columns = (
            "business_id", "site_id", "trip_id", "trip_ref", "operating_date",
            "lr_id", "lr_ref", "receiver_label", "receiver_name", "receiver_identifier",
            "sender_name", "goods_type", "container_type", "container_quantity_units",
            "lr_parcel_quantity_units", "lr_financial_totals_row",
            "lr_bhada_recorded_inr", "lr_hamali_recorded_inr",
            "lr_bhada_collected_inr", "lr_bhada_outstanding_inr",
        )
        rows = []
        for lr in lrs:
            rent = _decimal_value(lr.get("rent"))
            paid = lr_payment_totals.get(lr["_id"], Decimal("0.00"))
            outstanding = max(rent - paid, Decimal("0.00"))
            for index, container in enumerate(lr["containers"]):
                financial_row = index == 0
                rows.append((
                    business_id, site_id, trip_id, trip["trip_ref"], trip["operating_date"],
                    lr["_id"], lr["lr_ref"], lr["receiver_label"], lr["receiver_name"],
                    lr.get("receiver_identifier", ""), lr["sender_name"], lr["goods_type"],
                    container["type"], container["quantity"], lr["total_quantity"],
                    "true" if financial_row else "false",
                    format(rent, ".2f") if financial_row and lr.get("rent") is not None else "",
                    format(_decimal_value(lr.get("hamali")), ".2f")
                    if financial_row and lr.get("hamali") is not None else "",
                    format(paid, ".2f") if financial_row else "",
                    format(outstanding, ".2f") if financial_row and lr.get("rent") is not None else "",
                ))
    elif section == "receivers":
        grouped = {}
        for lr in lrs:
            key = lr["receiver_identity_id"]
            group = grouped.setdefault(key, {
                "label": lr["receiver_label"], "name": lr["receiver_name"],
                "lr_refs": [], "goods": [], "containers": [], "quantity": 0,
                "rent": Decimal("0.00"), "hamali": Decimal("0.00"),
                "paid": Decimal("0.00"), "details": [],
            })
            rent = _decimal_value(lr.get("rent"))
            hamali = _decimal_value(lr.get("hamali"))
            paid = lr_payment_totals.get(lr["_id"], Decimal("0.00"))
            group["lr_refs"].append(lr["lr_ref"])
            group["goods"].append(lr["goods_type"])
            group["containers"].extend(
                f"{line['type']} x {line['quantity']}" for line in lr["containers"])
            group["quantity"] += lr["total_quantity"]
            group["rent"] += rent
            group["hamali"] += hamali
            group["paid"] += paid
            for payment in payments_by_lr.get(lr["_id"], []):
                if payment.get("posting_status") == "rejected":
                    continue
                group["details"].append(
                    f"{payment['date']} {payment['method']} {format(_decimal_value(payment['amount']), '.2f')} {payment.get('reference', '')}".strip())
        columns = (
            "business_id", "site_id", "trip_id", "receiver_label", "receiver_name",
            "lr_references", "goods_types", "container_quantities", "parcel_quantity_units",
            "bhada_recorded_inr", "hamali_recorded_inr", "bhada_collected_inr",
            "bhada_outstanding_inr", "payment_details",
        )
        rows = []
        for group in sorted(grouped.values(), key=lambda item: item["label"].casefold()):
            rows.append((
                business_id, site_id, trip_id, group["label"], group["name"],
                "; ".join(group["lr_refs"]), "; ".join(group["goods"]),
                "; ".join(group["containers"]), group["quantity"],
                format(group["rent"], ".2f"), format(group["hamali"], ".2f"),
                format(group["paid"], ".2f"),
                format(max(group["rent"] - group["paid"], Decimal("0.00")), ".2f"),
                "; ".join(group["details"]),
            ))
    elif section == "summary":
        rent = sum((_decimal_value(lr.get("rent")) for lr in lrs), Decimal("0.00"))
        hamali = sum((_decimal_value(lr.get("hamali")) for lr in lrs), Decimal("0.00"))
        paid = sum(lr_payment_totals.values(), Decimal("0.00"))
        columns = (
            "business_id", "site_id", "trip_id", "trip_ref", "site_code",
            "operating_date", "lr_count", "parcel_quantity_units", "bhada_recorded_inr",
            "bhada_collected_inr", "bhada_outstanding_inr", "hamali_recorded_inr",
        )
        rows = [(
            business_id, site_id, trip_id, trip["trip_ref"], site["code"],
            trip["operating_date"], len(lrs), sum(lr["total_quantity"] for lr in lrs),
            format(rent, ".2f"), format(paid, ".2f"),
            format(max(rent - paid, Decimal("0.00")), ".2f"), format(hamali, ".2f"),
        )]
    else:
        columns = LEDGER_COLUMNS
        rows = []
        for lr in lrs:
            events = [p for p in payments_by_lr.get(lr["_id"], [])
                      if p.get("posting_status") != "rejected" and p.get("kind") == "payment"]
            last = events[-1] if events else {}
            rows.append((
                lr["_id"], lr["lr_ref"], trip["_id"], site_id, lr["receiver_label"],
                lr["goods_type"], format(_decimal_value(lr.get("rent")), ".2f")
                if lr.get("rent") is not None else "",
                format(_decimal_value(lr.get("hamali")), ".2f")
                if lr.get("hamali") is not None else "",
                format(lr_payment_totals.get(lr["_id"], Decimal("0.00")), ".2f"),
                last.get("date", ""), last.get("method", ""), last.get("reference", ""),
            ))
    suffix = (
        "import-template" if section == "import-template"
        else "goods-details" if section in ("goods-detail", "goods-details")
        else section
    )
    if section != "import-template":
        return _report_csv_response(
            f"{trip['trip_ref']}-{suffix}.csv", columns, rows,
        )
    return _csv_response(
        f"{trip['trip_ref']}-{suffix}.csv", business_id, site_id, trip_id, columns, rows,
    )


@router.get("/sites/{site_id}/ledger/entries")
async def list_site_ledger_entries(
    site_id: str,
    from_date: Optional[str] = None,
    to_date: Optional[str] = None,
    trip_id: Optional[str] = None,
    trip_ref: Optional[str] = None,
    search: Optional[str] = None,
    limit: int = Query(default=1000, ge=1, le=5000),
    offset: int = Query(default=0, ge=0),
    u=Depends(site_user),
):
    site = await _site_for_user(site_id, u, "lrs:read")
    query = _scoped(str(site["business_id"]), site_id)
    if trip_id:
        query["trip_id"] = trip_id
    clauses = []
    if from_date or to_date:
        dates = {}
        if from_date:
            dates["$gte"] = _validate_iso_date(from_date).isoformat()
        if to_date:
            dates["$lte"] = _validate_iso_date(to_date).isoformat()
        if trip_id:
            clauses.append({"$or": [
                {"operating_date": dates},
                {"operating_date": {"$exists": False}, "receipt_date": dates},
            ]})
        else:
            clauses.append({"$or": [
                {"receipt_date": dates},
                {"receipt_date": {"$exists": False}, "operating_date": dates},
            ]})
    if trip_ref and trip_ref.strip():
        query["trip_ref"] = {"$regex": re.escape(trip_ref.strip()), "$options": "i"}
    term = (search or "").strip()
    if term:
        escaped = re.escape(term)
        clauses.append({"$or": [
            {"sender_name": {"$regex": escaped, "$options": "i"}},
            {"receiver_name": {"$regex": escaped, "$options": "i"}},
        ]})
    if clauses:
        query["$and"] = clauses
    total = await db.site_lrs.count_documents(query)
    rows = await db.site_lrs.find(query).sort(
        [("operating_date", -1), ("sequence", -1)],
    ).skip(offset).limit(limit).to_list(limit)
    by_trip = {}
    for row in rows:
        by_trip.setdefault(row["trip_id"], []).append(row)
    trips = {}
    for trip_id in by_trip:
        trip = await db.site_trips.find_one(_scoped(
            str(site["business_id"]), site_id, _id=trip_id,
        ))
        if trip:
            trips[trip_id] = trip
    result = []
    for trip_id, trip_lrs in by_trip.items():
        trip = trips.get(trip_id)
        if not trip:
            continue
        enriched = await _lr_rows(site, trip_id, trip_lrs, u)
        for entry in enriched:
            entry["site_name"] = site.get("name", "")
            entry["trip_ref"] = trip.get("trip_ref", "")
            entry["trip_status"] = trip.get("status", "open")
            entry["truck_no"] = trip.get("truck_no", "")
            entry["driver_name"] = trip.get("driver_name", "")
            entry["total_rent"] = format(
                _decimal_value(entry.get("rent")) + _decimal_value(entry.get("hamali"))
                + _decimal_value(entry.get("receipt_fee")), ".2f",
            ) if _can_view_site_lr_charges(u, site_id) else None
            result.append(entry)
    result.sort(key=lambda row: (row.get("receipt_date", row.get("operating_date", "")),
                                 row.get("sequence", 0)), reverse=True)
    return {"rows": result, "total": total, "limit": limit, "offset": offset}


@router.get("/sites/{site_id}/ledger/export")
async def export_site_ledger_csv(
    site_id: str,
    from_date: str = Query(...),
    to_date: str = Query(...),
    u=Depends(site_user),
):
    site = await _site_for_user(site_id, u, "lrs:read")
    if u.get("role") == "site_manager" and "ledger:export" not in (
        (u.get("site_permissions") or {}).get(site_id, [])
    ):
        raise HTTPException(403, "Ledger export permission is required for this site")
    start = _validate_iso_date(from_date).isoformat()
    end = _validate_iso_date(to_date).isoformat()
    if start > end:
        raise HTTPException(422, "Start date must be before or equal to end date")
    date_range = {"$gte": start, "$lte": end}
    query = _scoped(str(site["business_id"]), site_id,
                    **{"$or": [
                        {"receipt_date": date_range},
                        {"receipt_date": {"$exists": False}, "operating_date": date_range},
                    ]})
    lrs = await db.site_lrs.find(query).sort(
        [("operating_date", 1), ("sequence", 1)],
    ).limit(10001).to_list(10001)
    if len(lrs) > 10000:
        raise HTTPException(413, "This ledger contains more than 10,000 receipts; export a shorter period")
    trips = {}
    for trip_id in {row["trip_id"] for row in lrs}:
        trip = await db.site_trips.find_one(_scoped(str(site["business_id"]), site_id, _id=trip_id))
        if trip:
            trips[trip_id] = trip
    columns = (
        "Date", "Trip number", "LR number", "Sender", "Sender address",
        "Receiver", "Receiver address", "Goods", "Quantity", "Bhada",
        "Hamali", "Receipt fee", "Total charges", "Status",
    )
    rows = []
    for lr in lrs:
        trip = trips.get(lr["trip_id"])
        if not trip:
            continue
        goods = lr.get("goods_rows") or lr.get("containers") or []
        rent = _decimal_value(lr.get("rent"))
        hamali = _decimal_value(lr.get("hamali"))
        receipt_fee = _decimal_value(lr.get("receipt_fee", Decimal128(Decimal("2.00"))))
        can_read_finances = _can_view_site_lr_charges(u, site_id)
        rows.append((
            lr.get("receipt_date", lr.get("operating_date", "")), trip.get("trip_ref", ""),
            lr.get("lr_ref", ""), lr.get("sender_name", ""), lr.get("sender_address", ""),
            lr.get("receiver_name", ""), lr.get("receiver_address", ""),
            "; ".join(line.get("type", "") for line in goods),
            sum(line.get("quantity", 0) for line in goods),
            format(rent, ".2f") if can_read_finances else "",
            format(hamali, ".2f") if can_read_finances else "",
            format(receipt_fee, ".2f") if can_read_finances else "",
            format(rent + hamali + receipt_fee, ".2f") if can_read_finances else "",
            "VOID" if lr.get("voided") else "Active",
        ))
    return _report_csv_response(f"{site.get('code', site_id)}-ledger-{start}-{end}.csv", columns, rows)


@router.get("/sites/{site_id}/ledger/daily")
async def export_site_daily_ledger(
    site_id: str, operating_date: str = Query(...), u=Depends(site_user),
):
    site = await _site_for_user(site_id, u)
    if u.get("role") == "site_manager" and "ledger:export" not in (
        (u.get("site_permissions") or {}).get(site_id, [])
    ):
        raise HTTPException(403, "Daily ledger export permission is required for this site")

    day = _validate_iso_date(operating_date).isoformat()
    business_id = str(site["business_id"])
    trip_query = _scoped(business_id, site_id, operating_date=day)
    trips = await db.site_trips.find(trip_query).sort("sequence", 1).limit(2001).to_list(2001)
    if len(trips) > 2000:
        raise HTTPException(413, "This site has more than 2,000 bookings for the selected day")

    trip_by_id = {trip["_id"]: trip for trip in trips}
    trip_ids = list(trip_by_id)
    lrs = []
    payments = []
    if trip_ids:
        lrs = await db.site_lrs.find(_scoped(
            business_id, site_id, trip_id={"$in": trip_ids},
        )).sort([("trip_id", 1), ("sequence", 1)]).limit(10001).to_list(10001)
        if len(lrs) > 10000:
            raise HTTPException(413, "This daily ledger has more than 10,000 LRs; export a smaller day")
        lrs.sort(key=lambda lr: (
            trip_by_id.get(lr.get("trip_id"), {}).get("sequence", 0),
            lr.get("sequence", 0),
        ))
        lr_ids = [lr["_id"] for lr in lrs]
        if lr_ids:
            payments = await db.site_payments.find(_scoped(
                business_id, site_id, trip_id={"$in": trip_ids}, lr_id={"$in": lr_ids},
                posting_status={"$ne": "rejected"},
            )).sort("created_at", 1).limit(50001).to_list(50001)
            if len(payments) > 50000:
                raise HTTPException(413, "This daily ledger has more than 50,000 payment events")

    payments_by_lr = {}
    totals_by_lr = {}
    for payment in payments:
        lr_id = payment["lr_id"]
        amount = _decimal_value(payment.get("amount"))
        if payment.get("kind") == "reversal":
            amount = -amount
        totals_by_lr[lr_id] = totals_by_lr.get(lr_id, Decimal("0.00")) + amount
        payments_by_lr.setdefault(lr_id, []).append(
            f"{payment.get('date', '')} {payment.get('method', '')} "
            f"{payment.get('kind', 'payment')} {format(amount, '.2f')} "
            f"{payment.get('reference', '')}".strip()
        )

    columns = (
        "operating_date", "site_name", "site_code", "trip_ref", "truck_no", "driver_name",
        "lr_ref", "sender_name", "receiver_name", "receiver_identifier", "goods_type",
        "container_quantities", "parcel_quantity_units", "bhada_recorded_inr",
        "hamali_recorded_inr", "receipt_fee_inr", "bhada_collected_inr", "bhada_outstanding_inr",
        "payment_status", "payment_details",
    )
    rows = []
    for lr in lrs:
        trip = trip_by_id.get(lr.get("trip_id"))
        if not trip:
            continue
        paid = totals_by_lr.get(lr["_id"], Decimal("0.00"))
        rent = _decimal_value(lr.get("rent"))
        receipt_fee = _decimal_value(lr.get("receipt_fee", Decimal128(Decimal("2.00"))))
        rows.append((
            day, site.get("name", ""), site.get("code", ""),
            trip.get("trip_ref", ""), trip.get("truck_no", ""), trip.get("driver_name", ""),
            lr.get("lr_ref", ""), lr.get("sender_name", ""),
            lr.get("receiver_name", ""), lr.get("receiver_identifier", ""),
            lr.get("goods_type", ""),
            "; ".join(f"{line.get('type', '')} x {line.get('quantity', 0)}"
                      for line in lr.get("containers", [])),
            lr.get("total_quantity", 0),
            format(rent, ".2f") if lr.get("rent") is not None else "",
            format(_decimal_value(lr.get("hamali")), ".2f")
            if lr.get("hamali") is not None else "",
            format(receipt_fee, ".2f"),
            format(paid, ".2f"),
            format(max(rent + receipt_fee - paid, Decimal("0.00")), ".2f")
            if lr.get("rent") is not None else "",
            _payment_status(
                rent + receipt_fee if lr.get("rent") is not None else None, paid,
            ),
            "; ".join(payments_by_lr.get(lr["_id"], [])),
        ))
    return _report_csv_response(
        f"{site.get('code', site_id)}-daily-ledger-{day}.csv", columns, rows,
    )


@router.get("/sites/{site_id}/ledger/export.xlsx")
async def export_site_ledger_xlsx(
    site_id: str,
    from_date: str = Query(...),
    to_date: str = Query(...),
    trip_id: Optional[str] = None,
    u=Depends(site_user),
):
    site = await _site_for_user(site_id, u, "lrs:read")
    if u.get("role") == "site_manager" and "ledger:export" not in (
        (u.get("site_permissions") or {}).get(site_id, [])
    ):
        raise HTTPException(403, "Ledger export permission is required for this site")
    start = _validate_iso_date(from_date).isoformat()
    end = _validate_iso_date(to_date).isoformat()
    if start > end:
        raise HTTPException(422, "Start date must be before or equal to end date")
    date_range = {"$gte": start, "$lte": end}
    date_clause = {"$or": [
        {"operating_date": date_range},
        {"operating_date": {"$exists": False}, "receipt_date": date_range},
    ]} if trip_id else {"$or": [
        {"receipt_date": date_range},
        {"receipt_date": {"$exists": False}, "operating_date": date_range},
    ]}
    query = _scoped(
        str(site["business_id"]), site_id, **({"trip_id": trip_id} if trip_id else {}),
        **date_clause,
    )
    lrs = await db.site_lrs.find(query).sort(
        [("operating_date", 1), ("sequence", 1)],
    ).limit(10001).to_list(10001)
    if len(lrs) > 10000:
        raise HTTPException(413, "This ledger contains more than 10,000 receipts; export a shorter period")
    trips = {}
    for trip_id in {row["trip_id"] for row in lrs}:
        trip = await db.site_trips.find_one(_scoped(str(site["business_id"]), site_id, _id=trip_id))
        if trip:
            trips[trip_id] = trip

    can_read_finances = _can_view_site_lr_charges(u, site_id)
    paid_by_trip = {}
    if u.get("role") == "owner":
        for ledger_trip_id in trips:
            paid_by_trip[ledger_trip_id] = await _payments_by_lr(
                site, ledger_trip_id,
                [row["_id"] for row in lrs if row.get("trip_id") == ledger_trip_id],
            )
    ledger_rows = []
    totals = {"rent": Decimal("0.00"), "hamali": Decimal("0.00"),
              "receipt_fee": Decimal("0.00")}
    for lr in lrs:
        trip = trips.get(lr["trip_id"])
        if not trip:
            continue
        goods = lr.get("goods_rows") or lr.get("containers") or []
        rent = _decimal_value(lr.get("rent"))
        hamali = _decimal_value(lr.get("hamali"))
        receipt_fee = _decimal_value(lr.get("receipt_fee", Decimal128(Decimal("2.00"))))
        if not lr.get("voided"):
            totals["rent"] += rent
            totals["hamali"] += hamali
            totals["receipt_fee"] += receipt_fee
        row_values = (
            lr.get("receipt_date", lr.get("operating_date", "")),
            trip.get("trip_ref", ""), lr.get("lr_ref", ""),
            _ledger_hindi_text(lr.get("sender_name_hindi")),
            _ledger_hindi_text(lr.get("receiver_name_hindi")),
            ", ".join(
                " - ".join(part for part in (
                    _hindi_digits(line.get("quantity", 0)),
                    "/".join(goods_part for goods_part in (
                        _ledger_hindi_text(line.get("type_hindi") or line.get("type")),
                        _ledger_hindi_text(line.get("description_hindi") or line.get("description")),
                    ) if goods_part),
                ) if part)
                for line in goods
            ),
            rent if can_read_finances else "",
            hamali if can_read_finances else "",
            receipt_fee,
            _entry_by(lr),
        )
        if u.get("role") == "owner":
            paid = paid_by_trip.get(lr["trip_id"], {}).get(lr["_id"], Decimal("0.00"))
            row_values += (
                "✓" if not lr.get("voided") and lr.get("rent") is not None
                and paid >= rent + receipt_fee else "",
            )
        ledger_rows.append(row_values)

    total_row = (
        "कुल (रद्द रसीद छोड़कर)", "", "", "", "", "",
        totals["rent"] if can_read_finances else "",
        totals["hamali"] if can_read_finances else "",
        totals["receipt_fee"],
        f"Grand total: ₹{totals['rent'] + totals['hamali'] + totals['receipt_fee']:.2f}"
        if can_read_finances else "",
    )
    if u.get("role") == "owner":
        total_row += ("",)
    ledger_rows.append(total_row)

    filename = f"{site.get('code', site_id)}-ledger-{start}-{end}.xlsx"
    columns = (
        "दिनांक", "ट्रिप नंबर", "रसीद नंबर", "भेजने वाला", "प्राप्तकर्ता",
        "माल एवं मात्रा", "भाड़ा", "हमाली", "रसीद शुल्क", "प्रविष्टि करने वाला",
    )
    widths = (14, 24, 24, 24, 24, 48, 14, 14, 14, 24)
    if u.get("role") == "owner":
        columns += ("Amount paid",)
        widths += (14,)
    return _xlsx_response(
        filename, site.get("name", "Booking ledger"), f"{start} to {end}",
        [
            ("बही", columns, ledger_rows, widths),
        ],
    )


def _parse_ledger_csv(content, site, trip):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise HTTPException(422, "Ledger CSV must be UTF-8 encoded")
    reader = csv.reader(io.StringIO(text, newline=""))
    try:
        marker = next(reader)
        scope = next(reader)
        columns = next(reader)
    except (csv.Error, StopIteration):
        raise HTTPException(422, "Ledger CSV metadata is malformed")
    if marker != ["#FLEET_MANAGER_TRIP_LEDGER", LEDGER_TEMPLATE_VERSION]:
        raise HTTPException(422, "Unsupported ledger template version")
    expected_scope = {
        key: value for key, value in (part.split("=", 1) for part in scope if "=" in part)
    }
    if expected_scope != {
        "business_id": str(site["business_id"]), "site_id": site["_id"], "trip_id": trip["_id"],
    }:
        raise HTTPException(422, "Ledger file business, site, or trip scope does not match this upload")
    if len(columns) != len(set(columns)):
        raise HTTPException(422, "Ledger CSV has duplicate column names")
    if set(columns) != set(LEDGER_COLUMNS):
        raise HTTPException(422, "Ledger CSV columns do not match the supported template")
    rows = []
    try:
        for row_number, values in enumerate(reader, start=1):
            if not values:
                continue
            if len(values) != len(columns):
                rows.append({
                    "row_number": row_number, "source": {}, "skip": False,
                    "errors": ["Row has a different number of values than the template columns"],
                })
            elif all(not value.strip() for value in values):
                continue
            else:
                row = dict(zip(columns, values))
                rows.append({
                    "row_number": row_number,
                    "source": {key: row[key].strip() for key in LEDGER_COLUMNS},
                    "skip": False, "errors": [],
                })
            if len(rows) > 5000:
                raise HTTPException(413, "Ledger CSV cannot contain more than 5,000 rows")
    except csv.Error:
        raise HTTPException(422, "Ledger CSV could not be parsed")
    if not rows:
        raise HTTPException(422, "Ledger CSV contains no LR rows")
    return rows


async def _validate_import_rows(site, trip, raw_rows):
    business_id = str(site["business_id"])
    ids = {r.get("source", {}).get("lr_id") for r in raw_rows if not r.get("skip")}
    ids.discard(None)
    docs = await db.site_lrs.find(_scoped(
        business_id, site["_id"], trip_id=trip["_id"], _id={"$in": list(ids)},
    )).to_list(5000) if ids else []
    lr_by_id = {lr["_id"]: lr for lr in docs}
    paid_map = await _payments_by_lr(site, trip["_id"], list(lr_by_id))
    payment_detail_rows = await db.site_payments.aggregate([
        {"$match": _scoped(
            business_id, site["_id"], trip_id=trip["_id"], lr_id={"$in": list(lr_by_id)},
            posting_status={"$ne": "rejected"}, kind="payment",
        )},
        {"$sort": {"created_at": 1}},
        {"$group": {"_id": "$lr_id", "date": {"$last": "$date"},
                    "method": {"$last": "$method"}, "reference": {"$last": "$reference"}}},
    ]).to_list(5000) if lr_by_id else []
    latest_payment = {row["_id"]: row for row in payment_detail_rows}
    seen = set()
    checked = []
    for raw in raw_rows:
        item = {"row_number": raw["row_number"], "source": raw.get("source", {}),
                "skip": bool(raw.get("skip")), "errors": []}
        if item["skip"]:
            item["errors"] = []
            item["skipped"] = True
            checked.append(item)
            continue
        source = item["source"]
        if not source:
            item["errors"].append("Correct or skip this malformed CSV row")
            checked.append(item)
            continue
        if source.get("site_id") != site["_id"] or source.get("trip_id") != trip["_id"]:
            item["errors"].append("Row belongs to a different site or trip")
        lr_id = source.get("lr_id", "")
        lr = lr_by_id.get(lr_id)
        if not lr:
            item["errors"].append("Unknown LR ID or LR belongs to another site/trip")
        elif lr_id in seen:
            item["errors"].append("Duplicate LR row")
        else:
            seen.add(lr_id)
            if source.get("lr_ref") != lr["lr_ref"]:
                item["errors"].append("LR reference does not match the immutable LR ID")
        if lr:
            item["matched_lr_ref"] = lr["lr_ref"]
            item["matched_receiver"] = lr["receiver_label"]
            item["current_rent"] = format(_decimal_value(lr.get("rent")), ".2f") if lr.get("rent") is not None else ""
            item["current_hamali"] = format(_decimal_value(lr.get("hamali")), ".2f") if lr.get("hamali") is not None else ""
            item["current_paid_total"] = format(paid_map.get(lr_id, Decimal("0.00")), ".2f")
            values = {}
            for key, label in (("rent", "Bhada"), ("hamali", "Hamali"),
                               ("paid_total", "Paid total")):
                try:
                    values[key] = _decimal(source.get(key, ""), label, allow_none=True)
                except HTTPException as exc:
                    item["errors"].append(exc.detail)
                    values[key] = None
            rent = values.get("rent")
            current_paid = paid_map.get(lr_id, Decimal("0.00"))
            requested_paid = values.get("paid_total")
            if rent is not None and requested_paid is not None and rent < requested_paid:
                item["errors"].append("Bhada cannot be less than the paid total")
            if requested_paid is not None and requested_paid < current_paid:
                item["errors"].append("Paid total cannot be reduced; use an auditable payment reversal")
            if requested_paid is not None and requested_paid > current_paid:
                if not source.get("payment_date"):
                    item["errors"].append("Payment date is required for a new payment")
                else:
                    try:
                        _validate_iso_date(source["payment_date"])
                    except HTTPException as exc:
                        item["errors"].append(exc.detail)
                if source.get("payment_method") not in ("Cash", "UPI", "Bank", "Cheque", "Other"):
                    item["errors"].append("Choose a supported payment method for the new payment")
            elif requested_paid is not None and lr_id in latest_payment:
                previous = latest_payment[lr_id]
                supplied = (
                    ("payment_date", "date"),
                    ("payment_method", "method"),
                    ("transaction_reference", "reference"),
                )
                if any(source.get(key, "") and source.get(key, "") != str(previous.get(value) or "")
                       for key, value in supplied):
                    item["errors"].append(
                        "Existing payment details are immutable; record a new payment or an owner reversal"
                    )
            item["changes"] = {
                "rent": format(rent, ".2f") if rent is not None else None,
                "hamali": format(values["hamali"], ".2f") if values["hamali"] is not None else None,
                "paid_total": format(requested_paid, ".2f") if requested_paid is not None else None,
            }
        checked.append(item)
    missing = sorted(set(lr_by_id) - seen)
    for lr_id in missing:
        lr = lr_by_id[lr_id]
        checked.append({
            "row_number": None, "source": {"lr_id": lr_id},
            "skip": True, "skipped": True, "missing": True,
            "errors": [f"LR {lr['lr_ref']} is missing from the upload"],
        })
    has_row_errors = any(row.get("errors") and not row.get("missing") and not row.get("skip")
                         for row in checked)
    return checked, {"missing_lr_ids": missing, "has_errors": has_row_errors,
                     "valid_rows": sum(1 for row in checked if row.get("changes") and not row["errors"]),
                     "skipped_rows": sum(1 for row in checked if row.get("skip"))}


def _import_response(doc: Optional[dict[str, Any]]) -> dict[str, Any]:
    out = _doc(doc)
    if out is None:
        raise HTTPException(404, "Ledger import not found")
    acknowledged = set(out.get("acknowledged_missing_lr_ids") or [])
    missing = set((out.get("summary") or {}).get("missing_lr_ids") or [])
    out["ready_to_commit"] = (
        out["status"] in ("preview", "applying")
        and not any(row.get("errors") and not row.get("missing") and not row.get("skip")
                    for row in out.get("rows", []))
        and missing.issubset(acknowledged)
    )
    return out


@router.post("/sites/{site_id}/trips/{trip_id}/ledger-imports")
async def preview_site_ledger_import(
    site_id: str, trip_id: str, file: UploadFile = File(...), u=Depends(site_user),
):
    _owner(u)
    site = await _site_for_user(site_id, u)
    trip = await _trip(site, trip_id)
    if file.filename is None or not file.filename.lower().endswith(".csv"):
        raise HTTPException(415, "Upload a CSV file generated from this trip's ledger export")
    content = await file.read(MAX_LEDGER_BYTES + 1)
    if len(content) > MAX_LEDGER_BYTES:
        raise HTTPException(413, "Ledger CSV must be 2 MB or smaller")
    if not content:
        raise HTTPException(422, "Ledger CSV is empty")
    digest = hashlib.sha256(content).hexdigest()
    business_id = str(site["business_id"])
    existing = await db.site_ledger_imports.find_one(_scoped(
        business_id, site_id, trip_id=trip_id, file_hash=digest,
    ))
    if existing:
        existing["already_processed"] = existing.get("status") == "applied"
        return _import_response(existing)
    safe_name = re.sub(r"[^A-Za-z0-9._-]", "_", file.filename)[:120] or "ledger.csv"
    storage_path = f"fleet-manager/{business_id}/site-ledgers/{new_id()}.csv"
    await run_in_threadpool(S.put_object, storage_path, content, "text/csv")
    fatal = ""
    try:
        raw_rows = _parse_ledger_csv(content, site, trip)
        checked, summary = await _validate_import_rows(site, trip, raw_rows)
        status = "preview" if not summary["has_errors"] else "needs_review"
    except HTTPException as exc:
        checked, summary, fatal, status = [], {
            "missing_lr_ids": [], "has_errors": True, "valid_rows": 0, "skipped_rows": 0,
        }, str(exc.detail), "needs_review"
    imported = {
        "_id": new_id(), "business_id": business_id, "site_id": site_id,
        "trip_id": trip_id, "file_hash": digest, "template_version": LEDGER_TEMPLATE_VERSION,
        "original_filename": safe_name, "original_path": storage_path, "size": len(content),
        "uploader": u["username"], "uploaded_at": now_iso(), "status": status,
        "rows": checked, "summary": summary, "fatal_error": fatal,
        "preview_decisions": [], "created_at": now_iso(),
    }
    try:
        await db.site_ledger_imports.insert_one(imported)
    except DuplicateKeyError:
        existing = await db.site_ledger_imports.find_one(_scoped(
            business_id, site_id, trip_id=trip_id, file_hash=digest,
        ))
        return _import_response(existing)
    await _audit(site, u, "ledger.import_previewed", "ledger_import", imported["_id"],
                 new={"file_hash": digest, "rows": len(checked), "status": status,
                      "fatal_error": fatal})
    return _import_response(imported)


@router.get("/sites/{site_id}/trips/{trip_id}/ledger-imports/{import_id}")
async def get_site_ledger_import(site_id: str, trip_id: str, import_id: str,
                                u=Depends(site_user)):
    _owner(u)
    site = await _site_for_user(site_id, u)
    await _trip(site, trip_id)
    imported = await db.site_ledger_imports.find_one(_scoped(
        str(site["business_id"]), site_id, trip_id=trip_id, _id=import_id,
    ))
    if not imported:
        raise HTTPException(404, "Ledger import not found")
    return _import_response(imported)


@router.put("/sites/{site_id}/trips/{trip_id}/ledger-imports/{import_id}/preview")
async def update_site_ledger_preview(
    site_id: str, trip_id: str, import_id: str, body: ImportPreviewEdit,
    u=Depends(site_user),
):
    _owner(u)
    site = await _site_for_user(site_id, u)
    trip = await _trip(site, trip_id)
    imported = await db.site_ledger_imports.find_one(_scoped(
        str(site["business_id"]), site_id, trip_id=trip_id, _id=import_id,
    ))
    if not imported:
        raise HTTPException(404, "Ledger import not found")
    if imported["status"] not in ("needs_review", "preview"):
        raise HTTPException(409, "Only an import awaiting review can be edited")
    edited = [row.model_dump() for row in body.rows]
    row_numbers = [row["row_number"] for row in edited]
    if len(row_numbers) != len(set(row_numbers)):
        raise HTTPException(422, "CSV row numbers must be unique")
    original_by_number = {row["row_number"]: row for row in imported["rows"] if row.get("row_number")}
    rows = []
    for row in edited:
        original = original_by_number.get(row["row_number"])
        if not original:
            raise HTTPException(422, "Preview cannot add rows not present in the uploaded file")
        source = dict(original.get("source") or {})
        for key in ("lr_id", "rent", "hamali", "paid_total",
                    "payment_date", "payment_method", "transaction_reference"):
            if row.get(key) is not None:
                source[key] = row[key]
        rows.append({"row_number": row["row_number"], "source": source,
                     "skip": row["skip"], "errors": []})
    checked, summary = await _validate_import_rows(site, trip, rows)
    status = "preview" if not summary["has_errors"] else "needs_review"
    await db.site_ledger_imports.update_one(
        {"_id": import_id, "status": imported["status"]},
        {"$set": {"rows": checked, "summary": summary, "status": status,
                  "acknowledged_missing_lr_ids": [],
                  "preview_decisions": [{"at": now_iso(), "by": u["username"],
                                         "rows": [r["row_number"] for r in edited]}]},
         "$unset": {"fatal_error": ""}},
    )
    await _audit(site, u, "ledger.import_preview_edited", "ledger_import", import_id,
                 new={"status": status, "valid_rows": summary["valid_rows"],
                      "skipped_rows": summary["skipped_rows"]})
    return _import_response(await db.site_ledger_imports.find_one({"_id": import_id}))


async def _apply_import_amount(imported, site, trip, lr, field, target, user):
    if target is None:
        return
    business_id = str(site["business_id"])
    effect_id = f"ledger:{imported['_id']}:{lr['_id']}:{field}"
    latest = await db.site_lrs.find_one(_scoped(
        business_id, site["_id"], _id=lr["_id"], trip_id=trip["_id"],
    ))
    if not latest:
        raise HTTPException(409, f"LR {lr['lr_ref']} no longer exists in this trip")
    prior_effect = next(
        (effect for effect in latest.get("ledger_import_effects", [])
         if isinstance(effect, dict) and effect.get("effect_id") == effect_id),
        None,
    )
    current = _decimal_value(latest.get(field)) if latest.get(field) is not None else Decimal("0.00")
    delta = Decimal(prior_effect["delta"]) if prior_effect else target - current
    marker_query = _scoped(
        business_id, site["_id"], _id=lr["_id"], trip_id=trip["_id"],
        ledger_import_effects={"$not": {"$elemMatch": {"effect_id": effect_id}}},
    )
    if not prior_effect:
        effect = {
            "effect_id": effect_id, "from": format(current, ".2f"),
            "to": format(target, ".2f"), "delta": format(delta, ".2f"),
        }
        update = await db.site_lrs.update_one(
            {**marker_query, field: latest.get(field)},
            {"$set": {field: Decimal128(target)},
             "$addToSet": {"ledger_import_effects": effect}},
        )
        if update.modified_count != 1:
            latest = await db.site_lrs.find_one(_scoped(
                business_id, site["_id"], _id=lr["_id"], trip_id=trip["_id"],
            ))
            prior_effect = next(
                (item for item in (latest or {}).get("ledger_import_effects", [])
                 if isinstance(item, dict) and item.get("effect_id") == effect_id),
                None,
            )
            if not prior_effect:
                raise HTTPException(409, f"LR {lr['lr_ref']} changed during reconciliation")
            delta = Decimal(prior_effect["delta"])
    if field == "rent" and delta:
        await _financial_event(
            effect_id, site, lr["party_id"], lr["operating_date"],
            f"Ledger reconciliation {lr['lr_ref']} bhada adjustment",
            delta, user=u,
        )
    audit_old = prior_effect["from"] if prior_effect else format(current, ".2f")
    await _audit(
        site, user, "ledger.lr_amount_reconciled", "lr", lr["_id"],
        old={field: audit_old},
        new={field: format(target, ".2f"), "import_id": imported["_id"]},
        event_id=f"ledger-import-audit:{imported['_id']}:{lr['_id']}:{field}",
    )


@router.post("/sites/{site_id}/trips/{trip_id}/ledger-imports/{import_id}/commit")
async def commit_site_ledger_import(
    site_id: str, trip_id: str, import_id: str, body: ImportCommit,
    u=Depends(site_user),
):
    _owner(u)
    if not body.confirm:
        raise HTTPException(422, "Confirm the ledger changes before applying them")
    site = await _site_for_user(site_id, u)
    trip = await _trip(site, trip_id)
    _ensure_trip_ledger_open(trip)
    business_id = str(site["business_id"])
    scope = _scoped(business_id, site_id, trip_id=trip_id, _id=import_id)
    imported = await db.site_ledger_imports.find_one(scope)
    if not imported:
        raise HTTPException(404, "Ledger import not found")
    if imported["status"] == "applied":
        return _import_response(imported)
    if imported["status"] not in ("preview", "applying"):
        raise HTTPException(409, "Resolve all ledger preview errors before applying")
    if any(row.get("errors") and not row.get("missing") and not row.get("skip")
           for row in imported.get("rows", [])):
        raise HTTPException(409, "Resolve all row errors before applying")
    if imported["status"] == "preview":
        raw_rows = [
            {"row_number": row["row_number"], "source": row.get("source", {}),
             "skip": row.get("skip", False), "errors": []}
            for row in imported.get("rows", []) if row.get("row_number") is not None
        ]
        fresh_rows, fresh_summary = await _validate_import_rows(site, trip, raw_rows)
        if fresh_summary["has_errors"]:
            await db.site_ledger_imports.update_one(
                {**scope, "status": "preview"},
                {"$set": {"rows": fresh_rows, "summary": fresh_summary,
                          "status": "needs_review", "preview_refreshed_at": now_iso()}},
            )
            raise HTTPException(409, "Ledger data changed since preview; review the refreshed discrepancies")
        await db.site_ledger_imports.update_one(
            {**scope, "status": "preview"},
            {"$set": {"rows": fresh_rows, "summary": fresh_summary,
                      "preview_refreshed_at": now_iso()}},
        )
        imported["rows"] = fresh_rows
        imported["summary"] = fresh_summary
    missing_ids = set((imported.get("summary") or {}).get("missing_lr_ids") or [])
    acknowledged_missing = set(body.acknowledged_missing_lr_ids)
    if not missing_ids.issubset(acknowledged_missing) or acknowledged_missing - missing_ids:
        raise HTTPException(422, "Explicitly acknowledge each missing LR, without adding unknown IDs")
    if imported["status"] == "preview":
        claimed = await db.site_ledger_imports.update_one(
            {**scope, "status": "preview"},
            {"$set": {"status": "applying", "apply_started_at": now_iso(),
                      "confirmed_by": u["username"],
                      "acknowledged_missing_lr_ids": sorted(acknowledged_missing)}},
        )
        if claimed.modified_count != 1:
            imported = await db.site_ledger_imports.find_one(scope)
            if imported and imported["status"] == "applied":
                return _import_response(imported)
            if not imported or imported["status"] != "applying":
                raise HTTPException(409, "Ledger import state changed; reload its preview")
    imported = await db.site_ledger_imports.find_one(scope)
    if not imported:
        raise HTTPException(404, "Ledger import not found")
    for row in imported.get("rows", []):
        if row.get("skip") or row.get("missing") or row.get("apply_status") == "applied":
            continue
        source = row.get("source") or {}
        lr_id = source.get("lr_id")
        lr = await db.site_lrs.find_one(_scoped(
            business_id, site_id, _id=lr_id, trip_id=trip_id,
        ))
        if not lr:
            raise HTTPException(409, f"LR {lr_id} changed after preview; review the import again")
        changes = row.get("changes") or {}
        rent = _decimal(changes.get("rent"), "Bhada", allow_none=True)
        hamali = _decimal(changes.get("hamali"), "Hamali", allow_none=True)
        target_paid = _decimal(changes.get("paid_total"), "Paid total", allow_none=True)
        if rent is not None and target_paid is not None and rent < target_paid:
            raise HTTPException(409, f"LR {lr['lr_ref']} bhada is now lower than the paid total")
        await _apply_import_amount(imported, site, trip, lr, "rent", rent, u)
        await _apply_import_amount(imported, site, trip, lr, "hamali", hamali, u)
        if target_paid is not None:
            paid_now = await _payments_by_lr(site, trip_id, [lr_id])
            current_paid = paid_now.get(lr_id, Decimal("0.00"))
            if target_paid < current_paid:
                raise HTTPException(409, f"LR {lr['lr_ref']} paid total changed after preview")
            if target_paid > current_paid:
                await _record_lr_payment(
                    site, trip, lr, target_paid - current_paid,
                    source.get("payment_date"), source.get("payment_method"),
                    source.get("transaction_reference", ""),
                    f"ledger_import:{import_id}:{lr_id}", u,
                )
        await db.site_lrs.update_one(
            _scoped(business_id, site_id, _id=lr_id, trip_id=trip_id),
            {"$set": {"reconciled": True, "reconciled_by": u["username"],
                      "reconciled_at": now_iso(), "ledger_import_id": import_id}},
        )
        await db.site_ledger_imports.update_one(
            {**scope, "rows.row_number": row.get("row_number")},
            {"$set": {"rows.$.apply_status": "applied",
                      "rows.$.applied_at": now_iso()}},
        )
    unverified = await db.site_lrs.count_documents(_scoped(
        business_id, site_id, trip_id=trip_id, reconciled={"$ne": True},
    ))
    await db.site_trips.update_one(
        _scoped(business_id, site_id, _id=trip_id),
        {"$set": {"reconciled": unverified == 0, "updated_at": now_iso()}},
    )
    await _audit(site, u, "ledger.import_applied", "ledger_import", import_id,
                 new={"unreconciled_lrs": unverified,
                      "valid_rows": imported.get("summary", {}).get("valid_rows", 0)},
                 event_id=f"ledger-import-applied-audit:{import_id}")
    await db.site_ledger_imports.update_one(
        {**scope, "status": "applying"},
        {"$set": {"status": "applied", "applied_by": u["username"],
                  "applied_at": now_iso(), "unreconciled_lrs": unverified}},
    )
    return _import_response(await db.site_ledger_imports.find_one(scope))


@router.get("/sites/{site_id}/ledger-imports/{import_id}/original")
async def download_site_ledger_original(site_id: str, import_id: str, u=Depends(site_user)):
    _owner(u)
    site = await _site_for_user(site_id, u)
    imported = await db.site_ledger_imports.find_one(_scoped(
        str(site["business_id"]), site_id, _id=import_id,
    ))
    if not imported:
        raise HTTPException(404, "Ledger import not found")
    try:
        content, media_type = await run_in_threadpool(S.get_object, imported["original_path"])
    except FileNotFoundError:
        raise HTTPException(404, "Original ledger file is no longer available on this server")
    return Response(
        content=content, media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{imported["original_filename"]}"',
                 "Cache-Control": "no-store"},
    )
